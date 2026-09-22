import shutil
import tempfile
import json
import threading
import unittest
import urllib.request
from http.server import HTTPServer
from pathlib import Path
from unittest.mock import patch

from scripts import tenant_store
from scripts import connections_registry
from scripts import jecs_api_server
from scripts.connections_registry import list_connections, secret_reference
from scripts.digits_ingest import parse_digits_export
from scripts.jecs_api_server import (
    _public_features,
    _set_feature_override,
    build_fleet_cost_reconciliation,
    build_connections_payload,
    build_dashboard_operations_payload,
)

ROOT = Path(__file__).resolve().parents[1]
EXPORT = ROOT / "data/fleet_reviews/2026-09-07/three-month-reconciliation/JEC-June-August-Rental-Reconciliation.xlsx"


def _digits_bytes():
    """Rebuild a tenant-style Digits export from the reviewed charge rows."""
    from openpyxl import load_workbook, Workbook
    import io
    sheet = load_workbook(EXPORT, data_only=True)["All Digits charges"]
    rows = [r for r in sheet.iter_rows(values_only=True) if any(v is not None for v in r)]
    book = Workbook()
    for row in rows:
        book.active.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class TenantUploadTests(unittest.TestCase):
    def test_dashboard_operations_uses_current_connected_sources(self):
        payload = build_dashboard_operations_payload('jecs')
        available_weeks = sorted(
            folder.name for folder in (ROOT / 'data/scorecard_data').glob('2026-wk??')
            if any(folder.glob('DSP_Overview_Dashboard_*.csv'))
        )
        self.assertEqual(payload['performance']['period'], available_weeks[-1])
        self.assertGreater(len(payload['performance']['drivers']), 0)
        self.assertEqual(payload['fleet']['summary']['registeredFleet'], 50)
        self.assertEqual(payload['connections']['summary']['connectionTotal'], 5)
        self.assertTrue(payload['sources'])
        self.assertTrue(all('feeds' in source for source in payload['sources']
                            if source['id'] not in ('amazon', 'fleet', 'accounting')))

    def setUp(self):
        self._tenant_tmp = tempfile.TemporaryDirectory()
        tenant_root = Path(self._tenant_tmp.name)
        self.addCleanup(self._tenant_tmp.cleanup)
        self._patches = [
            patch.object(tenant_store, "TENANT_ROOT", tenant_root),
            patch.object(connections_registry, "TENANT_ROOT", tenant_root),
            patch.object(jecs_api_server, "TENANT_ROOT", tenant_root),
        ]
        for item in self._patches:
            item.start()
            self.addCleanup(item.stop)

    def test_secret_reference_is_tenant_scoped_and_holds_no_value(self):
        reference = secret_reference("jecs", "amazon", "storage-state")
        self.assertEqual(reference, "secret://jecs/aws-secrets/amazon/storage-state")
        self.assertNotIn("password", reference)

    def test_new_tenant_never_receives_jecs_operational_data(self):
        payload = build_dashboard_operations_payload('funk')
        self.assertEqual(payload['tenant']['id'], 'funk')
        self.assertTrue(payload['needsData'])
        self.assertEqual(payload['performance']['history'], [])
        self.assertEqual(payload['performance']['drivers'], [])
        self.assertEqual(payload['fleet']['vehicles'], [])
        self.assertEqual(payload['fleet']['summary']['registeredFleet'], 0)
        self.assertEqual(payload['costs']['summary'], {})

    def test_new_tenant_rosters_attendance_and_compliance_are_empty(self):
        server = HTTPServer(('127.0.0.1', 0), jecs_api_server.JecsAPIHandler)
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)

        def get(path):
            request = urllib.request.Request(
                f'http://127.0.0.1:{server.server_port}{path}',
                headers={'x-tenant-id': 'funk'},
            )
            with urllib.request.urlopen(request, timeout=5) as response:
                return json.load(response)

        drivers = get('/api/drivers')
        vans = get('/api/vans')
        attendance = get('/api/time-attendance/exceptions')
        compliance = get('/api/fleet-compliance')

        self.assertTrue(drivers['needsData'])
        self.assertEqual(drivers['data'], [])
        self.assertTrue(vans['needsData'])
        self.assertEqual(vans['data'], [])
        self.assertTrue(attendance['needsData'])
        self.assertEqual(attendance['exceptions'], [])
        self.assertTrue(compliance['needsData'])
        self.assertEqual(compliance['vehicles'], [])
        self.assertEqual(compliance['unmatchedPmIssues'], [])
        self.assertEqual(compliance['sources'], [])
        self.assertEqual(compliance['summary']['readinessRate'], 0)

    def test_new_tenant_performance_disputes_routes_and_reimbursements_are_empty(self):
        server = HTTPServer(('127.0.0.1', 0), jecs_api_server.JecsAPIHandler)
        worker = threading.Thread(target=server.serve_forever, daemon=True)
        worker.start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)

        def get(path):
            request = urllib.request.Request(
                f'http://127.0.0.1:{server.server_port}{path}',
                headers={'x-tenant-id': 'funk'},
            )
            with urllib.request.urlopen(request, timeout=5) as response:
                return json.load(response)

        performance = get('/api/performance/dashboard')
        routes = get('/api/route-performance')
        disputes = get('/api/disputes')
        candidates = get('/api/disputes/candidates')
        reimbursements = get('/api/modules')
        module_cases = get('/api/modules/fixed_monthly/cases')

        self.assertTrue(performance['needsData'])
        self.assertEqual(performance['drivers'], [])
        self.assertEqual(performance['dspPerformance']['overallScore'], 0)
        self.assertTrue(routes['needsData'])
        self.assertEqual(routes['routes'], [])
        self.assertEqual(disputes, [])
        self.assertEqual(candidates, [])
        self.assertTrue(reimbursements['needsData'])
        self.assertEqual(reimbursements['modules'], [])
        self.assertEqual(reimbursements['cases'], [])
        self.assertTrue(module_cases['needsData'])
        self.assertEqual(module_cases['cases'], [])

    def test_connection_catalog_exposes_references_not_credential_values(self):
        payload = build_connections_payload("jecs")
        self.assertEqual(payload["summary"]["total"], 6)
        self.assertEqual(payload["summary"]["connectionTotal"], 5)
        self.assertLessEqual(payload["summary"]["connected"], 5)
        amazon = next(c for c in payload["connections"] if c["id"] == "amazon")
        self.assertIn("Fleet Costs", amazon["feeds"])
        self.assertIn("Fleet Compliance", amazon["feeds"])
        self.assertNotIn("PAVE", amazon["description"])
        pave = next(c for c in payload["connections"] if c["id"] == "pave")
        self.assertEqual(pave["authKind"], "browser_session")
        self.assertIn("Wear & Tear", pave["feeds"])

        # A connection may carry a secret *reference*, never a secret value,
        # and must not expose any credential-bearing field at all.
        credential_fields = ("password", "secret", "token", "clientSecret",
                             "storageState", "cookies", "privateKey")
        for connection in payload["connections"]:
            for field in credential_fields:
                self.assertNotIn(field, connection,
                                 f"{connection['id']} exposed credential field {field}")
            reference = connection["secretReference"]
            if reference is not None:
                self.assertTrue(reference.startswith("secret://"))

        # An API connection advertises which secret NAMES it needs, never values.
        digits_api = next(c for c in payload["connections"] if c["id"] == "digits_api")
        self.assertEqual(digits_api["authKind"], "api_credentials")
        self.assertEqual(sorted(digits_api["secretNames"]), ["client-id", "client-secret"])
        for value in digits_api["secretNames"]:
            self.assertNotIn("=", value)
            self.assertLess(len(value), 40)

        financial = next(c for c in payload["connections"] if c["id"] == "financial_charges")
        self.assertEqual(financial["authKind"], "manual_upload")
        self.assertEqual(financial["status"], "not_connected")
        self.assertIn("QuickBooks", financial["acceptedProviders"])

        amazon = next(c for c in payload["connections"] if c["id"] == "amazon")
        pave = next(c for c in payload["connections"] if c["id"] == "pave")
        self.assertEqual(amazon["reconnectLabel"], "Amazon")
        self.assertTrue(amazon["reconnectAvailable"])
        self.assertEqual(pave["reconnectLabel"], "PAVE")
        self.assertTrue(pave["reconnectAvailable"])

    @patch("scripts.connections_registry.has_secret")
    def test_existing_digits_keys_enable_the_write_only_connection_controls(self, has_secret):
        has_secret.side_effect = lambda tenant, integration, name: (
            tenant == "jecs" and integration == "digits"
            and name in {"development-client-id", "development-client-secret"}
        )
        digits = next(c for c in list_connections("jecs") if c["id"] == "digits_api")
        self.assertTrue(digits["configured"])
        self.assertEqual(digits["environment"], "development")
        self.assertEqual(digits["status"], "pending")
        self.assertEqual([f["name"] for f in digits["credentialFields"]],
                         ["clientId", "clientSecret"])
        self.assertTrue(all(f["configured"] for f in digits["credentialFields"]))
        self.assertNotIn("value", str(digits["credentialFields"]))

    def test_local_feature_management_lists_and_persists_every_feature(self):
        features = _public_features("jecs")
        self.assertEqual(len(features), 24)
        self.assertIn("super_admin", {feature["id"] for feature in features})
        self.assertEqual(sum(f["status"] == "planned" for f in features), 4)
        changed = _set_feature_override("jecs", "fuel", False)
        self.assertFalse(changed["enabled"])
        self.assertFalse(next(f for f in _public_features("jecs") if f["id"] == "fuel")["enabled"])

    def test_uploads_are_isolated_between_tenants(self):
        payload = _digits_bytes()
        mine = tenant_store.put_file("jecs", "digits", "d.xlsx", payload, "jason@jeclogs.com")
        theirs = tenant_store.put_file("rival-dsp", "digits", "d.xlsx", payload, "ops@rival.com")
        self.assertTrue(mine["storageKey"].startswith("tenants/jecs/digits/"))
        self.assertTrue(theirs["storageKey"].startswith("tenants/rival-dsp/digits/"))
        self.assertEqual(len(tenant_store.list_files("jecs", "digits")), 1)
        self.assertEqual(len(tenant_store.list_files("rival-dsp", "digits")), 1)

    def test_identical_reupload_is_idempotent(self):
        payload = _digits_bytes()
        first = tenant_store.put_file("jecs", "digits", "d.xlsx", payload, "jason@jeclogs.com")
        second = tenant_store.put_file("jecs", "digits", "d.xlsx", payload, "jason@jeclogs.com")
        self.assertEqual(first["id"], second["id"])
        self.assertEqual(len(tenant_store.list_files("jecs", "digits")), 1)

    def test_parser_reads_net_charge_not_debit(self):
        parsed = parse_digits_export(_digits_bytes(), "d.xlsx")
        self.assertTrue(parsed["ok"])
        self.assertEqual(parsed["columnMapping"]["netCharge"], "Net charge")
        self.assertEqual(parsed["summary"]["included"], 44)
        self.assertEqual(parsed["summary"]["includedTotal"], 79904.99)
        self.assertEqual(parsed["periods"], ["2026-06", "2026-07", "2026-08"])

    def test_unreadable_upload_is_rejected_not_guessed(self):
        parsed = parse_digits_export(b"not,a,digits,file", "junk.csv")
        self.assertFalse(parsed["ok"])
        self.assertEqual(parsed["charges"], [])

    def test_confirmed_upload_drives_fleet_costs_and_matches_workbook(self):
        workbook_view = build_fleet_cost_reconciliation("jecs")
        record = tenant_store.put_file("jecs", "digits", "d.xlsx", _digits_bytes(), "jason@jeclogs.com")
        tenant_store.update_file("jecs", record["id"], status="confirmed", confirmedAt="2026-09-20T21:30:00+00:00")
        uploaded_view = build_fleet_cost_reconciliation("jecs")

        self.assertEqual(uploaded_view["dataSources"][0]["kind"], "tenant_upload")
        self.assertEqual(uploaded_view["summary"]["threeMonthIncludedCost"],
                         workbook_view["summary"]["threeMonthIncludedCost"])
        self.assertEqual(uploaded_view["summary"]["threeMonthDifference"], -6697.56)
        self.assertEqual([m["difference"] for m in uploaded_view["months"]], [332.82, 3877.53, -10907.91])

    def test_other_tenant_never_sees_reference_workbook_data(self):
        view = build_fleet_cost_reconciliation("rival-dsp")
        self.assertTrue(view["needsData"])
        self.assertEqual(view["charges"], [])
        self.assertEqual(view["months"], [])
        self.assertNotIn("79904.99", str(view))


if __name__ == "__main__":
    unittest.main()
