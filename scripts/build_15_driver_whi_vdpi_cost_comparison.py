#!/usr/bin/env python3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data/policies/15-driver-WHI-vs-VDPI-cost-comparison.xlsx"
PDF = ROOT / "data/policies/15-driver-WHI-vs-VDPI-cost-comparison.pdf"

DRIVERS = 15
RATE = 20.50
SHIFTS = 4
FUNDED_HOURS = 10
WEEKS_MONTH = 4.333
OT_THRESHOLD = 40
OT_MULTIPLIER = 1.5
POOL_RATE = 0.90
INDIVIDUAL_CAP = 650
SCENARIOS = [7.5, 8.0, 9.0, 10.0, 10.5, 11.0, 12.0]

NAVY = "183153"
BLUE = "2F6690"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "E9F6EC"
LIGHT_GOLD = "FFF3CD"
LIGHT_RED = "FCE8E6"
WHITE = "FFFFFF"
GRID = "C7D0D9"


def calculations(avg_hours):
    weekly_hours = avg_hours * SHIFTS
    straight_hours = min(weekly_hours, OT_THRESHOLD)
    overtime_hours = max(0, weekly_hours - OT_THRESHOLD)
    actual_wages_driver_week = straight_hours * RATE + overtime_hours * RATE * OT_MULTIPLIER
    whi_topup_driver_week = max(0, FUNDED_HOURS - avg_hours) * SHIFTS * RATE
    whi_group_week = (actual_wages_driver_week + whi_topup_driver_week) * DRIVERS
    reference_whi_group_month = whi_topup_driver_week * DRIVERS * WEEKS_MONTH
    vdpi_bonus_group_month = min(reference_whi_group_month * POOL_RATE, INDIVIDUAL_CAP * DRIVERS)
    vdpi_group_week = actual_wages_driver_week * DRIVERS + vdpi_bonus_group_month / WEEKS_MONTH
    return {
        "weekly_hours": weekly_hours,
        "ot_hours": overtime_hours,
        "whi_topup_driver_week": whi_topup_driver_week,
        "whi_group_week": whi_group_week,
        "vdpi_group_week": vdpi_group_week,
        "vdpi_bonus_group_month": vdpi_bonus_group_month,
        "week_diff": vdpi_group_week - whi_group_week,
        "whi_group_month": whi_group_week * WEEKS_MONTH,
        "vdpi_group_month": vdpi_group_week * WEEKS_MONTH,
        "month_diff": (vdpi_group_week - whi_group_week) * WEEKS_MONTH,
    }


def build_xlsx():
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    thin = Side(style="thin", color=GRID)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=NAVY)
    input_fill = PatternFill("solid", fgColor=LIGHT_GOLD)
    calc_fill = PatternFill("solid", fgColor="F3F5F7")

    ws = wb.active
    ws.title = "Assumptions"
    ws.merge_cells("A1:D1")
    ws["A1"] = "15-Driver WHI vs. VDPI Cost Comparison"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30
    rows = [
        ("Input", "Value", "Purpose"),
        ("Drivers", DRIVERS, "Assumes all listed drivers qualify"),
        ("Hourly rate", RATE, "Current DA base rate"),
        ("Regular shifts per week", SHIFTS, "WHI applies to the first four eligible days"),
        ("Amazon-funded hours per route", FUNDED_HOURS, "Funding boundary, not a cap on worked or paid hours"),
        ("Weeks per month", WEEKS_MONTH, "Monthly-equivalent factor"),
        ("Weekly overtime threshold", OT_THRESHOLD, "Actual hours only"),
        ("Overtime multiplier", OT_MULTIPLIER, "Replace if payroll rules differ"),
        ("VDPI pool funding percentage", POOL_RATE, "Maximum share of Reference WHI Top-Up"),
        ("VDPI monthly individual cap", INDIVIDUAL_CAP, "Per qualifying driver"),
        ("Qualifying drivers", DRIVERS, "Change to model partial qualification"),
    ]
    for r, values in enumerate(rows, 3):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 3:
                cell.fill = header_fill
                cell.font = Font(bold=True, color=WHITE)
            elif c == 2:
                cell.fill = input_fill
    ws["B5"].number_format = '$0.00'
    ws["B8"].number_format = '0.000'
    ws["B10"].number_format = '0.00x'
    ws["B11"].number_format = '0%'
    ws["B12"].number_format = '$#,##0.00'
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 62
    ws.freeze_panes = "A4"

    comp = wb.create_sheet("Cost Comparison")
    comp.merge_cells("A1:M1")
    comp["A1"] = "Weekly Payroll and Monthly Cost — 15 Drivers"
    comp["A1"].font = Font(size=18, bold=True, color=WHITE)
    comp["A1"].fill = header_fill
    comp["A1"].alignment = Alignment(horizontal="center")
    comp.merge_cells("A2:M2")
    comp["A2"] = "WHI tops eligible shifts to 10 hours. Cost-aligned VDPI pays actual wages plus no more than 90% of the Reference WHI Top-Up."
    comp["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    headers = [
        "Avg Hrs/Shift", "Weekly Actual Hrs/DA", "Weekly OT Hrs/DA", "WHI Top-Up/DA/Week",
        "WHI Group Weekly", "VDPI Group Weekly Eq.", "Weekly Difference",
        "WHI Group Monthly", "VDPI Group Monthly Eq.", "Monthly Difference",
        "VDPI Monthly Payout", "WHI Cost Retained", "Cost Result",
    ]
    for c, h in enumerate(headers, 1):
        cell = comp.cell(4, c, h)
        cell.fill = header_fill
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for idx, hours in enumerate(SCENARIOS, 5):
        comp.cell(idx, 1, hours).fill = input_fill
        comp.cell(idx, 2, f"=A{idx}*Assumptions!$B$6")
        comp.cell(idx, 3, f"=MAX(0,B{idx}-Assumptions!$B$9)")
        comp.cell(idx, 4, f"=MAX(0,Assumptions!$B$7-A{idx})*Assumptions!$B$6*Assumptions!$B$5")
        actual = f"(MIN(B{idx},Assumptions!$B$9)*Assumptions!$B$5+MAX(0,B{idx}-Assumptions!$B$9)*Assumptions!$B$5*Assumptions!$B$10)"
        comp.cell(idx, 5, f"=({actual}+D{idx})*Assumptions!$B$4")
        vdpi_pool = f"MIN(D{idx}*Assumptions!$B$4*Assumptions!$B$8*Assumptions!$B$11,Assumptions!$B$12*Assumptions!$B$13)"
        comp.cell(idx, 6, f"={actual}*Assumptions!$B$4+({vdpi_pool}/Assumptions!$B$8)")
        comp.cell(idx, 7, f"=F{idx}-E{idx}")
        comp.cell(idx, 8, f"=E{idx}*Assumptions!$B$8")
        comp.cell(idx, 9, f"=F{idx}*Assumptions!$B$8")
        comp.cell(idx, 10, f"=I{idx}-H{idx}")
        comp.cell(idx, 11, f"={vdpi_pool}")
        comp.cell(idx, 12, f"=MAX(0,D{idx}*Assumptions!$B$4*Assumptions!$B$8-K{idx})")
        comp.cell(idx, 13, f'=IF(J{idx}>0,"ERROR: above WHI",IF(J{idx}<0,"PASS: below WHI","PASS: equal"))')
        for c in range(1, 14):
            cell = comp.cell(idx, c)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if c != 1:
                cell.fill = calc_fill
        for c in range(4, 13):
            comp.cell(idx, c).number_format = '$#,##0.00;[Red]-$#,##0.00'
    for c, width in enumerate([14, 18, 16, 19, 19, 21, 18, 20, 22, 19, 20, 22, 18], 1):
        comp.column_dimensions[get_column_letter(c)].width = width
    comp.freeze_panes = "A5"
    comp.auto_filter.ref = f"A4:M{4+len(SCENARIOS)}"

    timing = wb.create_sheet("Payout Timing")
    timing.merge_cells("A1:D1")
    timing["A1"] = "Cost-Aligned VDPI Funding Rules"
    timing["A1"].font = Font(size=18, bold=True, color=WHITE)
    timing["A1"].fill = header_fill
    timing["A1"].alignment = Alignment(horizontal="center")
    payout_rows = [
        ("Rule", "Calculation", "Result", "Notes"),
        ("Reference WHI Top-Up", "Sum of otherwise eligible top-up hours to 10 × base rate", "Variable", "Calculated after monthly payroll closes"),
        ("Maximum VDPI pool", "Lesser of 90% of Reference WHI Top-Up or $650 × qualifying DAs", "Variable", "Never exceeds WHI; includes a 10% cost reserve"),
        ("Payment timing", "One monthly payment after certification", "No quarterly bonus", "Any holiday/rescue cash counts against the same pool unless separately funded"),
        ("10+ hour shifts", "Reference WHI Top-Up = $0", "$0 VDPI pool", "Actual hours and overtime are still paid normally"),
    ]
    for r, values in enumerate(payout_rows, 3):
        for c, value in enumerate(values, 1):
            cell = timing.cell(r, c, value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 3:
                cell.fill = header_fill
                cell.font = Font(bold=True, color=WHITE)
            elif c == 3 and isinstance(value, (int, float)):
                cell.fill = calc_fill
                cell.number_format = '$#,##0.00'
    for c, width in enumerate([28, 42, 20, 55], 1):
        timing.column_dimensions[get_column_letter(c)].width = width

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddFooter.center.text = "WHI vs. VDPI Cost Comparison"
        sheet.oddFooter.right.text = "Page &P of &N"

    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)


def para(text, style):
    return Paragraph(text, style)


def build_pdf():
    width, height = landscape(letter)
    c = Canvas(str(PDF), pagesize=(width, height))
    c.setTitle("15-Driver WHI vs. VDPI Cost Comparison")
    c.setAuthor("JEC Logistics Solutions LLC")
    navy = colors.HexColor("#183153")
    blue = colors.HexColor("#2F6690")
    gray = colors.HexColor("#5E6872")
    grid = colors.HexColor("#C7D0D9")
    pale_blue = colors.HexColor("#EAF2F8")
    pale_gold = colors.HexColor("#FFF3CD")
    pale_red = colors.HexColor("#FCE8E6")
    pale_green = colors.HexColor("#E9F6EC")
    title = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=17, leading=20, textColor=navy, alignment=TA_CENTER)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.2, leading=10, textColor=gray, alignment=TA_CENTER)
    section = ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=10.2, leading=12, textColor=navy)
    head = ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=7.2, leading=8.3, textColor=colors.white, alignment=TA_CENTER)
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=7.2, leading=8.5, textColor=colors.HexColor("#222222"), alignment=TA_CENTER)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#222222"))
    callout = ParagraphStyle("callout", fontName="Helvetica-Bold", fontSize=8.2, leading=10, textColor=navy, alignment=TA_CENTER)
    fine = ParagraphStyle("fine", fontName="Helvetica", fontSize=6.8, leading=8, textColor=gray, alignment=TA_CENTER)

    c.setFillColor(navy)
    c.rect(0, height - 0.38 * inch, width, 0.38 * inch, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 8)
    c.drawString(0.4 * inch, height - 0.24 * inch, "JEC LOGISTICS SOLUTIONS LLC")
    c.drawRightString(width - 0.4 * inch, height - 0.24 * inch, "OWNER COST COMPARISON • DRAFT")

    p = para("WHI vs. Proposed VDPI — Cost for 15 Drivers", title)
    _, ph = p.wrap(width - 0.8 * inch, 0.35 * inch)
    p.drawOn(c, 0.4 * inch, height - 0.84 * inch)
    p = para("$20.50/hour • 4 eligible shifts/week • 4.333 weeks/month • 90% WHI cost cap • $650 monthly individual cap", sub)
    p.wrap(width - 0.8 * inch, 0.2 * inch)
    p.drawOn(c, 0.4 * inch, height - 1.02 * inch)

    x = 0.42 * inch
    y_top = height - 1.18 * inch
    table_rows = [[
        para("Avg.<br/>hours/shift", head), para("Weekly actual<br/>hours/DA", head), para("OT hours<br/>per DA", head),
        para("WHI group<br/>weekly cost", head), para("VDPI group<br/>weekly equivalent", head), para("Weekly<br/>difference", head),
        para("WHI group<br/>monthly cost", head), para("VDPI group<br/>monthly equivalent", head), para("Monthly<br/>difference", head),
    ]]
    for hours in SCENARIOS:
        d = calculations(hours)
        table_rows.append([
            para(f"{hours:.1f}", cell), para(f"{d['weekly_hours']:.1f}", cell), para(f"{d['ot_hours']:.1f}", cell),
            para(f"${d['whi_group_week']:,.0f}", cell), para(f"${d['vdpi_group_week']:,.0f}", cell), para(f"${d['week_diff']:,.0f}", cell),
            para(f"${d['whi_group_month']:,.0f}", cell), para(f"${d['vdpi_group_month']:,.0f}", cell), para(f"${d['month_diff']:,.0f}", cell),
        ])
    table = Table(table_rows, colWidths=[0.78*inch,0.96*inch,0.72*inch,1.12*inch,1.22*inch,0.90*inch,1.18*inch,1.28*inch,1.05*inch], repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0,0), (-1,0), navy), ("GRID", (0,0), (-1,-1), 0.45, grid),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F3F5F7")]),
        ("TOPPADDING", (0,0), (-1,-1), 3.2), ("BOTTOMPADDING", (0,0), (-1,-1), 3.2),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("BACKGROUND", (0,2), (-1,2), pale_gold), ("BACKGROUND", (0,4), (-1,4), pale_red),
    ]
    table.setStyle(TableStyle(style_cmds))
    tw, th = table.wrap(width - 0.84 * inch, 3.6 * inch)
    table.drawOn(c, x, y_top - th)

    box_y = y_top - th - 1.20 * inch
    gap = 0.15 * inch
    box_w = (width - 0.84 * inch - gap * 2) / 3
    boxes = [
        (pale_blue, "One monthly variable payout", "Pool = lesser of <b>90% of Reference WHI Top-Up</b> or <b>$650 × qualifying DAs</b>. There is no quarterly cash bonus."),
        (pale_gold, "What happens above 10 hours?", "WHI top-up becomes <b>$0</b>. Both models still pay all actual wages and overtime. The cost-aligned VDPI pool is also <b>$0</b>."),
        (pale_red, "No stacking", "Holiday, rescue, attendance, and retention cash must count against the same monthly pool unless a written amendment identifies separate outside funding."),
    ]
    for i, (fill, heading, text) in enumerate(boxes):
        bx = x + i * (box_w + gap)
        c.setFillColor(fill); c.setStrokeColor(grid)
        c.roundRect(bx, box_y, box_w, 0.95 * inch, 7, fill=1, stroke=1)
        q = para(heading, section); q.wrap(box_w - 0.24*inch, 0.22*inch); q.drawOn(c, bx+0.12*inch, box_y+0.68*inch)
        q = para(text, body); q.wrap(box_w - 0.24*inch, 0.55*inch); q.drawOn(c, bx+0.12*inch, box_y+0.12*inch)

    # Detailed over-10-hours example fills the lower half and makes the payroll mechanics explicit.
    peak = calculations(10.5)
    detail_y = 1.28 * inch
    detail_h = 2.05 * inch
    c.setFillColor(pale_blue); c.setStrokeColor(grid)
    c.roundRect(x, detail_y, width - 0.84*inch, detail_h, 7, fill=1, stroke=1)
    q = para("10.5-Hour Peak Example — 15 Drivers", section)
    q.wrap(width - 1.12*inch, 0.24*inch); q.drawOn(c, x+0.14*inch, detail_y+detail_h-0.30*inch)
    q = para("Each DA works 42 actual hours: 40 regular hours plus 2 overtime hours. WHI provides no top-up because every shift exceeded 10 hours.", body)
    q.wrap(width - 1.12*inch, 0.25*inch); q.drawOn(c, x+0.14*inch, detail_y+detail_h-0.56*inch)
    detail_rows = [[
        para("Cost item", head), para("One DA / week", head), para("15 DAs / weekly payroll", head),
        para("15 DAs / monthly equivalent", head), para("Explanation", head),
    ], [
        para("Actual wages + OT", cell), para("$881.50", cell), para(f"${peak['whi_group_week']:,.2f}", cell),
        para(f"${peak['whi_group_month']:,.2f}", cell), para("Same under both models", cell),
    ], [
        para("WHI top-up", cell), para("$0.00", cell), para("$0.00", cell), para("$0.00", cell),
        para("No shift was below 10 hours", cell),
    ], [
        para("VDPI bonus", cell), para("$0.00", cell),
        para(f"${peak['week_diff']:,.2f}", cell), para(f"${peak['month_diff']:,.2f}", cell),
        para("Additional to actual wages", cell),
    ], [
        para("VDPI total", cell), para("$881.50", cell), para(f"${peak['vdpi_group_week']:,.2f}", cell),
        para(f"${peak['vdpi_group_month']:,.2f}", cell), para("Actual wages + bonus equivalent", cell),
    ]]
    detail_table = Table(detail_rows, colWidths=[1.25*inch,1.18*inch,1.60*inch,1.70*inch,3.48*inch], repeatRows=1)
    detail_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), blue), ("GRID", (0,0), (-1,-1), 0.45, grid),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F3F5F7")]),
        ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("BACKGROUND", (0,3), (-1,3), pale_gold),
    ]))
    dw, dh = detail_table.wrap(width - 1.12*inch, 1.25*inch)
    detail_table.drawOn(c, x+0.14*inch, detail_y+0.13*inch)

    call_y = 0.55 * inch
    c.setFillColor(pale_green); c.setStrokeColor(colors.HexColor("#87B891"))
    c.roundRect(x, call_y, width - 0.84*inch, 0.55*inch, 7, fill=1, stroke=1)
    q = para("COST CONCLUSION: The revised VDPI cannot exceed WHI. It uses no more than 90% of the Reference WHI Top-Up, is capped at $650 per qualifying DA, has no quarterly bonus, and pays $0 when all eligible shifts are 10 hours or longer.", callout)
    _, qh = q.wrap(width - 1.12*inch, 0.42*inch)
    q.drawOn(c, x+0.14*inch, call_y+0.27*inch-qh/2)
    q = para("Planning comparison only. Excludes employer payroll burden and any additional regular-rate/overtime adjustment attributable to nondiscretionary bonuses. All actual hours must be recorded and paid.", fine)
    q.wrap(width - 0.9*inch, 0.15*inch); q.drawOn(c, 0.45*inch, 0.25*inch)
    c.showPage(); c.save()


def main():
    build_xlsx()
    build_pdf()
    print(XLSX)
    print(PDF)


if __name__ == "__main__":
    main()
