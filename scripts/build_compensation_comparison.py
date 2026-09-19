from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = [
    ROOT / "data" / "compensation-model-comparison.xlsx",
    Path("/Users/claw/Downloads/veteran-driver-compensation-model-comparison.xlsx"),
]

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
GOLD = "FFC000"
LIGHT_GOLD = "FFF2CC"
RED = "C00000"
LIGHT_RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_FILL = PatternFill("solid", fgColor=LIGHT_GOLD)
FORMULA_FILL = PatternFill("solid", fgColor="F2F2F2")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FILL = PatternFill("solid", fgColor=BLUE)
THIN = Side(style="thin", color="B7B7B7")


def title(ws, text, subtitle=None):
    ws.merge_cells("A1:H1")
    ws["A1"] = text
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells("A2:H2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(italic=True, color="555555")
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[2].height = 30


def section(ws, row, text, end_col=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    c = ws.cell(row, 1, text)
    c.fill = SECTION_FILL
    c.font = Font(bold=True, color=WHITE)


def money(cell):
    cell.number_format = '$#,##0.00;[Red]-$#,##0.00'


def percent(cell):
    cell.number_format = '0.0%'


def build():
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    readme = wb.active
    readme.title = "Read Me"
    title(readme, "Veteran Driver Compensation Comparison", "Editable comparison of a 10-hour guarantee and a capped Veteran Performance Pool.")
    section(readme, 4, "How to use this workbook")
    instructions = [
        "1. Open Assumptions and replace the yellow cells with your payroll assumptions.",
        "2. Open Driver Inputs and enter eligible drivers, pay rates, shifts, and expected hours under both models.",
        "3. Use Model B Avg Hours/Shift to test whether drivers stretch their shifts.",
        "4. Enter preventable-damage estimates only if you want those costs included in the comparison.",
        "5. Review Summary for total employer cost, savings, pool payout, and the cost-cap test.",
        "6. Review Stretch Test to see how the performance pool responds as paid hours rise.",
    ]
    for i, text in enumerate(instructions, 5):
        readme.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        readme.cell(i, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[i].height = 24
    section(readme, 12, "Model logic")
    logic = [
        ("Model A", "Actual wages plus enough guaranteed-hour pay to reach the selected hours per completed shift."),
        ("Model B", "Actual wages plus a fixed performance award. The total award is capped and reduced dollar-for-dollar by unexplained labor-hour inflation and the selected damage deduction."),
        ("Guardrail", "The workbook flags whether Model B is above or below Model A and whether the pool stays within its preset budget."),
        ("Safety", "The efficiency band is a review flag—not permission to withhold earned wages. Document route size, rescues, late dispatch, weather, traffic, vehicle issues, and station delays."),
        ("Payroll", "The payroll-burden and bonus/overtime-uplift fields are editable. Confirm the final policy and regular-rate treatment with payroll/employment counsel."),
    ]
    for r, (label, text) in enumerate(logic, 13):
        readme.cell(r, 1, label).font = Font(bold=True)
        readme.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        readme.cell(r, 2, text).alignment = Alignment(wrap_text=True)
        readme.row_dimensions[r].height = 32
    readme.column_dimensions["A"].width = 18
    for col in "BCDEFGH":
        readme.column_dimensions[col].width = 15
    readme.freeze_panes = "A4"

    assumptions = wb.create_sheet("Assumptions")
    title(assumptions, "Assumptions", "Yellow cells are user inputs. Defaults are examples only.")
    section(assumptions, 4, "Program assumptions")
    assumption_rows = [
        (5, "Weeks in analysis period", 4.333, "Average weeks per month"),
        (6, "Guarantee hours per completed shift", 10.0, "Model A"),
        (7, "Pool budget as % of historical guarantee equivalent", 0.70, "Model B ceiling"),
        (8, "Target performance award per driver per month", 250.0, "Before pool cap and deductions"),
        (9, "Quarterly holdback % of earned award", 0.25, "Part of the same award—not extra"),
        (10, "Payroll burden %", 0.12, "Taxes/insurance estimate; replace with actual"),
        (11, "Bonus overtime/regular-rate uplift %", 0.00, "Replace with payroll estimate"),
        (12, "Pool deduction for preventable damage %", 1.00, "Use only after preventability review"),
        (13, "Safe efficiency lower bound vs baseline", 0.85, "Too low triggers fast-route review"),
        (14, "Safe efficiency upper bound vs baseline", 1.05, "Above triggers slow-route review"),
    ]
    for row, label, value, note in assumption_rows:
        assumptions.cell(row, 1, label)
        assumptions.cell(row, 2, value).fill = INPUT_FILL
        assumptions.cell(row, 3, note)
        assumptions.cell(row, 3).alignment = Alignment(wrap_text=True)
    for row in (7, 9, 10, 11, 12, 13, 14):
        percent(assumptions.cell(row, 2))
    money(assumptions["B8"])
    assumptions["B5"].number_format = "0.000"
    assumptions["B6"].number_format = "0.00"
    section(assumptions, 16, "Recommended cost rule")
    assumptions.merge_cells("A17:H17")
    assumptions["A17"] = "Actual Pool = MIN(base pool budget, eligible target awards) − unexplained excess labor cost − selected preventable-damage deduction"
    assumptions["A17"].alignment = Alignment(wrap_text=True)
    assumptions.row_dimensions[17].height = 32
    assumptions.merge_cells("A18:H19")
    assumptions["A18"] = "This model never changes the obligation to pay for all hours worked. The pool deduction is a prospective bonus calculation and should be reviewed for payroll and employment-law compliance before adoption."
    assumptions["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    assumptions.column_dimensions["A"].width = 58
    assumptions.column_dimensions["B"].width = 18
    assumptions.column_dimensions["C"].width = 45
    assumptions.freeze_panes = "A5"

    ws = wb.create_sheet("Driver Inputs")
    ws["A1"] = "Driver-Level Inputs and Calculations"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = HEADER_FILL
    ws.merge_cells("A1:AG1")
    ws["A2"] = "Enter or replace yellow input cells. Gray cells calculate automatically. Model B hours are the main stretch-risk input."
    ws.merge_cells("A2:AG2")
    ws["A2"].alignment = Alignment(wrap_text=True)
    headers = [
        "Include?", "Driver", "Hourly Rate", "Shifts/Week", "Baseline Avg Hrs/Shift",
        "Model A Avg Hrs/Shift", "Model B Avg Hrs/Shift", "Bonus Eligible?",
        "Model A Damage Cost", "Model B Damage Cost", "Period Shifts", "Baseline Wages",
        "A Worked Hours", "A Guarantee Top-Up Hrs", "A Wage + Guarantee", "A Payroll Burden",
        "MODEL A TOTAL", "B Worked Hours", "B Excess Hrs vs Baseline", "B Excess Labor Cost",
        "B Target Award", "B Allocated Award", "B Quarterly Holdback", "B Current Award",
        "B Actual Wages", "B Bonus OT Uplift", "B Payroll Burden", "MODEL B TOTAL",
        "B − A Difference", "B vs A %", "Efficiency Status", "Pool Reduction", "Notes"
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(4, col, h)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color=WHITE)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = Border(bottom=THIN, right=THIN)
    ws.row_dimensions[4].height = 58
    yes_no = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(yes_no)
    yes_no.add("A5:A34")
    yes_no.add("H5:H34")
    sample_hours = [8.75, 8.60, 9.10, 8.25, 9.00, 8.80, 8.50, 9.20, 8.70, 8.95]
    for row in range(5, 35):
        idx = row - 4
        active = idx <= 10
        defaults = {
            1: "Y" if active else "N",
            2: f"Driver {idx}" if active else "",
            3: 22.00 if active else None,
            4: 4.0 if active else None,
            5: sample_hours[idx - 1] if active else None,
            6: sample_hours[idx - 1] - 0.10 if active else None,
            7: sample_hours[idx - 1] + 0.35 if active else None,
            8: "Y" if active else "N",
            9: 0.0 if active else None,
            10: 0.0 if active else None,
        }
        for col in range(1, 11):
            ws.cell(row, col, defaults.get(col)).fill = INPUT_FILL
        ws.cell(row, 11, f'=IF(A{row}="Y",D{row}*Assumptions!$B$5,0)')
        ws.cell(row, 12, f'=IF(A{row}="Y",C{row}*D{row}*Assumptions!$B$5*E{row},0)')
        ws.cell(row, 13, f'=IF(A{row}="Y",D{row}*Assumptions!$B$5*F{row},0)')
        ws.cell(row, 14, f'=IF(A{row}="Y",MAX(0,K{row}*Assumptions!$B$6-M{row}),0)')
        ws.cell(row, 15, f'=IF(A{row}="Y",C{row}*(M{row}+N{row}),0)')
        ws.cell(row, 16, f'=O{row}*Assumptions!$B$10')
        ws.cell(row, 17, f'=O{row}+P{row}+IF(A{row}="Y",I{row},0)')
        ws.cell(row, 18, f'=IF(A{row}="Y",D{row}*Assumptions!$B$5*G{row},0)')
        ws.cell(row, 19, f'=IF(A{row}="Y",MAX(0,R{row}-(D{row}*Assumptions!$B$5*E{row})),0)')
        ws.cell(row, 20, f'=S{row}*C{row}')
        ws.cell(row, 21, f'=IF(AND(A{row}="Y",H{row}="Y"),Assumptions!$B$8*Assumptions!$B$5/4.333,0)')
        ws.cell(row, 22, f'=IFERROR(U{row}*Summary!$B$18,0)')
        ws.cell(row, 23, f'=V{row}*Assumptions!$B$9')
        ws.cell(row, 24, f'=V{row}-W{row}')
        ws.cell(row, 25, f'=R{row}*C{row}')
        ws.cell(row, 26, f'=V{row}*Assumptions!$B$11')
        ws.cell(row, 27, f'=(Y{row}+V{row}+Z{row})*Assumptions!$B$10')
        ws.cell(row, 28, f'=Y{row}+V{row}+Z{row}+AA{row}+IF(A{row}="Y",J{row},0)')
        ws.cell(row, 29, f'=AB{row}-Q{row}')
        ws.cell(row, 30, f'=IFERROR(AB{row}/Q{row}-1,0)')
        ws.cell(row, 31, f'=IF(A{row}<>"Y","Excluded",IF(G{row}<E{row}*Assumptions!$B$13,"Review: unusually fast",IF(G{row}>E{row}*Assumptions!$B$14,"Review: unexplained slow","Within safe band")))')
        ws.cell(row, 32, f'=IFERROR(T{row}+J{row}*Assumptions!$B$12,0)')
        for col in range(11, 33):
            ws.cell(row, col).fill = FORMULA_FILL
        for col in (3, 9, 10, 12, 15, 16, 17, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 32):
            money(ws.cell(row, col))
        percent(ws.cell(row, 30))
        for col in (4, 5, 6, 7, 11, 13, 14, 18, 19):
            ws.cell(row, col).number_format = "0.00"
    ws.auto_filter.ref = "A4:AG34"
    ws.freeze_panes = "K5"
    widths = [10, 20, 12, 12, 16, 17, 17, 14, 16, 16, 13, 14, 13, 17, 16, 14, 16, 13, 16, 15, 14, 15, 16, 15, 14, 14, 14, 16, 15, 12, 24, 15, 26]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.conditional_formatting.add("AC5:AC34", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    ws.conditional_formatting.add("AC5:AC34", CellIsRule(operator="lessThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
    ws.conditional_formatting.add("AE5:AE34", FormulaRule(formula=['ISNUMBER(SEARCH("Review",AE5))'], fill=PatternFill("solid", fgColor=LIGHT_GOLD)))

    summary = wb.create_sheet("Summary")
    title(summary, "Cost Comparison Summary", "Calculated from Driver Inputs for the selected analysis period.")
    section(summary, 4, "Core comparison")
    rows = [
        (5, "Included drivers", '=COUNTIF(\'Driver Inputs\'!A5:A34,"Y")'),
        (6, "Analysis weeks", '=Assumptions!B5'),
        (7, "Historical guarantee equivalent", '=SUMPRODUCT((\'Driver Inputs\'!A5:A34="Y")*(\'Driver Inputs\'!C5:C34)*\'Driver Inputs\'!K5:K34*IF(Assumptions!B6>\'Driver Inputs\'!E5:E34,Assumptions!B6-\'Driver Inputs\'!E5:E34,0))'),
        (8, "Base performance-pool budget", '=B7*Assumptions!B7'),
        (9, "Eligible target awards", '=SUM(\'Driver Inputs\'!U5:U34)'),
        (10, "Unexplained excess labor cost", '=SUM(\'Driver Inputs\'!T5:T34)'),
        (11, "Selected damage deduction", '=SUM(\'Driver Inputs\'!J5:J34)*Assumptions!B12'),
        (12, "Actual performance-pool payout", '=MAX(0,MIN(B8,B9)-B10-B11)'),
        (13, "Current-period award paid now", '=B12*(1-Assumptions!B9)'),
        (14, "Quarterly holdback accrued", '=B12*Assumptions!B9'),
        (15, "Model A total employer cost", '=SUM(\'Driver Inputs\'!Q5:Q34)'),
        (16, "Model B total employer cost", '=SUM(\'Driver Inputs\'!AB5:AB34)'),
        (17, "Model B minus Model A", '=B16-B15'),
        (18, "Award allocation ratio", '=IFERROR(B12/B9,0)'),
        (19, "Model B savings %", '=IFERROR((B15-B16)/B15,0)'),
        (20, "Cost-cap status", '=IF(B16<=B15,"PASS: Model B is at or below Model A","REVIEW: Model B exceeds Model A")'),
    ]
    for row, label, formula in rows:
        summary.cell(row, 1, label)
        summary.cell(row, 2, formula).fill = FORMULA_FILL
    for row in range(7, 18):
        money(summary.cell(row, 2))
    summary["B18"].number_format = "0.0%"
    summary["B19"].number_format = "0.0%"
    summary["B20"].font = Font(bold=True)
    summary.conditional_formatting.add("B20", FormulaRule(formula=['LEFT(B20,4)="PASS"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
    summary.conditional_formatting.add("B20", FormulaRule(formula=['LEFT(B20,6)="REVIEW"'], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    section(summary, 22, "Interpretation")
    summary.merge_cells("A23:H23")
    summary["A23"] = "If Model B hours rise, excess labor cost reduces the bonus pool. If hours increase beyond the available pool, Model B can still exceed Model A; the red cost-cap status makes that visible."
    summary["A23"].alignment = Alignment(wrap_text=True)
    summary.row_dimensions[23].height = 34
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 25
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Total Employer Cost"
    chart.y_axis.title = "Cost"
    data = Reference(summary, min_col=2, min_row=15, max_row=16)
    cats = Reference(summary, min_col=1, min_row=15, max_row=16)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 13
    summary.add_chart(chart, "D5")

    stretch = wb.create_sheet("Stretch Test")
    title(stretch, "Model B Stretch Sensitivity", "Shows the aggregate effect if every included driver works progressively longer than their frozen baseline.")
    stretch_headers = ["Added Hours/Shift", "Projected Pool Wages", "Excess Labor Cost", "Projected Pool Payout", "Projected Model B Total", "Model A Total", "B − A", "Status"]
    for col, h in enumerate(stretch_headers, 1):
        c = stretch.cell(4, col, h)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color=WHITE)
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    increments = [0, .25, .5, .75, 1.0, 1.25, 1.5, 2.0]
    for row, inc in enumerate(increments, 5):
        stretch.cell(row, 1, inc).fill = INPUT_FILL
        stretch.cell(row, 2, f'=SUMPRODUCT((\'Driver Inputs\'!$A$5:$A$34="Y")*\'Driver Inputs\'!$C$5:$C$34*\'Driver Inputs\'!$D$5:$D$34*Assumptions!$B$5*(\'Driver Inputs\'!$E$5:$E$34+A{row}))')
        stretch.cell(row, 3, f'=SUMPRODUCT((\'Driver Inputs\'!$A$5:$A$34="Y")*\'Driver Inputs\'!$C$5:$C$34*\'Driver Inputs\'!$D$5:$D$34*Assumptions!$B$5*A{row})')
        stretch.cell(row, 4, f'=MAX(0,MIN(Summary!$B$8,Summary!$B$9)-C{row}-Summary!$B$11)')
        stretch.cell(row, 5, f'=(B{row}+D{row}*(1+Assumptions!$B$11))*(1+Assumptions!$B$10)+SUM(\'Driver Inputs\'!$J$5:$J$34)')
        stretch.cell(row, 6, '=Summary!$B$15')
        stretch.cell(row, 7, f'=E{row}-F{row}')
        stretch.cell(row, 8, f'=IF(G{row}<=0,"At/below Model A","Above Model A")')
        for col in range(2, 8):
            money(stretch.cell(row, col))
    stretch.conditional_formatting.add("G5:G12", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    stretch.conditional_formatting.add("G5:G12", CellIsRule(operator="lessThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))
    for col, width in enumerate([18, 22, 20, 22, 22, 18, 16, 20], 1):
        stretch.column_dimensions[get_column_letter(col)].width = width
    stretch.freeze_panes = "A5"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddFooter.center.text = "Veteran Driver Compensation Comparison"
        sheet.oddFooter.right.text = "Page &P of &N"

    for path in OUTPUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)


if __name__ == "__main__":
    build()
