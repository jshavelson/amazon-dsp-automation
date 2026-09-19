#!/usr/bin/env python3
"""Shared evidence and case contracts for independently billable reconciliation modules."""
from __future__ import annotations

import csv
import hashlib
import json
import os
import tempfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


SCHEMA_VERSION = 1


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def evidence(path: Path, evidence_type: str) -> dict[str, Any]:
    resolved = path.resolve(strict=True)
    return {
        "type": evidence_type,
        "path": str(resolved),
        "sha256": sha256(resolved),
        "bytes": resolved.stat().st_size,
    }


def money(value: Any, default: Decimal | None = None) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return default
    try:
        return Decimal(str(value).replace("$", "").replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid monetary value: {value!r}") from exc


def parse_date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date: {value!r}")


def normalize_key(value: Any) -> str:
    return "".join(ch.lower() for ch in str(value or "") if ch.isalnum())


def pick(row: dict[str, Any], aliases: Iterable[str], *, required: bool = False) -> Any:
    indexed = {normalize_key(key): value for key, value in row.items()}
    for alias in aliases:
        key = normalize_key(alias)
        if key in indexed and str(indexed[key] or "").strip() != "":
            return indexed[key]
    if required:
        raise ValueError(f"Missing required field; accepted names: {', '.join(aliases)}")
    return None


def load_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    if suffix == ".json":
        document = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(document, dict):
            for key in ("rows", "data", "items", "records"):
                if isinstance(document.get(key), list):
                    document = document[key]
                    break
        if not isinstance(document, list) or not all(isinstance(row, dict) for row in document):
            raise ValueError(f"{path} must contain a list of objects")
        return document
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        return [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
    raise ValueError(f"Unsupported evidence format: {path.suffix}; use CSV, JSON, or XLSX")


def case_document(
    *,
    module_id: str,
    external_key: str,
    findings: list[dict[str, Any]],
    evidence_items: list[dict[str, Any]],
    blocking_evidence: list[str] | None = None,
    tenant_id: str = "jecs",
) -> dict[str, Any]:
    candidates = [finding for finding in findings if finding.get("dispute_candidate")]
    blockers = list(blocking_evidence or [])
    if blockers:
        status, disposition = "needs_evidence", "evidence_incomplete"
    elif candidates:
        status, disposition = "approval_required", "potential_underpayment"
    else:
        status, disposition = "closed", "reconciled"
    return {
        "schema_version": SCHEMA_VERSION,
        "module_id": module_id,
        "tenant_id": tenant_id,
        "external_key": external_key,
        "status": status,
        "disposition": disposition,
        "created_at": utc_now(),
        "evidence": evidence_items,
        "findings": findings,
        "candidate_count": len(candidates),
        "blocking_evidence": blockers,
        "external_action_authorized": False,
        "approval": None,
        "submission": None,
    }


def json_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(item) for item in value]
    return value


def write_json_atomic(path: Path, document: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    safe_document = json_safe(document)
    if path.exists() and isinstance(safe_document, dict) and "created_at" in safe_document:
        try:
            existing = json.loads(path.read_text(encoding="utf-8"))
            old_compare = {key: value for key, value in existing.items() if key != "created_at"}
            new_compare = {key: value for key, value in safe_document.items() if key != "created_at"}
            if old_compare == new_compare:
                safe_document["created_at"] = existing.get("created_at", safe_document["created_at"])
        except (OSError, json.JSONDecodeError):
            pass
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(safe_document, handle, indent=2)
            handle.write("\n")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
