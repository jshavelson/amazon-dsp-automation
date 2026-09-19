#!/usr/bin/env python3
"""Reconcile Amazon Capacity & Reliability with route-coded WST executions."""
from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from reconciliation_core import case_document, evidence, write_json_atomic


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def number(value: Any) -> int:
    if value in (None, "", "-"):
        return 0
    return int(float(str(value).replace("%", "").replace(",", "")))


def amazon_reporting_week(day: date) -> tuple[int, int]:
    """Amazon operating weeks run Sunday-Saturday; ISO weeks run Monday-Sunday."""
    shifted = day + timedelta(days=1)
    iso = shifted.isocalendar()
    return iso.year, iso.week


def amazon_rows(path: Path, year: int, week: int) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    matches: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for values in sheet.iter_rows(values_only=True):
            if len(values) < 14:
                continue
            day = as_date(values[0])
            if not day or amazon_reporting_week(day) != (year, week):
                continue
            matches.append({
                "date": day.isoformat(),
                "final_scheduled": number(values[4]),
                "completed_routes": number(values[5]),
                "amazon_paid_cancels": number(values[7]),
                "reliability_target": number(values[9]),
                "dsp_dropped_routes": number(values[11]),
            })
    if not matches:
        raise ValueError(f"No Amazon C&R daily rows found for {year}-W{week:02d}")
    return sorted(matches, key=lambda row: row["date"])


def worksheet_list(document: Any) -> list[dict[str, Any]]:
    if isinstance(document, list):
        return [row for row in document if isinstance(row, dict)]
    if isinstance(document, dict):
        for key in ("worksheets", "data", "workOrders"):
            if isinstance(document.get(key), list):
                return [row for row in document[key] if isinstance(row, dict)]
    return []


def wst_rows(directory: Path, year: int, week: int) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    daily: dict[str, dict[str, Any]] = {}
    evidence_items: list[dict[str, Any]] = []
    for path in sorted(directory.glob("*-worksheets.json")):
        match = re.match(r"(\d{4}-\d{2}-\d{2})", path.name)
        if not match:
            continue
        day = date.fromisoformat(match.group(1))
        if amazon_reporting_week(day) != (year, week):
            continue
        rows = worksheet_list(json.loads(path.read_text(encoding="utf-8")))
        eligible = [
            row for row in rows
            if not bool(row.get("suppressed"))
            and str(row.get("status") or "").upper() == "SUBMITTED"
            and bool(row.get("routeCode"))
        ]
        excluded = [row for row in rows if row not in eligible]
        daily[day.isoformat()] = {
            "wst_completed_routes": len(eligible),
            "excluded_work_orders": len(excluded),
            "service_types": sorted({str(row.get("serviceTypeName") or "Unlabeled") for row in eligible}),
        }
        evidence_items.append(evidence(path, "wst_daily_worksheets"))
    return daily, evidence_items


def reconcile(capacity_path: Path, wst_directory: Path, year: int, week: int) -> dict[str, Any]:
    amazon = amazon_rows(capacity_path, year, week)
    wst, wst_evidence = wst_rows(wst_directory, year, week)
    findings: list[dict[str, Any]] = []
    blockers: list[str] = []
    for row in amazon:
        local = wst.get(row["date"])
        if local is None:
            blockers.append(f"WST worksheet export for {row['date']}")
            findings.append({"date": row["date"], "status": "missing_evidence", "dispute_candidate": False, **row})
            continue
        difference = local["wst_completed_routes"] - row["completed_routes"]
        candidate = difference > 0 or row["dsp_dropped_routes"] > 0
        findings.append({
            "date": row["date"],
            "status": "review" if candidate else "pass",
            "direction": "potential_amazon_undercount" if difference > 0 else ("reported_dsp_drop" if row["dsp_dropped_routes"] else "matched"),
            "difference": difference,
            "dispute_candidate": candidate,
            **row,
            **local,
        })
    case = case_document(
        module_id="capacity_reliability",
        external_key=f"{year}-W{week:02d}",
        findings=findings,
        evidence_items=[evidence(capacity_path, "amazon_capacity_reliability"), *wst_evidence],
        blocking_evidence=sorted(set(blockers)),
    )
    case["summary"] = {
        "amazon_completed_routes": sum(row["completed_routes"] for row in amazon),
        "wst_route_coded_executions": sum((wst.get(row["date"]) or {}).get("wst_completed_routes", 0) for row in amazon),
        "dsp_dropped_routes": sum(row["dsp_dropped_routes"] for row in amazon),
        "days_reviewed": len(amazon),
    }
    return case


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--capacity-report", type=Path, required=True)
    parser.add_argument("--wst-dir", type=Path, required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--week", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = reconcile(args.capacity_report.resolve(strict=True), args.wst_dir.resolve(strict=True), args.year, args.week)
    write_json_atomic(args.output, result)
    print(json.dumps({"output": str(args.output), "status": result["status"], "candidate_count": result["candidate_count"]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
