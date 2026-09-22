#!/usr/bin/env python3
"""
JECS Unified API Server for DSP Operations.

Provides endpoints for:
- Daily Entry Form (drivers, vans, entries)
- Driver Performance Dashboard
- Fleet Cost Optimization
- Dispute Detection
- Route Monitoring
- Payroll Reconciliation

Usage:
    python3 scripts/jecs_api_server.py
    
Then open: http://localhost:8000/amazon-dsp-kpi-dashboard.html
"""

import json
import csv
import imaplib
import email
import math
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import threading
import ssl
import base64
import uuid
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
from pathlib import Path
from http.server import HTTPServer, BaseHTTPRequestHandler
import urllib.parse
import urllib.error
import urllib.request
from email import policy
from email.header import decode_header, make_header
from pypdf import PdfReader

import assistant_service

ROOT = Path(__file__).resolve().parents[1]
TENANT_ROOT = ROOT / "data/tenants"
DB_PATH = ROOT / "data/dsp_operations.db"
DASHBOARD_DIR = ROOT / "data/dashboards"
DASHBOARD_PATH = DASHBOARD_DIR / "amazon-dsp-kpi-dashboard.html"
FLEET_REVIEW_DIR = ROOT / "data/fleet_reviews"
WEAR_TEAR_TARGET_PERCENT = 80.0
WEAR_TEAR_STRETCH_PERCENT = 85.0


def _read_json(path):
    with open(path, 'r', encoding='utf-8') as source:
        return json.load(source)


def _test_adp_connection(tenant):
    """Validate the saved ADP mTLS client and read-only worker scope.

    Secrets are supplied to curl over stdin. Certificate material is written
    only to owner-readable temporary files and removed before this returns.
    """
    client_id = get_secret(tenant, 'adp', 'production-client-id')
    client_secret = get_secret(tenant, 'adp', 'production-client-secret')
    certificate = get_secret(tenant, 'adp', 'production-certificate-pem')
    private_key = get_secret(tenant, 'adp', 'production-private-key-pem')
    paths = []

    def request(url, method='GET', headers=None, body=None):
        arguments = [
            'curl', '-q', '--silent', '--show-error', '--http1.1',
            '--config', '-', '--cert', paths[0], '--key', paths[1],
            '--max-time', '45', '--write-out', '\n__ADP_HTTP_STATUS__:%{http_code}',
        ]
        if body is not None:
            arguments.extend(['--data-raw', body])
        config_lines = [
            f'url = "{url}"', f'request = "{method}"',
            'header = "Accept: application/json"',
        ]
        for name, value in (headers or {}).items():
            config_lines.append(f'header = "{name}: {value}"')
        completed = subprocess.run(
            arguments, input=('\n'.join(config_lines) + '\n').encode(),
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=50, check=False)
        marker = b'\n__ADP_HTTP_STATUS__:'
        marker_at = completed.stdout.rfind(marker)
        if completed.returncode or marker_at < 0:
            raise RuntimeError('ADP HTTPS request failed')
        status = int(completed.stdout[marker_at + len(marker):].strip())
        try:
            document = json.loads(completed.stdout[:marker_at])
        except (TypeError, ValueError):
            raise RuntimeError(f'ADP returned an invalid response (HTTP {status})')
        return status, document

    try:
        for material in (certificate, private_key):
            handle = tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', delete=False)
            paths.append(handle.name)
            try:
                handle.write(material)
                handle.flush()
            finally:
                handle.close()
        authorization = base64.b64encode(
            f'{client_id}:{client_secret}'.encode()).decode('ascii')
        token_status, token = request(
            'https://accounts.adp.com/auth/oauth/v2/token', method='POST',
            headers={
                'Authorization': f'Basic {authorization}',
                'Content-Type': 'application/x-www-form-urlencoded',
            }, body='grant_type=client_credentials')
        access_token = token.get('access_token') if isinstance(token, dict) else None
        if token_status != 200 or not access_token:
            raise RuntimeError(f'ADP OAuth authentication failed (HTTP {token_status})')
        worker_status, workers = request(
            'https://api.adp.com/hr/v2/workers?$top=1',
            headers={'Authorization': f'Bearer {access_token}'})
        if worker_status != 200:
            raise RuntimeError(f'ADP worker access failed (HTTP {worker_status})')
        return len(workers.get('workers', [])) if isinstance(workers, dict) else 0
    finally:
        for path in paths:
            try:
                os.unlink(path)
            except FileNotFoundError:
                pass


def _normalized_unit(value):
    return re.sub(r'[^A-Z0-9]', '', (value or '').upper())


def _days_until(value, today):
    if not value:
        return None
    try:
        return (datetime.fromisoformat(value).date() - today).days
    except (TypeError, ValueError):
        return None


def _latest_pave_upload_preview(tenant="jecs"):
    record = tenant_store.latest_confirmed(tenant, "pave")
    if not record:
        return None, None
    preview = parse_pave_export(tenant_store.read_bytes(tenant, record["id"]), record["originalFilename"])
    return record, preview if preview.get("ok") else None


def build_fleet_compliance_payload(tenant="jecs"):
    """Reconcile Fleet Portal, dispatch readiness, DVIC, and PM evidence."""
    if tenant != DEFAULT_TENANT and not tenant_store.latest_confirmed(tenant, "pave"):
        return {
            'tenant': tenant, 'asOf': None, 'generatedAt': datetime.now(timezone.utc).isoformat(),
            'needsData': True, 'message': 'No tenant-scoped fleet source has been ingested',
            'summary': {
                'registeredFleet': 0, 'operational': 0, 'grounded': 0, 'ready': 0,
                'readinessRate': 0, 'pmDue': 0, 'pmDueSoon': 0,
                'pmSourceIssues': 0, 'pmUnmatchedVehicles': 0,
                'inspectionVehicles': 0, 'inspectionCoverageRate': 0,
                'openMaintenanceIssues': 0, 'statusCounts': {}, 'ownershipCounts': {},
            },
            'vehicles': [], 'rows': [], 'unmatchedPmIssues': [],
            'wearAndTear': None, 'paveAssessments': None, 'sources': [],
            'reconciliation': {
                'currentRosterVinCount': 0, 'portalVinCount': 0,
                'vinSetsMatch': None, 'note': 'Tenant fleet data has not been ingested',
            },
        }
    roster_path = FLEET_REVIEW_DIR / "2026-09-07/vehicles-1.json"
    inspection_path = FLEET_REVIEW_DIR / "2026-09-07/inspection-stats-1.json"
    pm_path = FLEET_REVIEW_DIR / "2026-09-07/pm-stats-1.json"
    maintenance_path = FLEET_REVIEW_DIR / "2026-09-07/maintenance-issues-1.json"
    snapshots = sorted(FLEET_REVIEW_DIR.glob("*/dashboard-fleet-snapshot.json"))
    snapshot_path = snapshots[-1]

    vehicles = _read_json(roster_path).get('data', {}).get('vehicles', [])
    inspection_data = _read_json(inspection_path).get('inspectionsStatList', [])
    pm_data = _read_json(pm_path).get('pmIssueStatusCount', [])
    maintenance_data = _read_json(maintenance_path)
    snapshot = _read_json(snapshot_path)
    wear_files = sorted(FLEET_REVIEW_DIR.glob("*/wear-and-tear-compliance.json"))
    wear_data = _read_json(wear_files[-1]) if wear_files else None
    lsc_files = sorted(FLEET_REVIEW_DIR.glob("*/wear-and-tear-lsc-cases.json"))
    lsc_data = _read_json(lsc_files[-1]) if lsc_files else {'cases': []}
    pave_upload, pave_preview = _latest_pave_upload_preview(tenant)
    today = datetime.now().date()

    note = snapshot.get('reconciliation', {}).get('note', '')
    grounded_match = re.search(r'marked grounded:\s*(.*?)\.', note, re.IGNORECASE)
    grounded_units = []
    if grounded_match:
        grounded_units = [re.sub(r'^and\s+', '', item.strip(), flags=re.IGNORECASE) for item in re.split(r',\s*|\s+and\s+', grounded_match.group(1)) if item.strip()]
    grounded_keys = {_normalized_unit(item) for item in grounded_units}

    inspections_by_vin = {}
    for item in inspection_data:
        total = sum(int(stat.get('totalInspectionsDone') or 0) for stat in item.get('inspectionStats', []))
        inspections_by_vin[item.get('vehicleIdentifier')] = {
            'count': total,
            'types': sorted({stat.get('inspectionType') for stat in item.get('inspectionStats', []) if stat.get('inspectionType')})
        }

    pm_by_vin = {}
    unmatched_pm_issues = []
    for item in pm_data:
        entries = []
        for stat in item.get('pmIssueStats', []):
            for issue in stat.get('pmIssues', []):
                entries.append({
                    'status': stat.get('status'),
                    'issueType': issue.get('issueType'),
                    'serviceStatus': issue.get('serviceStatus'),
                    'issueId': issue.get('pmIssueId')
                })
        pm_by_vin[item.get('vehicleIdentifier')] = entries

    roster_vins = {vehicle.get('vin') for vehicle in vehicles}
    for vin, entries in pm_by_vin.items():
        if vin not in roster_vins:
            unmatched_pm_issues.append({'vin': vin, 'issues': entries})

    overrides = {item.get('vin'): item for item in snapshot.get('ownershipOverrides', [])}
    rows = []
    for vehicle in vehicles:
        vin = vehicle.get('vin')
        unit = vehicle.get('dspVehicleId') or vin
        override = overrides.get(vin)
        ownership = (override or {}).get('currentClassification') or vehicle.get('vehicleOwnershipType') or 'UNKNOWN'
        grounded = _normalized_unit(unit) in grounded_keys
        inspection = inspections_by_vin.get(vin, {'count': 0, 'types': []})
        pm_issues = pm_by_vin.get(vin, [])
        pm_statuses = {item.get('status') for item in pm_issues}
        ownership_days = _days_until(vehicle.get('ownershipEndDate'), today)
        registration_days = _days_until(vehicle.get('registrationExpiryDate'), today)
        health = vehicle.get('healthStatuses') or {}
        health_exceptions = [key for key, value in health.items() if value != 'OPERATIONAL']

        issues = []
        severity = 0
        if grounded:
            issues.append({'category': 'Readiness', 'severity': 'critical', 'label': 'Grounded by dispatch'})
            severity = max(severity, 100)
        if 'DUE' in pm_statuses:
            issues.append({'category': 'Maintenance', 'severity': 'critical', 'label': 'Preventive maintenance due'})
            severity = max(severity, 90)
        if ownership_days is not None and ownership_days < 0:
            issues.append({'category': 'Ownership', 'severity': 'critical', 'label': f'Ownership term ended {abs(ownership_days)} days ago'})
            severity = max(severity, 85)
        if registration_days is not None and registration_days < 0:
            issues.append({'category': 'Registration', 'severity': 'critical', 'label': f'Registration expired {abs(registration_days)} days ago'})
            severity = max(severity, 85)
        if 'DUE_SOON' in pm_statuses:
            issues.append({'category': 'Maintenance', 'severity': 'warning', 'label': 'Preventive maintenance due soon'})
            severity = max(severity, 60)
        if ownership_days is not None and 0 <= ownership_days <= 30:
            issues.append({'category': 'Ownership', 'severity': 'warning', 'label': f'Ownership term ends in {ownership_days} days'})
            severity = max(severity, 55)
        if registration_days is not None and 0 <= registration_days <= 30:
            issues.append({'category': 'Registration', 'severity': 'warning', 'label': f'Registration expires in {registration_days} days'})
            severity = max(severity, 55)
        for category in health_exceptions:
            issues.append({'category': 'Portal health', 'severity': 'warning', 'label': f'{category.replace("_", " ").title()} requires attention'})
            severity = max(severity, 60)
        if inspection['count'] == 0:
            issues.append({'category': 'Inspection evidence', 'severity': 'info', 'label': 'No DVIC record in the September 7 evidence pull'})
            severity = max(severity, 20)

        if grounded:
            compliance_status = 'grounded'
            next_action = 'Keep out of service; confirm repair disposition and dispatch release.'
        elif any(item['severity'] == 'critical' for item in issues):
            compliance_status = 'action_required'
            next_action = 'Resolve expired documentation or overdue maintenance before assignment.'
        elif any(item['severity'] == 'warning' for item in issues):
            compliance_status = 'monitor'
            next_action = 'Schedule the due item and confirm completion evidence.'
        elif inspection['count'] == 0:
            compliance_status = 'evidence_gap'
            next_action = 'Confirm DVIC completion in the current inspection window.'
        else:
            compliance_status = 'ready'
            next_action = 'No immediate action.'

        rows.append({
            'vin': vin,
            'unit': unit,
            'year': vehicle.get('year'),
            'make': vehicle.get('make'),
            'model': vehicle.get('model'),
            'registrationNumber': vehicle.get('registrationNo'),
            'registrationState': vehicle.get('registeredState'),
            'registrationExpiryDate': vehicle.get('registrationExpiryDate'),
            'registrationDaysRemaining': registration_days,
            'ownership': ownership,
            'provider': vehicle.get('vehicleProvider'),
            'ownershipEndDate': vehicle.get('ownershipEndDate'),
            'ownershipDaysRemaining': ownership_days,
            'operationalStatus': 'GROUNDED' if grounded else 'OPERATIONAL',
            'portalOperationalStatus': vehicle.get('operationalStatus'),
            'complianceStatus': compliance_status,
            'priority': severity,
            'inspectionCount': inspection['count'],
            'inspectionTypes': inspection['types'],
            'pmIssues': pm_issues,
            'healthStatuses': health,
            'issues': issues,
            'nextAction': next_action,
            'serviceTier': vehicle.get('serviceTier'),
            'lastRouteCompletedInDays': vehicle.get('lastRouteCompletedInDays')
        })

    pave_by_vin = {item['vin']: item for item in (pave_preview or {}).get('latestByVin', [])}
    rows_by_vin = {row['vin']: row for row in rows}
    for vin, assessment in pave_by_vin.items():
        row = rows_by_vin.get(vin)
        if not row:
            continue
        row.update({
            'paveGrade': assessment.get('grade'),
            'paveGradeLabel': assessment.get('gradeLabel'),
            'paveConditionScore': assessment.get('conditionScore'),
            'paveHasNewDamage': assessment.get('hasNewDamage'),
            'paveGroundingRisk': assessment.get('groundingRisk'),
            'paveAssessedAt': assessment.get('createdAt'),
            'paveSessionKey': assessment.get('sessionKey'),
        })
        if assessment.get('groundingRisk'):
            row['issues'].append({'category': 'PAVE', 'severity': 'critical', 'label': 'Latest PAVE assessment identifies grounding risk'})
            row['priority'] = max(row['priority'], 4)
            if row['complianceStatus'] != 'grounded':
                row['complianceStatus'] = 'action_required'
            row['nextAction'] = 'Review PAVE grounding evidence and repair before dispatch.'
        elif assessment.get('grade') == 2:
            row['issues'].append({'category': 'PAVE', 'severity': 'warning', 'label': 'Latest PAVE grade is Poor'})
            row['priority'] = max(row['priority'], 3)
            if row['complianceStatus'] == 'ready':
                row['complianceStatus'] = 'action_required'
            row['nextAction'] = 'Repair documented PAVE damage and complete a replacement assessment.'

    rows.sort(key=lambda item: (-item['priority'], item['unit']))
    status_counts = {}
    ownership_counts = {}
    for row in rows:
        status_counts[row['complianceStatus']] = status_counts.get(row['complianceStatus'], 0) + 1
        ownership_counts[row['ownership']] = ownership_counts.get(row['ownership'], 0) + 1

    wear_and_tear = None
    if wear_data:
        eligible_denominator = int(wear_data.get('eligibleVehicleCount') or len(rows))
        current_percent = float(wear_data.get('currentPercent') or 0)
        # Management target is an application policy, not a value inherited
        # from an older evidence snapshot. Keep it stable across exports.
        target_percent = WEAR_TEAR_TARGET_PERCENT
        stretch_percent = WEAR_TEAR_STRETCH_PERCENT
        current_count = int(wear_data.get('wearTearPassingCount') or round(current_percent / 100 * eligible_denominator))
        target_count = math.ceil(target_percent / 100 * eligible_denominator)
        stretch_count = math.ceil(stretch_percent / 100 * eligible_denominator)
        rows_by_vin = {row['vin']: row for row in rows}

        lsc_cases = []
        cases_by_vin = {}
        for case in lsc_data.get('cases', []):
            linked_vehicle = rows_by_vin.get(case.get('vin'))
            normalized_case = {
                **case,
                'unit': linked_vehicle.get('unit') if linked_vehicle else None,
                'operationalStatus': linked_vehicle.get('operationalStatus') if linked_vehicle else None,
            }
            lsc_cases.append(normalized_case)
            if case.get('vin'):
                cases_by_vin.setdefault(case['vin'], []).append(case.get('caseNumber'))

        repair_candidates = []
        for candidate in wear_data.get('poorGradeVehicles', []):
            linked_vehicle = rows_by_vin.get(candidate.get('vin'))
            repair_candidates.append({
                **candidate,
                'unit': linked_vehicle.get('unit') if linked_vehicle else 'Roster match required',
                'operationalStatus': linked_vehicle.get('operationalStatus') if linked_vehicle else 'UNKNOWN',
                'ownership': linked_vehicle.get('ownership') if linked_vehicle else 'UNKNOWN',
                'provider': linked_vehicle.get('provider') if linked_vehicle else None,
                'caseNumbers': cases_by_vin.get(candidate.get('vin'), []),
                'recommendedAction': (
                    'Complete repairs and the case-required replacement FCA before requesting ungrounding.'
                    if cases_by_vin.get(candidate.get('vin'))
                    else 'Repair to Fair+ (grade 3 or better), complete a new FCA, and verify the rolling metric posts.'
                )
            })
        repair_candidates.sort(key=lambda item: (
            item.get('operationalStatus') != 'OPERATIONAL',
            item.get('lastPave') or '',
            item.get('unit') or ''
        ))
        wear_and_tear = {
            **wear_data,
            'targetPercent': target_percent,
            'stretchPercent': stretch_percent,
            'planningDenominator': eligible_denominator,
            'estimatedCurrentCompliant': current_count,
            'targetCompliant': target_count,
            'minimumAdditionalCompliant': max(target_count - current_count, 0),
            'stretchCompliant': stretch_count,
            'stretchAdditionalCompliant': max(stretch_count - current_count, 0),
            'percentagePointGap': max(target_percent - current_percent, 0),
            'repairCandidates': repair_candidates,
            'lscCases': lsc_cases,
            'openLscCaseCount': sum(1 for case in lsc_cases if case.get('status') != 'closed'),
            'actions': [
                {
                    'id': 'prioritize-operational-poor',
                    'title': 'Start with operational grade-2 vehicles',
                    'detail': 'Repair and reassess operational units first so the target can improve without waiting for grounded-vehicle release.'
                },
                {
                    'id': 'complete-minimum',
                    'title': f'Close at least {max(target_count - current_count, 0)} additional compliant assessments',
                    'detail': f'The exact report threshold is {target_count} of {eligible_denominator} Fair+ vehicles. Upload complete FCA evidence and confirm the rolling metric posts.'
                },
                {
                    'id': 'build-buffer',
                    'title': f'Schedule {max(stretch_count - current_count, 0)} completions for a {stretch_percent:.0f}% buffer',
                    'detail': f'{stretch_count} of {eligible_denominator} equals {round(stretch_count / eligible_denominator * 100, 1)}%, protecting against rejected evidence and posting lag.'
                },
                {
                    'id': 'close-lsc-evidence',
                    'title': 'Advance open LSC cases without crediting them as complete',
                    'detail': 'Attach repair/FCA evidence to each case and keep the vehicle outside the compliant numerator until Amazon accepts the assessment.'
                }
            ]
        }

    return {
        'asOf': snapshot.get('asOf'),
        'generatedAt': datetime.now().isoformat(timespec='seconds'),
        'summary': {
            'registeredFleet': len(rows),
            'operational': sum(1 for row in rows if row['operationalStatus'] == 'OPERATIONAL'),
            'grounded': sum(1 for row in rows if row['operationalStatus'] == 'GROUNDED'),
            'readinessRate': round(sum(1 for row in rows if row['operationalStatus'] == 'OPERATIONAL') / len(rows) * 100, 1) if rows else 0,
            'pmDue': sum(1 for row in rows if any(item.get('status') == 'DUE' for item in row['pmIssues'])),
            'pmDueSoon': sum(1 for row in rows if any(item.get('status') == 'DUE_SOON' for item in row['pmIssues'])),
            'pmSourceIssues': sum(len(entries) for entries in pm_by_vin.values()),
            'pmUnmatchedVehicles': len(unmatched_pm_issues),
            'inspectionVehicles': sum(1 for row in rows if row['inspectionCount'] > 0),
            'inspectionCoverageRate': round(sum(1 for row in rows if row['inspectionCount'] > 0) / len(rows) * 100, 1) if rows else 0,
            'openMaintenanceIssues': maintenance_data.get('totalIssuesCount', 0),
            'statusCounts': status_counts,
            'ownershipCounts': ownership_counts
        },
        'vehicles': rows,
        'unmatchedPmIssues': unmatched_pm_issues,
        'wearAndTear': wear_and_tear,
        'paveAssessments': ({
            **pave_preview['summary'],
            'filename': pave_upload['originalFilename'],
            'uploadedAt': pave_upload['uploadedAt'],
            'confirmedAt': pave_upload['confirmedAt'],
            'unmatchedVins': sorted(set(pave_by_vin) - set(rows_by_vin)),
        } if pave_preview else None),
        'sources': [
            {'label': 'Dispatch readiness', 'asOf': snapshot.get('asOf'), 'path': snapshot.get('readinessSource')},
            {'label': 'Fleet roster and ownership', 'asOf': snapshot.get('ownershipAsOf'), 'path': snapshot.get('ownershipClassificationSource')},
            {'label': 'DVIC inspection evidence', 'asOf': '2026-09-07', 'path': str(inspection_path.relative_to(ROOT))},
            {'label': 'Preventive maintenance', 'asOf': '2026-09-07', 'path': str(pm_path.relative_to(ROOT))},
            *([{'label': 'Quarterly Wear & Tear report', 'asOf': wear_data.get('reportedAt', '')[:10], 'path': str(wear_files[-1].relative_to(ROOT))}] if wear_data else []),
            *([{'label': 'Wear & Tear LSC case register', 'asOf': lsc_data.get('asOf'), 'path': str(lsc_files[-1].relative_to(ROOT))}] if lsc_files else []),
            *([{'label': 'PAVE Fleet Dashboard CSV', 'asOf': pave_preview['summary'].get('latestAt', '')[:10], 'path': pave_upload['storageKey']}] if pave_preview else []),
        ],
        'reconciliation': snapshot.get('reconciliation', {})
    }


def get_connection():
    """Get a database connection."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn



RENTAL_RECON_PATH = FLEET_REVIEW_DIR / "2026-09-07/three-month-reconciliation/JEC-June-August-Rental-Reconciliation.xlsx"
_FLEET_COST_CACHE = {}


def build_fleet_cost_reconciliation(tenant=None):
    """Fleet cost reconciliation for a tenant.

    The "what I paid" side prefers the tenant's confirmed Digits upload and
    falls back to the owner-reviewed workbook. Amazon coverage still comes from
    the reconciliation invoices until the Amazon Payments connector lands.
    """
    tenant = tenant or DEFAULT_TENANT
    mtime = RENTAL_RECON_PATH.stat().st_mtime
    confirmed = latest_financial_upload(tenant)

    # The reviewed workbook is JECS-specific reference data. Any other tenant
    # sees only what they have uploaded themselves.
    if not confirmed and tenant != DEFAULT_TENANT:
        return {
            'tenant': tenant, 'needsData': True, 'period': None, 'asOf': None,
            'summary': {}, 'months': [], 'vendors': [], 'amazonClasses': [],
            'invoiceBridge': [], 'charges': [], 'fleet': [], 'notes': [],
            'dataSources': [{'side': 'What I paid', 'label': 'No Digits export uploaded yet',
                             'kind': 'missing', 'reference': None, 'asOf': None}],
            'caveats': ['Upload a Digits export on the Connections screen to populate this view.'],
        }

    cache_key = (tenant, mtime, (confirmed or {}).get('contentSha256'))
    cached = _FLEET_COST_CACHE.get('payload')
    if cached and _FLEET_COST_CACHE.get('key') == cache_key:
        return cached
    from openpyxl import load_workbook
    wb = load_workbook(RENTAL_RECON_PATH, data_only=True, read_only=True)

    def rows(sheet):
        return [r for r in wb[sheet].iter_rows(values_only=True) if any(v is not None for v in r)]

    def num(v):
        return round(float(v), 2) if isinstance(v, (int, float)) else None

    months = ['June', 'July', 'August']
    month_status = {'June': 'final', 'July': 'final', 'August': 'advance'}
    monthly = {r[0]: r for r in rows('Monthly reconciliation')[1:]}

    def series(label):
        return [num(monthly[label][i]) for i in (1, 2, 3)]

    cost = series('Included fleet expenses')
    coverage = series('Total rental/LMR/lease coverage')
    difference = series('Difference after included charges')
    vendors = {
        'Enterprise Rent-A-Car': series('Enterprise Rent-A-Car'),
        'Hertz': series('Hertz'),
        'MerchAuto9150 Corp': series('MerchAuto9150 Corp'),
        'Element Fleet': series('Element Fleet'),
    }
    third_party = [round(e + h, 2) for e, h in zip(vendors['Enterprise Rent-A-Car'], vendors['Hertz'])]
    rental_lease_cov = series('Amazon rental + lease coverage')
    lmr_cov = series('Amazon LMR coverage')
    full_cov = series('Full Amazon fleet coverage')

    month_rows = []
    for i, m in enumerate(months):
        month_rows.append({
            'month': m, 'status': month_status[m], 'invoiceBasis': monthly['Invoice basis'][i + 1],
            'includedCost': cost[i], 'amazonCoverage': coverage[i], 'difference': difference[i],
            'coverageRate': round(coverage[i] / cost[i] * 100, 1) if cost[i] else None,
            'thirdPartyRentalCost': third_party[i], 'rentalLeaseCoverage': rental_lease_cov[i],
            'rentalLeaseBalance': round(rental_lease_cov[i] - third_party[i], 2),
            'lmrCost': vendors['MerchAuto9150 Corp'][i], 'lmrCoverage': lmr_cov[i],
            'lmrBalance': round(lmr_cov[i] - vendors['MerchAuto9150 Corp'][i], 2),
            'elementCost': vendors['Element Fleet'][i], 'acuraExcluded': num(monthly['Acura excluded'][i + 1]),
            'rawExportTotal': num(monthly['Raw export total'][i + 1]), 'fullAmazonCoverage': full_cov[i],
        })

    class_rows = []
    for r in rows('Amazon class coverage')[1:]:
        cat = r[0]
        cov = [num(r[1]), num(r[2]), num(r[3])]
        days = [num(r[4]), num(r[5]), num(r[6])]
        class_rows.append({
            'category': cat, 'coverage': cov, 'vehicleDays': days,
            'perVehicleDay': [round(c / d, 2) if d else None for c, d in zip(cov, days)],
            'group': 'lmr' if 'Last Mile Rental' in cat else 'rental_lease' if cat in ('DSP Leased Van', 'Rental Van') else 'branded',
        })

    bridge = []
    for r in rows('Invoice bridge')[1:]:
        bridge.append({'period': r[0], 'finalGross': num(r[1]), 'priorAdvanceDeducted': num(r[2]),
                       'netReconciliation': num(r[3]), 'invoiceIssued': r[4],
                       'note': None if isinstance(r[1], (int, float)) else str(r[1])})

    charges = []
    for r in rows('All Digits charges')[1:]:
        charges.append({'month': r[0], 'datePosted': str(r[2])[:10] if r[2] else None, 'vendor': r[3], 'account': r[4],
                        'netCharge': num(r[7]), 'treatment': r[8], 'memo': r[9], 'vin': r[10], 'invoice': r[11],
                        'serviceStart': str(r[12])[:10] if r[12] else None, 'serviceEnd': str(r[13])[:10] if r[13] else None})

    # Tenant-uploaded Digits export takes precedence for charges and the
    # monthly cost side. Amazon coverage is left untouched.
    charge_sources = [{
        'side': 'What I paid · June/July',
        'label': 'Reviewed reconciliation workbook',
        'kind': 'workbook',
        'reference': str(RENTAL_RECON_PATH.relative_to(ROOT)),
        'asOf': '2026-09-07',
    }]
    if confirmed:
        try:
            parsed = parse_financial_export(tenant_store.read_bytes(tenant, confirmed['id']),
                                            confirmed['originalFilename'], tenant=tenant)
        except Exception:
            parsed = {'ok': False}
        uploaded = [c for c in parsed.get('charges', []) if c['treatment'] == 'INCLUDE']
        by_month = {}
        for charge in uploaded:
            by_month.setdefault(charge['month'], []).append(charge)
        if parsed.get('ok') and any(month in by_month for month in months):
            uploaded_months = set(by_month).intersection(months)
            charges = [c for c in charges if c['month'] not in uploaded_months] + parsed['charges']
            cost = [round(sum(c['netCharge'] for c in by_month[m]), 2)
                    if m in uploaded_months else cost[index]
                    for index, m in enumerate(months)]
            for name in list(vendors):
                vendors[name] = [round(sum(c['netCharge'] for c in by_month[m]
                                           if c['vendor'] == name), 2)
                                 if m in uploaded_months else vendors[name][index]
                                 for index, m in enumerate(months)]
            third_party = [round(e + h, 2) for e, h in zip(vendors['Enterprise Rent-A-Car'], vendors['Hertz'])]
            difference = [round(cov - cst, 2) for cov, cst in zip(coverage, cost)]
            for index, row in enumerate(month_rows):
                row['includedCost'] = cost[index]
                row['difference'] = difference[index]
                row['coverageRate'] = round(coverage[index] / cost[index] * 100, 1) if cost[index] else None
                row['thirdPartyRentalCost'] = third_party[index]
                row['rentalLeaseBalance'] = round(rental_lease_cov[index] - third_party[index], 2)
                row['lmrCost'] = vendors['MerchAuto9150 Corp'][index]
                row['lmrBalance'] = round(lmr_cov[index] - vendors['MerchAuto9150 Corp'][index], 2)
                row['elementCost'] = vendors['Element Fleet'][index]
            charge_sources = [source for source in charge_sources
                              if any(month not in uploaded_months for month in ('June', 'July'))]
            charge_sources.append({
                'side': 'What I paid · ' + '/'.join(sorted(uploaded_months, key=months.index)),
                'label': 'Your confirmed ' + parsed.get('providerLabel', 'accounting') + ' export',
                'provider': parsed.get('provider'),
                'kind': 'tenant_upload',
                'reference': confirmed['originalFilename'],
                'uploadId': confirmed['id'],
                'contentSha256': confirmed['contentSha256'],
                'uploadedBy': confirmed['uploadedBy'],
                'asOf': confirmed.get('confirmedAt') or confirmed['uploadedAt'],
            })

    current_amazon = None
    payment_candidates = []
    for payment_path in sorted((FLEET_REVIEW_DIR / '2026-09-07' / 'payments').glob('*.csv')):
        try:
            with payment_path.open(newline='', encoding='utf-8-sig') as source:
                payment_rows = list(csv.DictReader(source))
        except (OSError, csv.Error):
            continue
        dated = [row for row in payment_rows if row.get('Date')]
        if dated:
            payment_candidates.append((max(row['Date'] for row in dated), payment_path, payment_rows))
    if payment_candidates:
        latest_date, payment_path, payment_rows = max(payment_candidates, key=lambda item: item[0])
        relevant = [row for row in payment_rows if row.get('Vehicle Description') in (
            'Rental Van', 'DSP Leased Van', 'Branded Last Mile Rental Van')]
        current_amazon = {
            'period': datetime.strptime(latest_date, '%Y/%m/%d').strftime('%B %Y'),
            'asOf': latest_date.replace('/', '-'),
            'status': 'advance' if any(row.get('Payment Type') == 'PrePayment' for row in payment_rows) else 'final',
            'invoiceNumber': next((row.get('Invoice Number') for row in payment_rows if row.get('Invoice Number')), None),
            'rentalLmrLeaseCoverage': round(sum(float(row.get('Amount') or 0) for row in relevant), 2),
            'fullFleetCoverage': round(sum(float(row.get('Amount') or 0) for row in payment_rows), 2),
            'source': str(payment_path.relative_to(ROOT)),
            'costStatus': 'Awaiting a confirmed tenant accounting export for this month',
        }

    fleet = []
    for r in rows('Current fleet plus two')[1:]:
        fleet.append({'unit': r[0], 'vin': r[1], 'year': r[2], 'make': r[3], 'model': r[4], 'status': r[6],
                      'operationalStatus': r[7], 'ownership': r[8], 'provider': r[9],
                      'ownershipStart': r[10], 'ownershipEnd': r[11]})

    notes = [{'topic': r[0], 'detail': r[1]} for r in rows('Notes')[1:]]
    total_cost = round(sum(cost), 2); total_cov = round(sum(coverage), 2)
    payload = {
        'period': 'June-August 2026', 'asOf': '2026-09-07',
        'source': str(RENTAL_RECON_PATH.relative_to(ROOT)),
        'summary': {
            'threeMonthIncludedCost': total_cost, 'threeMonthAmazonCoverage': total_cov,
            'threeMonthDifference': round(sum(difference), 2),
            'coverageRate': round(total_cov / total_cost * 100, 1) if total_cost else None,
            'augustDifference': difference[2],
            'thirdPartyRentalCost': round(sum(third_party), 2), 'rentalLeaseCoverage': round(sum(rental_lease_cov), 2),
            'lmrCost': round(sum(vendors['MerchAuto9150 Corp']), 2), 'lmrCoverage': round(sum(lmr_cov), 2),
            'elementCost': round(sum(vendors['Element Fleet']), 2), 'acuraExcluded': round(sum(series('Acura excluded')), 2),
            'fullAmazonCoverage': round(sum(full_cov), 2),
            'includedTransactions': sum(1 for c in charges if c['treatment'] == 'INCLUDE'),
            'excludedTransactions': sum(1 for c in charges if c['treatment'] != 'INCLUDE'),
            'unmatchedVinCharges': sum(1 for c in charges if c['treatment'] == 'INCLUDE' and not c['vin']),
        },
        'months': month_rows,
        'vendors': [{'vendor': k, 'monthly': v, 'total': round(sum(v), 2),
                     'coverageClass': 'Amazon LMR coverage' if k.startswith('MerchAuto') else 'Rental + lease coverage' if k in ('Enterprise Rent-A-Car', 'Hertz') else 'Unallocated'}
                    for k, v in vendors.items()],
        'amazonClasses': class_rows, 'invoiceBridge': bridge, 'charges': charges, 'fleet': fleet, 'notes': notes,
        'tenant': tenant, 'needsData': False,
        'currentPeriod': current_amazon,
        'dataSources': [
            *charge_sources,
            {'side': 'What Amazon paid', 'label': 'Amazon reconciliation invoices', 'kind': 'workbook',
             'reference': str(RENTAL_RECON_PATH.relative_to(ROOT)), 'asOf': '2026-09-07'},
            *([{'side': 'Latest Amazon payment', 'label': f"Amazon invoice {current_amazon['invoiceNumber']}",
                'kind': 'connector_artifact', 'reference': current_amazon['source'],
                'asOf': current_amazon['asOf']}] if current_amazon else []),
        ],
        'caveats': [
            'Posting-period comparison: Digits dates are posting dates, not confirmed rental service periods.',
            'June and July use final Amazon reconciliation invoices; August uses the advance because the final was not locally available.',
            'Acura ($1,400/month) is excluded from the fleet comparison at owner direction.',
            'Differences are not net profit/loss and do not prove underpayment; VIN/service-period matching is still open.',
        ],
    }
    wb.close()
    _FLEET_COST_CACHE.update(payload=payload, key=cache_key)
    return payload



MODULE_STATUS_PATH = ROOT / "data/dashboards/platform/module-status.json"
MODULE_REGISTRY_DIR = ROOT / "platform/modules"
CASE_GLOBS = ["data/payment_reconciliation/*/*case*.json", "data/fixed_monthly*/**/*case*.json", "data/reconciliation*/**/*case*.json"]


def _to_money(value):
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, str):
        cleaned = re.sub(r'[^0-9.\-]', '', value)
        try:
            return round(float(cleaned), 2) if cleaned else None
        except ValueError:
            return None
    return None


def _load_module_cases():
    cases, seen = [], set()
    for pattern in CASE_GLOBS:
        for path in ROOT.glob(pattern):
            if path in seen or not path.is_file():
                continue
            seen.add(path)
            try:
                data = _read_json(path)
            except Exception:
                continue
            if not isinstance(data, dict) or 'module_id' not in data:
                continue
            approval = data.get('approval') if isinstance(data.get('approval'), dict) else {}
            submission = data.get('submission') if isinstance(data.get('submission'), dict) else {}
            cases.append({
                'moduleId': data.get('module_id'), 'externalKey': data.get('external_key'),
                'status': data.get('status'), 'disposition': data.get('disposition'), 'createdAt': data.get('created_at'),
                'candidateCount': data.get('candidate_count', 0), 'recoveredRoutes': data.get('recovered_routes'),
                'recoveredValue': data.get('recovered_value'), 'blockingEvidence': data.get('blocking_evidence') or [],
                'externalActionAuthorized': bool(data.get('external_action_authorized')),
                'approvalStatus': approval.get('status'), 'approvedBy': approval.get('approved_by') or approval.get('approver'),
                'submissionStatus': submission.get('status'),
                'submissionConfirmation': submission.get('confirmation') or submission.get('confirmation_number'),
                'findings': data.get('findings') or [], 'evidence': data.get('evidence') or [],
                'sourcePath': str(path.relative_to(ROOT)),
            })
    for path in ROOT.glob('data/fleet_reviews/fixed-monthly/approval-queue/*/approval.json'):
        try:
            d = _read_json(path)
        except Exception:
            continue
        inv = d.get('invoice') or {}; cand = d.get('candidate') or {}; sub = d.get('submission') or {}; dec = d.get('decision') or {}
        cases.append({
            'moduleId': 'fixed_monthly', 'externalKey': d.get('approval_id'), 'status': d.get('status'),
            'disposition': sub.get('outcome') or dec.get('decision'), 'createdAt': d.get('created_at'),
            'candidateCount': 1 if cand else 0, 'recoveredRoutes': None, 'recoveredValue': _to_money(cand.get('estimated_value')),
            'blockingEvidence': [], 'externalActionAuthorized': dec.get('decision') == 'YES',
            'approvalStatus': dec.get('decision'), 'approvedBy': dec.get('from'),
            'submissionStatus': sub.get('outcome'), 'submissionConfirmation': sub.get('confirmation'),
            'findings': [{'label': 'Invoice', 'value': inv.get('invoice_number'), 'detail': f"{inv.get('service_month')} - {inv.get('service_start')} to {inv.get('service_end')}"},
                         {'label': 'Candidate', 'value': cand.get('candidate_id'), 'detail': cand.get('proposed_wording')},
                         *[{'label': 'Fact', 'value': None, 'detail': f} for f in (cand.get('facts') or [])]],
            'evidence': [{'label': k, 'path': (d.get(k) or {}).get('path'), 'sha256': (d.get(k) or {}).get('sha256')} for k in ('report', 'pdf', 'source_invoice') if d.get(k)],
            'sourcePath': str(path.relative_to(ROOT)),
        })
    cases.sort(key=lambda c: c.get('createdAt') or '', reverse=True)
    return cases


def build_reimbursement_review_payload():
    """Go HQ replacement view: module registry + status feed + every case on disk."""
    status = _read_json(MODULE_STATUS_PATH) if MODULE_STATUS_PATH.exists() else {'modules': []}
    status_by_id = {m['id']: m for m in status.get('modules', [])}
    cases = _load_module_cases()
    schedules = {'fixed_monthly': 'Daily 10:00 ET review + 15 min approvals', 'weekly_payments': 'Daily 11:00 ET review cycle', 'capacity_reliability': 'Daily 11:00 ET review cycle'}
    modules = []
    for manifest_path in sorted(MODULE_REGISTRY_DIR.glob('*/module.json')):
        manifest = _read_json(manifest_path); mid = manifest['id']; st = status_by_id.get(mid, {})
        mod_cases = [c for c in cases if c['moduleId'] == mid]; runner = manifest.get('runner')
        modules.append({
            'id': mid, 'displayName': manifest.get('displayName') or st.get('display_name') or mid,
            'description': manifest.get('description'), 'goHqFunction': manifest.get('goHqFunction') or manifest.get('replaces'),
            'billingSku': manifest.get('billingSku'), 'implementationStatus': st.get('implementation_status') or manifest.get('status'),
            'state': st.get('state') or manifest.get('status'), 'headline': st.get('headline'), 'detail': st.get('detail'),
            'runner': runner, 'runnerExists': bool(runner) and (ROOT / runner).exists(),
            'requiredInputs': manifest.get('requiredInputs') or manifest.get('inputs') or [],
            'hasSubmissionAdapter': mid == 'fixed_monthly', 'schedule': schedules.get(mid, 'On import'),
            'caseCount': len(mod_cases), 'openCaseCount': sum(1 for c in mod_cases if c['status'] not in ('closed', 'submitted')),
            'recoveredValue': round(sum(c['recoveredValue'] or 0 for c in mod_cases), 2),
            'latestCase': mod_cases[0] if mod_cases else None, 'verifiedRoutes': st.get('verified_routes'),
        })
    order = {'active': 0, 'ready_for_import': 1, 'foundation': 2}
    modules.sort(key=lambda m: (order.get(m['state'], 9), m['displayName']))
    return {
        'generatedAt': status.get('generated_at'), 'servedAt': datetime.now().isoformat(timespec='seconds'), 'tenant': status.get('tenant'),
        'summary': {'modules': len(modules), 'active': sum(1 for m in modules if m['state'] == 'active'),
                    'readyForImport': sum(1 for m in modules if m['state'] == 'ready_for_import'), 'cases': len(cases),
                    'openCases': sum(1 for c in cases if c['status'] not in ('closed', 'submitted')),
                    'recoveredValue': round(sum(c['recoveredValue'] or 0 for c in cases), 2),
                    'submitted': sum(1 for c in cases if c['submissionStatus'] or c['status'] == 'submitted')},
        'modules': modules, 'cases': cases,
        'schedules': [{'id': '1914db4c', 'name': 'Fixed Monthly Invoice Review', 'cadence': 'Daily 10:00 ET'},
                      {'id': '7554813f', 'name': 'Fixed Monthly Approval Processor', 'cadence': 'Every 15 min'},
                      {'id': '539fdf4d', 'name': 'Payment Reconciliation Review Cycle', 'cadence': 'Daily 11:00 ET'},
                      {'id': 'c4871993', 'name': 'Reconciliation Module Approval Monitor', 'cadence': 'Every 15 min'}],
    }




# ---------------------------------------------------------------- connections
# Importable both as "scripts.jecs_api_server" (tests) and as a direct script.
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import connections_registry, tenant_store
from scripts.financial_ingest import (
    PROVIDER_LABELS,
    load_vendor_rules,
    parse_financial_export,
    save_vendor_rules,
)
from scripts.pave_ingest import parse_pave_export
from scripts.secret_store import get_secret, has_secret, put_secret

DEFAULT_TENANT = "jecs"
SCORECARD_DATA_DIR = ROOT / "data/scorecard_data"

LOCAL_FEATURES = [
    ('dashboard', 'Dashboard', '/dashboard', 'implemented'),
    ('drivers', 'Drivers', '/drivers', 'implemented'),
    ('fleet_compliance', 'Fleet Compliance', '/fleet-compliance', 'implemented'),
    ('vans', 'Vans', '/vans', 'implemented'),
    ('route_monitor', 'Live Route Monitor', '/routes', 'implemented'),
    ('route_performance', 'Weekly Route Performance', '/route-performance', 'implemented'),
    ('disputes', 'Dispute Center', '/disputes', 'implemented'),
    ('payroll', 'Payroll', '/payroll', 'implemented'),
    ('weekly_evaluation', 'Weekly Evaluation', '/weekly-evaluation', 'implemented'),
    ('driver_performance', 'Driver Performance', '/performance', 'implemented'),
    ('connections', 'Connections', '/connections', 'implemented'),
    ('users', 'Users & Roles', '/users', 'implemented'),
    ('reimbursement_review', 'Reimbursement Review', '/reimbursement-review', 'implemented'),
    ('time_attendance', 'Time & Attendance', '/time-attendance', 'implemented'),
    ('fleet_costs', 'Fleet Costs', '/fleet-costs', 'implemented'),
    ('maintenance', 'Maintenance', '/maintenance', 'planned'),
    ('fuel', 'Fuel Tracking', '/fuel', 'planned'),
    ('settings', 'Settings', '/settings', 'implemented'),
    ('security', 'Security', '/security', 'planned'),
    ('notifications', 'Notifications', '/notifications', 'planned'),
    ('help', 'Help & Support', '/help', 'implemented'),
    ('feature_admin', 'Feature Management', '/admin/features', 'implemented'),
    ('super_admin', 'Tenant Administration', '/admin/tenants', 'implemented'),
    ('ai_admin', 'AI Assistant Setup', '/admin/ai', 'implemented'),
]


def _feature_state_path(tenant):
    return TENANT_ROOT / tenant / 'feature-overrides.json'


def _feature_overrides(tenant):
    path = _feature_state_path(tenant)
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text())
    except (ValueError, OSError):
        return {}


def _public_features(tenant):
    overrides = _feature_overrides(tenant)
    return [
        {'id': item[0], 'displayName': item[1], 'route': item[2], 'status': item[3],
         'permission': 'module.read', 'enabled': overrides.get(item[0], True)}
        for item in LOCAL_FEATURES
    ]


def _set_feature_override(tenant, feature_id, enabled):
    if feature_id not in {item[0] for item in LOCAL_FEATURES}:
        raise KeyError('unknown feature')
    path = _feature_state_path(tenant)
    path.parent.mkdir(parents=True, exist_ok=True)
    state = _feature_overrides(tenant)
    state[feature_id] = bool(enabled)
    path.write_text(json.dumps(state, indent=2, sort_keys=True))
    return next(item for item in _public_features(tenant) if item['id'] == feature_id)


def _member_state_path(tenant):
    return TENANT_ROOT / tenant / 'members.json'


def _local_members(tenant):
    path = _member_state_path(tenant)
    if path.exists():
        try:
            members = json.loads(path.read_text())
            if isinstance(members, list):
                return members
        except (ValueError, OSError):
            pass
    return [{'identitySubject': 'dev-user', 'email': 'dev@example.com', 'role': 'owner',
             'status': 'active', 'createdAt': '2026-09-21T00:00:00Z'}]


def _save_local_members(tenant, members):
    path = _member_state_path(tenant)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(members, indent=2, sort_keys=True))


def _local_tenant_state_path():
    return TENANT_ROOT / '_platform' / 'tenants.json'


def _local_tenants():
    path = _local_tenant_state_path()
    if path.exists():
        try:
            tenants = json.loads(path.read_text())
            if isinstance(tenants, list):
                return tenants
        except (ValueError, OSError):
            pass
    return [{'id': 'local-jecs', 'slug': 'jecs', 'displayName': 'JEC Logistics Solutions',
             'status': 'active', 'createdAt': '2026-09-21T00:00:00Z'}]


def _save_local_tenants(tenants):
    path = _local_tenant_state_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(tenants, indent=2, sort_keys=True))


LOCAL_MODULES = [
    {'id': 'executive_dashboard', 'displayName': 'Executive Operations Dashboard', 'status': 'in_development'},
    {'id': 'data_integrations', 'displayName': 'Secure Data Integrations', 'status': 'in_development'},
    {'id': 'fixed_monthly', 'displayName': 'Fixed Monthly Fleet Reconciliation', 'status': 'active'},
    {'id': 'weekly_payments', 'displayName': 'Weekly Variable and Incentive Review', 'status': 'active'},
]
_LOCAL_SUPPORT_SESSIONS = {}


def _valid_member_email(value):
    return bool(re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value or ''))


def _tenant_from_headers(headers):
    """Tenant comes from the authenticated context in production; the local
    server accepts the same header the React client sends."""
    tenant = (headers.get("x-tenant-id") or "").strip().lower() or DEFAULT_TENANT
    if not connections_registry.SLUG.match(tenant):
        raise ValueError("invalid tenant")
    return tenant


def _latest_upload_index(tenant):
    index = {}
    for record in tenant_store.list_files(tenant):
        if record["source"] not in index and record["status"] in ("uploaded", "parsed", "confirmed"):
            index[record["source"]] = record
    # Legacy "digits" uploads still satisfy the generalized financial source.
    if "financial_charges" not in index and "digits" in index:
        index["financial_charges"] = index["digits"]
    return index


def latest_financial_upload(tenant):
    """Most recent confirmed accounting export, whichever source id it used."""
    for source in connections_registry.FINANCIAL_SOURCES:
        record = tenant_store.latest_confirmed(tenant, source)
        if record:
            return record
    return None


def build_connections_payload(tenant):
    connections = connections_registry.list_connections(tenant, _latest_upload_index(tenant))
    # Manual uploads are data-ingestion sources, not persistent connections.
    # Keep them in the catalog but exclude them from connection-health KPIs.
    live_connections = [
        connection for connection in connections
        if connection["authKind"] != "manual_upload"
    ]
    counts = {}
    for connection in connections:
        counts[connection["status"]] = counts.get(connection["status"], 0) + 1
    active = sum(1 for connection in live_connections if connection["status"] == "healthy")
    health = "green" if live_connections and active == len(live_connections) else "yellow" if active else "red"
    return {
        "tenant": tenant,
        "servedAt": datetime.now().isoformat(timespec="seconds"),
        "summary": {
            "total": len(connections),
            "connectionTotal": len(live_connections),
            "connected": active,
            "active": active,
            "health": health,
            "needsAttention": sum(1 for c in live_connections if c["status"] in ("needs_reauth", "degraded")),
            "notConnected": sum(1 for c in live_connections if c["status"] == "not_connected"),
            "statusCounts": counts,
        },
        "connections": connections,
        "uploads": tenant_store.list_files(tenant),
        "secretPolicy": {
            "storage": "AWS Secrets Manager",
            "pathTemplate": connections_registry.secret_reference("{tenant}", "{connection}", "{name}"),
            "note": "The platform stores a reference only. Credential values are never written to the database, logs, or this API response.",
        },
    }


def _scorecard_number(value):
    text = str(value or '').strip().replace(',', '').replace('%', '')
    try:
        return float(text)
    except ValueError:
        return 0.0


def _dsp_scorecard_standing(week):
    """Read the fleet-level DSP score from the official weekly scorecard PDF.

    The Overview Dashboard CSV contains one score per delivery associate and
    must not be averaged into a DSP score. Missing official PDFs remain
    unavailable rather than silently substituting a DA average.
    """
    week_dir = SCORECARD_DATA_DIR / week
    official = sorted(
        path for path in week_dir.glob('*.pdf')
        if '_en_dspscorecard' in path.name.lower() and 'preview' not in path.name.lower()
    )
    previews = sorted(path for path in week_dir.glob('*.pdf') if 'dspscorecardpreview' in path.name.lower())
    # Amazon can republish a corrected scorecard after the original. Prefer the
    # newest republish timestamp; otherwise use the original official file.
    official.sort(key=lambda item: ('republish' in item.name.lower(), item.name), reverse=True)
    source = (official or previews or [None])[0]
    if source is None:
        return {'score': None, 'tier': None, 'source': None}
    try:
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(str(source)).pages)
    except Exception:
        return {'score': None, 'tier': None, 'source': source.name}
    match = re.search(r'Overall Standing:\s*([0-9]+(?:\.[0-9]+)?)\s*(?:\||,)\s*([^\n\r]+)', text, re.I)
    if not match:
        return {'score': None, 'tier': None, 'source': source.name}
    return {'score': float(match.group(1)), 'tier': ' '.join(match.group(2).split()), 'source': source.name}


def _scorecard_week(week):
    files = sorted((SCORECARD_DATA_DIR / week).glob('DSP_Overview_Dashboard_*.csv'))
    if not files:
        return None
    with files[-1].open(encoding='utf-8-sig', newline='') as source:
        rows = list(csv.DictReader(source))
    drivers = []
    for index, row in enumerate(rows):
        drivers.append({
            'driverId': row.get('Transporter ID') or f'{week}-{index + 1}',
            'driverName': (row.get('Delivery Associate') or row.get('Delivery Associate ') or '').strip(),
            'overallScore': _scorecard_number(row.get('Overall Score')),
            'deliveryScore': _scorecard_number(row.get('DCR')),
            'qualityScore': _scorecard_number(row.get('POD')),
            'score': _scorecard_number(row.get('Overall Score')),
            'grade': row.get('Overall Standing') or '',
        })
    if not drivers:
        return None
    average = lambda field: sum(_scorecard_number(row.get(field)) for row in rows) / len(rows)
    dsp_standing = _dsp_scorecard_standing(week)
    return {
        'period': week,
        'drivers': drivers,
        'overallScore': dsp_standing['score'],
        'overallStanding': dsp_standing['tier'],
        'dspScoreSource': dsp_standing['source'],
        'averageDaScore': average('Overall Score'),
        'pod': average('POD'),
        'dcr': average('DCR'),
        'cdf': sum(_scorecard_number(row.get('CDF DPMO')) for row in rows),
        'packages': int(sum(_scorecard_number(row.get('Packages Delivered')) for row in rows)),
        'activeDrivers': len(rows),
    }


_PERFORMANCE_DASHBOARD_CACHE = {'fingerprint': None, 'payload': None}


def _performance_source_fingerprint(weeks):
    """Cheaply detect scorecard source changes without reparsing every PDF."""
    sources = []
    for week in weeks:
        week_dir = SCORECARD_DATA_DIR / week
        for pattern in ('DSP_Overview_Dashboard_*.csv', '*.pdf'):
            for path in week_dir.glob(pattern):
                name = path.name.lower()
                if pattern == '*.pdf' and 'dspscorecard' not in name:
                    continue
                stat = path.stat()
                sources.append((str(path), stat.st_mtime_ns, stat.st_size))
    return tuple(sorted(sources))


def build_performance_dashboard_payload():
    weeks = sorted(
        item.name for item in SCORECARD_DATA_DIR.iterdir()
        if item.is_dir() and re.fullmatch(r'\d{4}-wk\d{2}', item.name)
    )[-13:]
    fingerprint = _performance_source_fingerprint(weeks)
    if (_PERFORMANCE_DASHBOARD_CACHE['fingerprint'] == fingerprint
            and _PERFORMANCE_DASHBOARD_CACHE['payload'] is not None):
        return _PERFORMANCE_DASHBOARD_CACHE['payload']
    history = [item for week in weeks if (item := _scorecard_week(week))]
    if not history:
        payload = {'period': None, 'generatedAt': datetime.now(timezone.utc).isoformat(),
                   'source': None, 'drivers': [], 'history': [],
                   'dspPerformance': {'overallScore': 0, 'deliveryScore': 0, 'safetyScore': 0,
                                      'qualityScore': 0, 'driverCount': 0, 'totalDeliveries': 0}}
        _PERFORMANCE_DASHBOARD_CACHE.update(fingerprint=fingerprint, payload=payload)
        return payload
    current = history[-1]
    drivers = current['drivers']
    # Trend consumers need weekly aggregates, not a duplicate driver roster
    # embedded in every historical point.
    summarized_history = [
        {key: value for key, value in item.items() if key != 'drivers'}
        for item in history
    ]
    payload = {
        'period': current['period'],
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'source': f"Amazon DSP scorecard {current['period']}",
        'drivers': drivers,
        'history': summarized_history,
        'dspPerformance': {
            'overallScore': current['overallScore'],
            'deliveryScore': current['dcr'],
            'safetyScore': 0,
            'qualityScore': current['pod'],
            'driverCount': current['activeDrivers'],
            'totalDeliveries': current['packages'],
        },
    }
    _PERFORMANCE_DASHBOARD_CACHE.update(fingerprint=fingerprint, payload=payload)
    return payload


def build_route_monitor_payload(period=None):
    """Export tenant route aggregates in the same schema used by the React screen."""
    conn = get_connection()
    cursor = conn.cursor()
    requested = period
    if not requested:
        cursor.execute("SELECT MAX(date) AS latest FROM amazon_routes")
        row = cursor.fetchone()
        requested = row['latest'] if row else None
    routes = []
    if requested:
        cursor.execute("""
            SELECT route_code, driver_id, date, stops, packages, status,
                   overall_score, pod, cdf, dsb, is_weekly_aggregate
              FROM amazon_routes
             WHERE date = ?
             ORDER BY route_code
        """, (requested,))
        for row in cursor.fetchall():
            route = dict(row)
            cursor.execute("SELECT name FROM drivers WHERE id = ?", (route.get('driver_id'),))
            driver = cursor.fetchone()
            route['driver_name'] = driver['name'] if driver else 'Unknown'
            routes.append(route)
    conn.close()
    return {
        'period': requested,
        'routes': routes,
        'routeCount': len(routes),
        'source': 'Amazon scorecard route aggregates',
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'needsData': not routes,
    }


def build_dashboard_operations_payload(tenant='jecs'):
    performance = (build_performance_dashboard_payload() if tenant == DEFAULT_TENANT else {
        'period': None, 'generatedAt': datetime.now(timezone.utc).isoformat(),
        'source': 'No tenant-scoped scorecard source available', 'needsData': True,
        'drivers': [], 'history': [],
        'dspPerformance': {'overallScore': None, 'deliveryScore': 0, 'safetyScore': 0,
                           'qualityScore': 0, 'driverCount': 0, 'totalDeliveries': 0},
    })
    fleet = build_fleet_compliance_payload(tenant)
    costs = build_fleet_cost_reconciliation(tenant)
    connections = build_connections_payload(tenant)
    sources = []
    for item in connections.get('connections', []):
        source_id = item.get('id')
        as_of = item.get('lastSuccessAt')
        status = item.get('status')
        if source_id == 'amazon':
            as_of = as_of or performance.get('period')
            status = status if performance.get('drivers') else 'needs_data'
        elif source_id == 'pave':
            as_of = as_of or fleet.get('generatedAt') or fleet.get('asOf')
            status = status if fleet.get('vehicles') else 'needs_data'
        elif item.get('authKind') == 'manual_upload':
            as_of = costs.get('asOf')
            status = 'needs_data' if costs.get('needsData') else 'current'
        sources.append({
            'id': source_id,
            'label': item.get('displayName') or item.get('name') or source_id,
            'asOf': as_of,
            'status': status,
            'feeds': item.get('feeds') or [],
        })
    # This endpoint is the dashboard's critical rendering path. Return only
    # fields consumed by that screen instead of serializing full driver,
    # vehicle, transaction, upload, and reconciliation records.
    dashboard_performance = {
        key: performance.get(key)
        for key in ('period', 'generatedAt', 'source', 'drivers', 'history', 'dspPerformance')
    }
    dashboard_fleet = {
        'asOf': fleet.get('asOf') or fleet.get('generatedAt'),
        'summary': fleet.get('summary') or {},
        'vehicles': [
            {'status': item.get('status'), 'ownership': item.get('ownership')}
            for item in (fleet.get('vehicles') or fleet.get('rows') or [])
        ],
    }
    dashboard_costs = {
        'asOf': costs.get('asOf'),
        'needsData': costs.get('needsData', False),
        'summary': costs.get('summary') or {},
    }
    dashboard_connections = {
        'summary': connections.get('summary') or {},
        'connections': [
            {key: item.get(key) for key in ('id', 'name', 'displayName', 'status', 'lastSuccessAt', 'kind', 'authKind')}
            for item in connections.get('connections', [])
        ],
    }
    return {
        'tenant': {'id': tenant, 'name': next((item.get('displayName') for item in _local_tenants()
                                               if item.get('slug') == tenant), tenant)},
        'generatedAt': datetime.now(timezone.utc).isoformat(),
        'needsData': tenant != DEFAULT_TENANT,
        'performance': dashboard_performance,
        'fleet': dashboard_fleet,
        'costs': dashboard_costs,
        'connections': dashboard_connections,
        'sources': sources,
    }


def build_assistant_snapshot(tenant, current_path='/dashboard'):
    """Return a bounded, read-only tenant snapshot for model grounding."""
    citations = [
        {'id': 'connections', 'label': 'Connection health', 'route': '/connections'},
        {'id': 'fleet-compliance', 'label': 'Fleet Compliance', 'route': '/fleet-compliance'},
        {'id': 'fleet-costs', 'label': 'Fleet Costs', 'route': '/fleet-costs'},
        {'id': 'disputes', 'label': 'Dispute Center', 'route': '/disputes'},
        {'id': 'driver-performance', 'label': 'Driver Performance', 'route': '/performance'},
        {'id': 'routes', 'label': 'Live Route Monitor', 'route': '/routes'},
        {'id': 'route-performance', 'label': 'Weekly Route Performance', 'route': '/route-performance'},
        {'id': 'payroll', 'label': 'Payroll', 'route': '/payroll'},
    ]
    connections = build_connections_payload(tenant)
    fleet = build_fleet_compliance_payload(tenant)
    fleet_rows = fleet.get('vehicles') or fleet.get('rows') or []
    priority_fleet = sorted(fleet_rows, key=lambda row: row.get('priority') or 0, reverse=True)[:12]
    costs = build_fleet_cost_reconciliation(tenant)
    conn = get_connection()
    cursor = conn.cursor()
    def rows(sql, params=()):
        cursor.execute(sql, params)
        return [dict(item) for item in cursor.fetchall()]
    try:
        drivers = rows("""SELECT ar.driver_id, d.name AS driver_name, ar.overall_score,
                                 ar.pod, ar.cdf, ar.dsb, ar.packages, ar.date
                            FROM amazon_routes ar JOIN drivers d ON d.id=ar.driver_id
                           WHERE ar.is_weekly_aggregate=1
                           ORDER BY ar.date DESC, ar.overall_score DESC LIMIT 20""")
        disputes = rows("""SELECT week, driver_id, metric, reason, status, priority
                             FROM disputes ORDER BY week DESC, priority ASC LIMIT 20""")
        route_period = rows("SELECT MAX(date) AS latest FROM amazon_routes")
        latest_date = route_period[0].get('latest') if route_period else None
        routes = rows("""SELECT route_code, driver_id, stops, packages, status
                           FROM amazon_routes WHERE date=? ORDER BY route_code LIMIT 80""", (latest_date,)) if latest_date else []
        payroll = rows("""SELECT COUNT(*) AS timecard_count,
                                  COALESCE(SUM(duration_hours),0) AS total_hours
                             FROM adp_timecards""")
    finally:
        conn.close()
    compact_fleet = [{key: row.get(key) for key in (
        'unit', 'vin', 'operationalStatus', 'complianceStatus', 'priority',
        'paveGradeLabel', 'paveGroundingRisk', 'paveHasNewDamage', 'nextAction')}
        for row in priority_fleet]
    return {
        'tenant': tenant,
        'generatedAt': datetime.now(timezone.utc).isoformat(timespec='seconds'),
        'currentPath': current_path,
        'connections': connections.get('summary'),
        'fleetCompliance': {
            'summary': fleet.get('summary'),
            'wearAndTear': fleet.get('wearAndTear'),
            'priorityVehicles': compact_fleet,
        },
        'fleetCosts': {
            key: costs.get(key) for key in ('period', 'summary', 'totals', 'months', 'dataSources')
        },
        'driverPerformance': drivers,
        'disputes': disputes,
        'routeMonitor': {'period': latest_date, 'routeCount': len(routes), 'routes': routes},
        'payroll': payroll[0] if payroll else {},
        'citations': citations,
    }


def reconcile_saved_browser_sessions(tenant="jecs"):
    """Reuse persistent provider sessions after an API restart.

    This runs headlessly and never opens an MFA prompt. A valid session is
    refreshed and marked healthy; an expired one is surfaced for deliberate
    user reconnection from the Connections screen.
    """
    checks = (
        ("amazon", ROOT / ".openclaw/amazon-logistics-storage-state.json",
         ["node", str(ROOT / "scripts/amazon_payments_session_check.mjs")]),
        ("pave", ROOT / ".openclaw/pave-storage-state.json",
         ["node", str(ROOT / "scripts/pave_login.mjs"), "--check"]),
    )
    for connection_id, state_path, command in checks:
        if not state_path.exists():
            continue
        checked_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        try:
            subprocess.run(
                command, cwd=str(ROOT), stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL, timeout=90, check=True)
            connections_registry.set_connection_state(
                tenant, connection_id, status="healthy",
                lastCheckedAt=checked_at, lastSuccessAt=checked_at,
                lastError=None)
        except (subprocess.SubprocessError, OSError):
            if connection_id == "pave":
                try:
                    environment = os.environ.copy()
                    environment["PAVE_USERNAME"] = get_secret(
                        tenant, "pave", "production-username")
                    environment["PAVE_PASSWORD"] = get_secret(
                        tenant, "pave", "production-password")
                    subprocess.run(
                        ["node", str(ROOT / "scripts/pave_login.mjs"), "--headless"],
                        cwd=str(ROOT), env=environment, stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL, timeout=90, check=True)
                    refreshed_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
                    connections_registry.set_connection_state(
                        tenant, connection_id, status="healthy",
                        lastCheckedAt=refreshed_at, lastSuccessAt=refreshed_at,
                        lastError=None)
                    continue
                except (KeyError, subprocess.SubprocessError, OSError):
                    pass
            connections_registry.set_connection_state(
                tenant, connection_id, status="needs_reauth",
                lastCheckedAt=checked_at,
                lastError=f'{connections_registry.BY_ID[connection_id]["displayName"]} saved session requires sign-in.')

    # API-provider checks run after browser sessions so a slow upstream API can
    # never prevent Amazon/PAVE state from being restored after startup.
    checked_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for environment in ("development", "production"):
        if (has_secret(tenant, "digits", f"{environment}-client-id")
                and has_secret(tenant, "digits", f"{environment}-client-secret")):
            try:
                from scripts.digits_api_client import list_entities
                list_entities(tenant, environment)
                connections_registry.set_connection_state(
                    tenant, "digits_api", status="healthy",
                    lastCheckedAt=checked_at, lastSuccessAt=checked_at,
                    lastError=None, environment=environment,
                    secretReference=connections_registry.secret_reference(
                        tenant, "digits_api", "credentials"),
                    configuredFields=["clientId", "clientSecret"])
            except Exception:
                connections_registry.set_connection_state(
                    tenant, "digits_api", status="degraded",
                    lastCheckedAt=checked_at,
                    lastError="Digits API credentials require attention.",
                    environment=environment,
                    secretReference=connections_registry.secret_reference(
                        tenant, "digits_api", "credentials"),
                    configuredFields=["clientId", "clientSecret"])
            break

    if all(has_secret(tenant, 'email-imap', name) for name in (
            'production-host', 'production-port', 'production-username', 'production-app-password')):
        try:
            sync_imap_reports(tenant)
        except Exception:
            checked_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
            connections_registry.set_connection_state(
                tenant, 'email_imap', status='degraded', lastCheckedAt=checked_at,
                lastError='IMAP credentials require attention.',
                secretReference=connections_registry.secret_reference(
                    tenant, 'email_imap', 'credentials'),
                configuredFields=['host', 'port', 'username', 'appPassword'], environment='production')


def _due(value, interval):
    if not value:
        return True
    try:
        return datetime.now(timezone.utc) - datetime.fromisoformat(value) >= interval
    except (TypeError, ValueError):
        return True


def sync_pave_export(tenant="jecs"):
    completed = subprocess.run(
        ["node", str(ROOT / "scripts/pave_export_download.mjs")],
        cwd=str(ROOT), capture_output=True, text=True, timeout=180, check=True)
    result = json.loads(completed.stdout.strip().splitlines()[-1])
    path = Path(result["path"])
    payload = path.read_bytes()
    preview = parse_pave_export(payload, result["filename"])
    if not preview.get("ok"):
        raise RuntimeError(preview.get("error") or "PAVE export parse failed")
    record = tenant_store.put_file(tenant, "pave", result["filename"], payload, "scheduled-pave-sync")
    record = tenant_store.update_file(
        tenant, record["id"], status="parsed", provider="pave",
        periodKey=preview.get("periodKey"), parseSummary={
            **preview["summary"], "provider": "pave", "providerLabel": "PAVE Fleet Dashboard"})
    tenant_store.supersede_others(tenant, "pave", record["id"], record.get("periodKey"))
    now = datetime.now(timezone.utc)
    record = tenant_store.update_file(
        tenant, record["id"], status="confirmed", confirmedAt=now.isoformat(timespec="seconds"))
    connections_registry.set_connection_state(
        tenant, "pave", status="healthy", lastSuccessAt=record["confirmedAt"],
        lastCheckedAt=record["confirmedAt"], lastAutomatedSyncAt=record["confirmedAt"],
        nextRunAt=(now + timedelta(days=1)).isoformat(timespec="seconds"), lastError=None,
        lastSyncSummary={"rows": preview["summary"]["rowsRead"], "uniqueVins": preview["summary"]["uniqueVins"]})
    return preview["summary"]


def sync_imap_reports(tenant="jecs"):
    host = get_secret(tenant, 'email-imap', 'production-host')
    port = int(get_secret(tenant, 'email-imap', 'production-port'))
    username = get_secret(tenant, 'email-imap', 'production-username')
    password = get_secret(tenant, 'email-imap', 'production-app-password')
    stored = 0
    with imaplib.IMAP4_SSL(host, port, ssl_context=ssl.create_default_context(), timeout=30) as client:
        client.login(username, password)
        status, _ = client.select('INBOX', readonly=True)
        if status != 'OK':
            raise RuntimeError('read-only INBOX access unavailable')
        since = (datetime.now(timezone.utc) - timedelta(days=2)).strftime('%d-%b-%Y')
        status, data = client.uid('search', None, 'SINCE', since)
        if status != 'OK':
            raise RuntimeError('IMAP search failed')
        for uid in (data[0] or b'').split():
            status, fetched = client.uid('fetch', uid, '(BODY.PEEK[])')
            if status != 'OK':
                continue
            raw = next((part[1] for part in fetched if isinstance(part, tuple)), b'')
            message = email.message_from_bytes(raw, policy=policy.default)
            for part in message.walk():
                filename = part.get_filename()
                if not filename:
                    continue
                try:
                    filename = str(make_header(decode_header(filename)))
                except Exception:
                    filename = str(filename)
                if Path(filename).suffix.lower() not in ('.csv', '.xlsx', '.xlsm', '.pdf'):
                    continue
                content = part.get_payload(decode=True) or b''
                if content:
                    tenant_store.put_file(tenant, 'email_reports', filename, content, 'scheduled-imap-sync')
                    stored += 1
        client.logout()
    now = datetime.now(timezone.utc)
    stamp = now.isoformat(timespec='seconds')
    connections_registry.set_connection_state(
        tenant, 'email_imap', status='healthy', lastSuccessAt=stamp,
        lastCheckedAt=stamp, lastAutomatedSyncAt=stamp,
        nextRunAt=(now + timedelta(minutes=15)).isoformat(timespec='seconds'), lastError=None,
        lastSyncSummary={'attachmentsArchived': stored})
    return stored


_CONNECTION_REFRESH_LOCK = threading.Lock()
_CONNECTION_REFRESH_JOBS = {}


def refresh_all_connection_data(tenant="jecs"):
    """Refresh every configured source and retain the last good data on failure."""
    started = datetime.now(timezone.utc).isoformat(timespec="seconds")
    _CONNECTION_REFRESH_JOBS[tenant] = {
        "status": "running", "startedAt": started, "finishedAt": None, "sources": []}
    results = []
    with _CONNECTION_REFRESH_LOCK:
        try:
            reconcile_saved_browser_sessions(tenant)
            current = {item["id"]: item for item in
                       connections_registry.list_connections(tenant, _latest_upload_index(tenant))}

            pave = current.get("pave", {})
            if pave.get("configured"):
                try:
                    summary = sync_pave_export(tenant)
                    results.append({"id": "pave", "status": "healthy",
                                    "message": f'{summary["rowsRead"]} rows; {summary["uniqueVins"]} VINs'})
                except Exception:
                    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
                    connections_registry.set_connection_state(
                        tenant, "pave", status="degraded", lastCheckedAt=stamp,
                        lastError="Manual PAVE refresh failed; saved data remains available.")
                    results.append({"id": "pave", "status": "degraded",
                                    "message": "Refresh failed; retained saved data"})

            imap = current.get("email_imap", {})
            if imap.get("configured"):
                try:
                    count = sync_imap_reports(tenant)
                    results.append({"id": "email_imap", "status": "healthy",
                                    "message": f"{count} report attachments checked"})
                except Exception:
                    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
                    connections_registry.set_connection_state(
                        tenant, "email_imap", status="degraded", lastCheckedAt=stamp,
                        lastError="Manual IMAP refresh failed; saved data remains available.")
                    results.append({"id": "email_imap", "status": "degraded",
                                    "message": "Refresh failed; retained saved data"})

            final_connections = connections_registry.list_connections(
                tenant, _latest_upload_index(tenant))
            completed_ids = {item["id"] for item in results}
            for item in final_connections:
                if item["id"] not in completed_ids:
                    results.append({"id": item["id"], "status": item["status"],
                                    "message": "Connection checked" if item.get("configured") else "Not configured"})
            status = "completed"
        except Exception:
            status = "failed"
            results.append({"id": "system", "status": "degraded",
                            "message": "Refresh could not be completed; saved data remains available"})
    _CONNECTION_REFRESH_JOBS[tenant] = {
        "status": status, "startedAt": started,
        "finishedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sources": results,
    }


def run_connection_scheduler(tenant="jecs"):
    reconcile_saved_browser_sessions(tenant)
    while True:
        connections = {item['id']: item for item in connections_registry.list_connections(tenant, _latest_upload_index(tenant))}
        pave = connections.get('pave', {})
        if pave.get('status') == 'healthy' and _due(pave.get('lastAutomatedSyncAt'), timedelta(days=1)):
            try:
                with _CONNECTION_REFRESH_LOCK:
                    sync_pave_export(tenant)
            except Exception:
                now = datetime.now(timezone.utc).isoformat(timespec='seconds')
                connections_registry.set_connection_state(
                    tenant, 'pave', status='degraded', lastCheckedAt=now,
                    lastError='Scheduled PAVE export failed; saved data remains available.')
        imap = connections.get('email_imap', {})
        if imap.get('configured') and _due(imap.get('lastAutomatedSyncAt'), timedelta(minutes=15)):
            try:
                with _CONNECTION_REFRESH_LOCK:
                    sync_imap_reports(tenant)
            except Exception:
                now = datetime.now(timezone.utc).isoformat(timespec='seconds')
                connections_registry.set_connection_state(
                    tenant, 'email_imap', status='degraded', lastCheckedAt=now,
                    lastError='Scheduled IMAP collection failed; verify the mailbox credentials.')
        threading.Event().wait(60)


def parse_multipart(body, content_type):
    """Minimal multipart/form-data reader for a single file field."""
    CR = chr(13).encode()
    LF = chr(10).encode()
    CRLF = CR + LF
    match = re.search(r'boundary=(?:"([^"]+)"|([^;]+))', content_type or "")
    if not match:
        raise ValueError("missing multipart boundary")
    boundary = (match.group(1) or match.group(2)).strip().encode()
    parts = body.split(b"--" + boundary)
    fields, files = {}, {}
    for part in parts:
        part = part.strip(CRLF)
        if not part or part == b"--":
            continue
        head, _, data = part.partition(CRLF + CRLF)
        headers = head.decode("utf-8", "replace")
        name = re.search(r'name="([^"]*)"', headers)
        filename = re.search(r'filename="([^"]*)"', headers)
        if not name:
            continue
        data = data.rstrip(CRLF)
        if filename and filename.group(1):
            files[name.group(1)] = (filename.group(1), data)
        else:
            fields[name.group(1)] = data.decode("utf-8", "replace")
    return fields, files


class JecsAPIHandler(BaseHTTPRequestHandler):
    """HTTP request handler for JECS API."""
    
    def log_message(self, format, *args):
        """Suppress default logging."""
        pass
    
    def end_headers(self):
        """Add CORS headers to all responses."""
        self.send_header('Access-Control-Allow-Origin', 'http://localhost:3000')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, PUT, DELETE, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')
        self.send_header('Access-Control-Allow-Credentials', 'true')
        super().end_headers()
    
    def do_OPTIONS(self):
        """Handle OPTIONS requests for CORS preflight."""
        self.send_response(200)
        self.end_headers()
    
    def do_GET(self):
        """Handle GET requests."""
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)
        
        try:
            # Daily Entry Form endpoints
            if path == '/api/drivers':
                self.handle_get_drivers()
            elif path == '/api/vans':
                self.handle_get_vans()
            elif path == '/api/daily-entries':
                self.handle_get_daily_entries(query)
            elif path == '/api/daily-entries/history':
                self.handle_get_history(query)
            
            # Driver Performance endpoints
            elif path == '/api/driver-performance':
                self.handle_get_driver_performance(query)
            
            # Fleet Cost Optimization endpoints
            elif path == '/api/fleet-optimization':
                self.handle_get_fleet_optimization(query)
            elif path == '/api/fleet-costs':
                self.handle_get_fleet_costs(query)
            elif path == '/api/fleet-costs/records':
                payload = build_fleet_cost_reconciliation(_tenant_from_headers(self.headers))
                self.send_paginated(payload.get('charges', []))
            elif path == '/api/fleet-costs/summary':
                self.handle_get_fleet_cost_summary(query)
            elif path == '/api/fleet-costs/fuel-analysis':
                self.handle_get_fuel_cost_analysis(query)
            elif path == '/api/fleet-costs/maintenance-analysis':
                self.handle_get_maintenance_cost_analysis(query)
            elif path == '/api/fleet-costs/trends':
                self.handle_get_fleet_cost_trends(query)
            elif path in ('/api/fleet-costs/budgets', '/api/fleet-costs/forecasts'):
                self.send_json([])
            
            # Dispute Detection endpoints
            elif path == '/api/disputes':
                self.handle_get_disputes(query)
            elif path == '/api/disputes/candidates':
                self.handle_get_dispute_candidates(query)
            
            # PAVE endpoints
            elif path == '/api/pave/vehicles':
                self.handle_get_pave_vehicles(query)
            elif path == '/api/fleet-compliance':
                self.handle_get_fleet_compliance()
            elif path == '/api/connections':
                self.send_json(build_connections_payload(_tenant_from_headers(self.headers)))
            elif path == '/api/connections/refresh':
                tenant = _tenant_from_headers(self.headers)
                self.send_json(_CONNECTION_REFRESH_JOBS.get(tenant, {
                    'status': 'idle', 'startedAt': None, 'finishedAt': None, 'sources': []}))
            elif path == '/api/dashboard/operations':
                self.send_json(build_dashboard_operations_payload(
                    _tenant_from_headers(self.headers)))
            elif path == '/api/uploads':
                tenant = _tenant_from_headers(self.headers)
                source = (query.get('source') or [None])[0]
                self.send_json({'tenant': tenant, 'uploads': tenant_store.list_files(tenant, source)})
            elif path == '/api/vendor-rules':
                tenant = _tenant_from_headers(self.headers)
                self.send_json({'tenant': tenant, 'rules': load_vendor_rules(tenant),
                                'providers': PROVIDER_LABELS})
            elif path == '/api/modules':
                self.send_json(build_reimbursement_review_payload())
            elif path == '/api/context':
                tenant = _tenant_from_headers(self.headers)
                support = _LOCAL_SUPPORT_SESSIONS.get(self.headers.get('x-support-session'))
                tenant_record = next((item for item in _local_tenants() if item['slug'] == tenant), None)
                user = ({'id': support['identitySubject'], 'email': support['email'], 'role': support['role'],
                         'tenantRole': support['role'], 'isPlatformAdmin': False}
                        if support and support.get('tenantSlug') == tenant else
                        {'id': 'dev-user', 'email': 'dev@example.com', 'role': 'platform_admin',
                         'tenantRole': 'owner', 'isPlatformAdmin': True})
                self.send_json({
                    'tenant': {'id': tenant, 'name': (tenant_record or {}).get('displayName', tenant)},
                    'user': user,
                    'permissions': ['module.read', 'integration.manage', 'member.manage',
                                    'feature.manage', 'impersonation.manage', 'tenant.manage',
                                    'tenant.provision', 'tenant.impersonate'],
                    'features': _public_features(tenant),
                    'modules': [],
                    'impersonation': ({'active': True, 'actorEmail': 'dev@example.com',
                                       'targetEmail': support['email'], 'targetRole': support['role'],
                                       'reason': support['reason'], 'expiresAt': support['expiresAt']}
                                      if support and support.get('tenantSlug') == tenant else None),
                })
            elif path == '/api/assistant/status':
                self.send_json(assistant_service.status(_tenant_from_headers(self.headers)))
            elif path == '/api/assistant/config':
                tenant = _tenant_from_headers(self.headers)
                credential = assistant_service.credentials(tenant)
                self.send_json({
                    'configured': bool(credential), 'provider': 'OpenAI',
                    'model': (credential or {}).get('model') or 'gpt-5-mini',
                    'organizationConfigured': bool((credential or {}).get('organization')),
                    'projectConfigured': bool((credential or {}).get('project')),
                    'storage': 'Local managed secret store (AWS Secrets Manager in production)',
                    'writeOnly': True,
                    'lastFour': (credential or {}).get('apiKey', '')[-4:] or None,
                })
            elif path == '/api/features':
                self.send_json({'features': _public_features(_tenant_from_headers(self.headers)),
                                'canManage': True})
            elif path == '/api/members':
                self.send_json({'members': _local_members(_tenant_from_headers(self.headers)),
                                'roles': ['owner', 'admin', 'reviewer', 'analyst', 'viewer']})
            elif path == '/api/super-admin/tenants':
                self.send_json({'tenants': _local_tenants(), 'modules': LOCAL_MODULES})
            elif path == '/api/super-admin/members':
                tenant_slug = (query.get('tenantSlug') or [None])[0]
                tenants = [item for item in _local_tenants()
                           if tenant_slug is None or item['slug'] == tenant_slug]
                members = [{**member, 'tenantSlug': tenant['slug'], 'tenantName': tenant['displayName']}
                           for tenant in tenants for member in _local_members(tenant['slug'])]
                self.send_json({'members': members, 'total': len(members)})
            elif path == '/api/super-admin/onboarding-checklist':
                self.send_json({'checklist': {'credentialNotice': 'Credentials use the tenant vault only.'}})
            elif path == '/api/super-admin/onboarding-checklist.pdf':
                self.send_binary((ROOT / 'platform/docs/tenant-onboarding-checklist.pdf').read_bytes(),
                                 'application/pdf', 'tenant-onboarding-checklist.pdf')
            elif re.fullmatch(r'/api/super-admin/tenants/[a-z][a-z0-9-]{2,62}/features', path):
                tenant_slug = path.split('/')[4]
                self.send_json({'tenantSlug': tenant_slug, 'features': _public_features(tenant_slug)})
            elif re.fullmatch(r'/api/super-admin/tenants/[a-z][a-z0-9-]{2,62}/connections', path):
                tenant_slug = path.split('/')[4]
                self.send_json({'tenantSlug': tenant_slug,
                                'connections': build_connections_payload(tenant_slug).get('connections', [])})
            elif re.fullmatch(r'/api/super-admin/tenants/[a-z][a-z0-9-]{2,62}', path):
                tenant_slug = path.split('/')[4]
                tenant = next((item for item in _local_tenants() if item['slug'] == tenant_slug), None)
                self.send_json_status(200 if tenant else 404, {'tenant': tenant} if tenant else {'error': 'tenant not found'})
            elif path.startswith('/api/modules/') and path.endswith('/cases'):
                module_id = path.split('/')[3]
                self.send_json({'cases': [c for c in _load_module_cases() if c['moduleId'] == module_id]})
            
            # Route Monitoring endpoints
            elif path == '/api/route-monitor':
                self.handle_get_live_route_monitor(query)
            elif path == '/api/route-performance':
                self.handle_get_route_monitor(query)
            elif path == '/api/routes':
                self.send_paginated([])

            elif path == '/api/performance/dashboard':
                tenant = _tenant_from_headers(self.headers)
                self.send_json(build_performance_dashboard_payload() if tenant == DEFAULT_TENANT else {
                    'period': None, 'generatedAt': datetime.now(timezone.utc).isoformat(),
                    'source': 'No tenant-scoped scorecard source available', 'needsData': True,
                    'drivers': [], 'history': [],
                    'dspPerformance': {'overallScore': None, 'deliveryScore': 0, 'safetyScore': 0,
                                       'qualityScore': 0, 'driverCount': 0, 'totalDeliveries': 0},
                })

            # Incumbent dashboard compatibility endpoints
            elif path == '/api/weekly-evaluations':
                self.handle_get_weekly_evaluations()
            elif path == '/api/time-attendance/exceptions':
                self.handle_get_time_attendance_exceptions()
            
            # Auth endpoints (mock for development)
            elif path == '/auth/me' or path == '/api/auth/me':
                self.handle_get_current_user()
            
            # Payroll Reconciliation endpoints
            elif path == '/api/payroll':
                self.handle_get_payroll(query)
            elif path == '/api/payroll/discrepancies':
                self.handle_get_payroll_discrepancies(query)
            elif path == '/api/payroll/periods':
                self.send_paginated([])
            
            # Dashboard files
            elif path == '/amazon-dsp-kpi-dashboard.html':
                self.handle_dashboard_html()
            elif path == '/daily-entry-form.html':
                self.handle_form_html()
            elif path.startswith('/data/dashboards/'):
                self.handle_dashboard_file()
            elif path == '/' or path == '/daily-entry':
                self.handle_redirect_to_dashboard()
            else:
                self.send_error(404, f"Not found: {path}")
        except Exception as e:
            self.send_error(500, f"Error: {str(e)}")
    
    def do_POST(self):
        """Handle POST requests."""
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)
        
        try:
            if path.startswith('/api/uploads/') and path.endswith('/reparse'):
                self.handle_post_upload_reparse(path.split('/')[3])
            elif path.startswith('/api/uploads/') and path.endswith('/confirm'):
                self.handle_post_upload_confirm(path.split('/')[3])
            elif path.startswith('/api/uploads/'):
                self.handle_post_upload(path.split('/')[3])
            elif path == '/api/vendor-rules':
                self.handle_post_vendor_rules()
            elif path.startswith('/api/connections/') and path.endswith('/test'):
                self.handle_post_connection_test(path.split('/')[3])
            elif path == '/api/connections/refresh':
                self.handle_post_connections_refresh()
            elif path == '/api/assistant/chat':
                self.handle_post_assistant_chat()
            elif path == '/api/assistant/config/test':
                self.handle_post_assistant_config_test()
            elif path == '/api/members/invitations':
                self.handle_post_member_invitation()
            elif re.fullmatch(r'/api/members/[^/]+/resend-invitation', path):
                self.handle_post_member_invitation_resend(urllib.parse.unquote(path.split('/')[3]))
            elif path == '/api/super-admin/tenants':
                self.handle_post_super_admin_tenant()
            elif path == '/api/super-admin/impersonate':
                self.handle_post_super_admin_impersonation()
            elif path == '/api/support/impersonation/end':
                _LOCAL_SUPPORT_SESSIONS.pop(self.headers.get('x-support-session'), None)
                self.send_json_status(200, {'ended': True})
            elif path.startswith('/api/connections/') and path.endswith('/reconnect'):
                self.handle_post_connection_reconnect(path.split('/')[3])
            elif path == '/api/daily-entries':
                self.handle_post_daily_entries()
            elif path == '/api/daily-entries/revert':
                self.handle_post_revert(query)
            # Auth endpoints (mock for development)
            elif path == '/auth/login' or path == '/api/auth/login':
                self.handle_post_login()
            elif path == '/auth/logout' or path == '/api/auth/logout':
                self.handle_post_logout()
            elif path == '/auth/refresh' or path == '/api/auth/refresh':
                self.handle_post_refresh()
            else:
                self.send_error(404, f"Not found: {path}")
        except Exception as e:
            self.send_error(500, f"Error: {str(e)}")

    def do_PUT(self):
        """Handle local-development configuration updates."""
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        try:
            if path.startswith('/api/connections/') and path.endswith('/credentials'):
                self.handle_put_connection_credentials(path.split('/')[3])
            elif path == '/api/assistant/config':
                self.handle_put_assistant_config()
            elif path.startswith('/api/features/'):
                self.handle_put_feature(path.split('/')[3])
            elif path.startswith('/api/members/'):
                self.handle_put_member(urllib.parse.unquote(path.split('/')[3]))
            elif re.fullmatch(r'/api/super-admin/tenants/[a-z][a-z0-9-]{2,62}/status', path):
                self.handle_put_super_admin_status(path.split('/')[4])
            elif re.fullmatch(r'/api/super-admin/tenants/[a-z][a-z0-9-]{2,62}/features/[a-z][a-z0-9_]{2,63}', path):
                parts = path.split('/')
                tenant_slug, feature_id = parts[4], parts[6]
                body = json.loads(self._read_body() or b'{}')
                if not isinstance(body.get('enabled'), bool):
                    raise ValueError('enabled must be a boolean')
                self.send_json_status(200, {'tenantSlug': tenant_slug,
                                            'feature': _set_feature_override(tenant_slug, feature_id, body['enabled'])})
            else:
                self.send_error(404, f"Not found: {path}")
        except KeyError as error:
            self.send_json_status(404, {'error': str(error)})
        except ValueError as error:
            self.send_json_status(400, {'error': str(error)})
        except Exception as error:
            self.send_json_status(500, {'error': str(error)})
    
    def do_OPTIONS(self):
        """Handle OPTIONS for CORS."""
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    # ========== Daily Entry Form Handlers ==========
    
    def handle_get_drivers(self):
        """Return list of all drivers."""
        tenant = _tenant_from_headers(self.headers)
        if tenant != DEFAULT_TENANT:
            self.send_json({
                'data': [], 'needsData': True,
                'message': 'No tenant-scoped driver roster has been ingested',
                'pagination': {'page': 1, 'limit': 0, 'total': 0, 'totalPages': 0},
            })
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, name, status FROM drivers ORDER BY name")
        drivers = [dict(row) for row in cursor.fetchall()]
        conn.close()
        self.send_json(drivers)
    
    def handle_get_vans(self):
        """Return list of all vans."""
        tenant = _tenant_from_headers(self.headers)
        if tenant != DEFAULT_TENANT:
            self.send_json({
                'data': [], 'needsData': True,
                'message': 'No tenant-scoped vehicle roster has been ingested',
            })
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, van_number, vin, make, model, year, status, ownership FROM vans ORDER BY van_number")
        vans = [dict(row) for row in cursor.fetchall()]
        conn.close()
        self.send_json(vans)
    
    def handle_get_daily_entries(self, query):
        """Return daily entries for a specific date."""
        date = query.get('date', [None])[0]
        if not date:
            self.send_error(400, "Missing date parameter")
            return
        
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                de.id, de.driver_id, d.name as driver_name,
                de.van_id, v.van_number,
                de.route_code, de.date, de.stops, de.packages,
                de.status, de.start_time, de.end_time, de.notes
            FROM daily_route_entries de
            LEFT JOIN drivers d ON de.driver_id = d.id
            LEFT JOIN vans v ON de.van_id = v.id
            WHERE de.date = ?
            ORDER BY d.name, de.route_code
        """, (date,))
        entries = [dict(row) for row in cursor.fetchall()]
        conn.close()
        self.send_json(entries)
    
    def handle_get_history(self, query):
        """Return history for a specific date."""
        date = query.get('date', [None])[0]
        if not date:
            self.send_error(400, "Missing date parameter")
            return
        
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                h.id, h.entry_id, h.driver_id, h.driver_name,
                h.van_id, h.van_number,
                h.route_code, h.date, h.stops, h.packages,
                h.status, h.start_time, h.end_time, h.notes,
                h.changed_by, h.changed_at, h.change_type
            FROM daily_entries_history h
            WHERE h.date = ?
            ORDER BY h.changed_at DESC
        """, (date,))
        history = [dict(row) for row in cursor.fetchall()]
        conn.close()
        self.send_json(history)
    
    def handle_post_daily_entries(self):
        """Save daily entries (auto-save)."""
        content_length = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(content_length)
        data = json.loads(body)
        
        entries = data.get('entries', [])
        date = data.get('date')
        
        if not date:
            self.send_error(400, "Missing date")
            return
        
        conn = get_connection()
        cursor = conn.cursor()
        
        history_entries = []
        
        for entry in entries:
            entry_id = entry.get('id')
            driver_id = entry.get('driver_id')
            van_id = entry.get('van_id')
            route_code = entry.get('route_code')
            stops = entry.get('stops')
            packages = entry.get('packages')
            status = entry.get('status')
            start_time = entry.get('start_time')
            end_time = entry.get('end_time')
            notes = entry.get('notes')
            deleted = entry.get('deleted', False)
            
            if deleted and entry_id and entry_id != 'new':
                cursor.execute("""
                    SELECT de.*, d.name as driver_name, v.van_number 
                    FROM daily_route_entries de
                    LEFT JOIN drivers d ON de.driver_id = d.id
                    LEFT JOIN vans v ON de.van_id = v.id
                    WHERE de.id = ?
                """, (entry_id,))
                old_entry = cursor.fetchone()
                if old_entry:
                    history_entries.append({
                        'entry_id': entry_id,
                        'driver_id': old_entry['driver_id'],
                        'driver_name': old_entry['driver_name'],
                        'van_id': old_entry['van_id'],
                        'van_number': old_entry['van_number'],
                        'route_code': old_entry['route_code'],
                        'date': old_entry['date'],
                        'stops': old_entry['stops'],
                        'packages': old_entry['packages'],
                        'status': old_entry['status'],
                        'start_time': old_entry['start_time'],
                        'end_time': old_entry['end_time'],
                        'notes': old_entry['notes'],
                        'changed_by': 'system',
                        'change_type': 'delete'
                    })
                cursor.execute("DELETE FROM daily_route_entries WHERE id = ?", (entry_id,))
            elif entry_id and entry_id != 'new':
                cursor.execute("""
                    SELECT de.*, d.name as driver_name, v.van_number 
                    FROM daily_route_entries de
                    LEFT JOIN drivers d ON de.driver_id = d.id
                    LEFT JOIN vans v ON de.van_id = v.id
                    WHERE de.id = ?
                """, (entry_id,))
                old_entry = cursor.fetchone()
                if old_entry:
                    history_entries.append({
                        'entry_id': entry_id,
                        'driver_id': old_entry['driver_id'],
                        'driver_name': old_entry['driver_name'],
                        'van_id': old_entry['van_id'],
                        'van_number': old_entry['van_number'],
                        'route_code': old_entry['route_code'],
                        'date': old_entry['date'],
                        'stops': old_entry['stops'],
                        'packages': old_entry['packages'],
                        'status': old_entry['status'],
                        'start_time': old_entry['start_time'],
                        'end_time': old_entry['end_time'],
                        'notes': old_entry['notes'],
                        'changed_by': 'system',
                        'change_type': 'update'
                    })
                cursor.execute("""
                    UPDATE daily_route_entries SET
                        driver_id = ?, van_id = ?, route_code = ?, stops = ?, packages = ?,
                        status = ?, start_time = ?, end_time = ?, notes = ?
                    WHERE id = ?
                """, (
                    driver_id, van_id, route_code, stops, packages,
                    status, start_time, end_time, notes, entry_id
                ))
            else:
                cursor.execute("""
                    INSERT INTO daily_route_entries (
                        driver_id, van_id, route_code, date, stops, packages,
                        status, start_time, end_time, notes, entered_by
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    driver_id, van_id, route_code, date, stops, packages,
                    status, start_time, end_time, notes, 'system'
                ))
                entry_id = cursor.lastrowid
                history_entries.append({
                    'entry_id': entry_id,
                    'driver_id': driver_id,
                    'driver_name': '',
                    'van_id': van_id,
                    'van_number': '',
                    'route_code': route_code,
                    'date': date,
                    'stops': stops,
                    'packages': packages,
                    'status': status,
                    'start_time': start_time,
                    'end_time': end_time,
                    'notes': notes,
                    'changed_by': 'system',
                    'change_type': 'create'
                })
        
        for history in history_entries:
            cursor.execute("""
                INSERT INTO daily_entries_history (
                    entry_id, driver_id, driver_name, van_id, van_number,
                    route_code, date, stops, packages, status, start_time, end_time,
                    notes, changed_by, change_type
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                history['entry_id'],
                history['driver_id'],
                history.get('driver_name', ''),
                history['van_id'],
                history.get('van_number', ''),
                history['route_code'],
                history['date'],
                history['stops'],
                history['packages'],
                history['status'],
                history['start_time'],
                history['end_time'],
                history['notes'],
                history['changed_by'],
                history['change_type']
            ))
        
        conn.commit()
        
        cursor.execute("""
            SELECT 
                de.id, de.driver_id, d.name as driver_name,
                de.van_id, v.van_number,
                de.route_code, de.date, de.stops, de.packages,
                de.status, de.start_time, de.end_time, de.notes
            FROM daily_route_entries de
            LEFT JOIN drivers d ON de.driver_id = d.id
            LEFT JOIN vans v ON de.van_id = v.id
            WHERE de.date = ?
            ORDER BY d.name, de.route_code
        """, (date,))
        updated_entries = [dict(row) for row in cursor.fetchall()]
        conn.close()
        
        self.send_json({'entries': updated_entries, 'saved': True})
    
    def handle_post_revert(self, query):
        """Revert to a specific history entry."""
        history_id = query.get('history_id', [None])[0]
        if not history_id:
            self.send_error(400, "Missing history_id parameter")
            return
        
        conn = get_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT * FROM daily_entries_history WHERE id = ?", (history_id,))
        history = cursor.fetchone()
        if not history:
            self.send_error(404, "History entry not found")
            return
        
        if history['change_type'] == 'delete':
            cursor.execute("""
                INSERT INTO daily_route_entries (
                    id, driver_id, van_id, route_code, date, stops, packages,
                    status, start_time, end_time, notes, entered_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                history['entry_id'], history['driver_id'], history['van_id'],
                history['route_code'], history['date'], history['stops'], history['packages'],
                history['status'], history['start_time'], history['end_time'],
                history['notes'], 'reverted'
            ))
        else:
            cursor.execute("""
                UPDATE daily_route_entries SET
                    driver_id = ?, van_id = ?, route_code = ?, stops = ?, packages = ?,
                    status = ?, start_time = ?, end_time = ?, notes = ?
                WHERE id = ?
            """, (
                history['driver_id'], history['van_id'], history['route_code'],
                history['stops'], history['packages'], history['status'],
                history['start_time'], history['end_time'], history['notes'],
                history['entry_id']
            ))
        
        cursor.execute("""
            INSERT INTO daily_entries_history (
                entry_id, driver_id, driver_name, van_id, van_number,
                route_code, date, stops, packages, status, start_time, end_time,
                notes, changed_by, change_type
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            history['entry_id'], history['driver_id'], history.get('driver_name', ''),
            history['van_id'], history.get('van_number', ''), history['route_code'],
            history['date'], history['stops'], history['packages'], history['status'],
            history['start_time'], history['end_time'], history['notes'],
            'system', 'revert'
        ))
        
        conn.commit()
        conn.close()
        
        self.send_json({'reverted': True})

    # ========== Driver Performance Handlers ==========
    
    def handle_get_driver_performance(self, query):
        """Return driver performance summary."""
        week = query.get('week', ['2026-wk37'])[0]
        
        conn = get_connection()
        cursor = conn.cursor()
        
        # Get driver performance from amazon_routes (weekly aggregates)
        cursor.execute("""
            SELECT 
                ar.driver_id,
                d.name as driver_name,
                ar.overall_score,
                ar.pod,
                ar.cdf,
                ar.dsb,
                ar.packages,
                ar.stops,
                ar.date
            FROM amazon_routes ar
            JOIN drivers d ON ar.driver_id = d.id
            WHERE ar.is_weekly_aggregate = 1
            ORDER BY ar.overall_score DESC NULLS LAST
            LIMIT 50
        """)
        routes = cursor.fetchall()
        
        # Build performance summary
        performance = []
        for route in routes:
            performance.append({
                'driver_id': route['driver_id'],
                'driver_name': route['driver_name'],
                'dcr': route['overall_score'] or 0,
                'pod': route['pod'] or 0,
                'cdf_negatives': route['cdf'] or 0,
                'dsb_defects': route['dsb'] or 0,
                'safety_events': 0,  # Will populate from safety table if available
                'packages_delivered': route['packages'] or 0,
                'score': route['overall_score'] or 0
            })
        
        # Sort by score (descending)
        performance.sort(key=lambda x: x['score'], reverse=True)
        
        conn.close()
        self.send_json(performance)
    
    def _calculate_driver_score(self, scorecard):
        """Calculate a composite score for a driver."""
        dcr = scorecard['dcr'] or 0
        pod = scorecard['pod'] or 0
        cdf = scorecard['cdf_negatives'] or 0
        dsb = scorecard['dsb_defects'] or 0
        safety = scorecard['safety_events'] or 0
        
        # Weighted score (DCR and POD are most important)
        score = (dcr * 0.4) + (pod * 0.4) - (cdf * 50) - (dsb * 100) - (safety * 200)
        return round(score, 2)

    # ========== Fleet Cost Optimization Handlers ==========
    
    def handle_get_fleet_optimization(self, query):
        """Return fleet cost optimization analysis."""
        conn = get_connection()
        cursor = conn.cursor()
        
        # Get fleet data
        cursor.execute("SELECT id, van_number, vin, make, model, year, status, ownership, notes, updated_at FROM vans")
        vans = [dict(row) for row in cursor.fetchall()]
        
        # Get ownership breakdown
        ownership = {'AMAZON_OWNED': 0, 'AMAZON_RENTAL': 0, 'RENTAL': 0, 'LEASE': 0}
        for van in vans:
            ownership[van['ownership']] = ownership.get(van['ownership'], 0) + 1
        
        # Get operational status
        operational = sum(1 for v in vans if v['status'] == 'OPERATIONAL')
        grounded = len(vans) - operational
        
        result = {
            'total_vans': len(vans),
            'operational': operational,
            'grounded': grounded,
            'ownership': ownership,
            'vans': vans
        }
        
        conn.close()
        self.send_json(result)

    def handle_get_fleet_costs(self, query):
        """Return the June-August rental cost vs Amazon reimbursement reconciliation."""
        self.send_json(build_fleet_cost_reconciliation(_tenant_from_headers(self.headers)))

    # ========== Dispute Detection Handlers ==========
    
    def handle_get_disputes(self, query):
        """Return existing disputes."""
        conn = get_connection()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT 
                week, driver_id, metric, reason, status, submitted_at, outcome, tba_ids, evidence_sources, priority, appeal_details
            FROM disputes
            ORDER BY week DESC, submitted_at DESC
        """)
        disputes = []
        for row in cursor.fetchall():
            dispute = dict(row)
            # Get driver name
            if dispute['driver_id']:
                cursor.execute("SELECT name FROM drivers WHERE id = ?", (dispute['driver_id'],))
                driver = cursor.fetchone()
                if driver:
                    dispute['driver_name'] = driver['name']
                else:
                    dispute['driver_name'] = 'Unknown'
            else:
                dispute['driver_name'] = 'Unknown'
            disputes.append(dispute)
        
        conn.close()
        self.send_json(disputes)
    
    def handle_get_dispute_candidates(self, query):
        """Return auto-detected dispute candidates."""
        week = query.get('week', ['2026-wk37'])[0]
        
        # Normalize week format: 2026-wk37 -> 2026-W37
        week = week.replace('-wk', '-W')
        
        conn = get_connection()
        cursor = conn.cursor()
        
        # Get driver performance data for the week
        cursor.execute("""
            SELECT driver_id, dcr, pod, cdf_negatives, dsb_defects, safety_events
            FROM driver_performance
            WHERE week = ?
        """, (week,))
        scorecards = cursor.fetchall()
        
        # Detect dispute candidates
        candidates = []
        for sc in scorecards:
            driver_id = sc['driver_id']
            if not driver_id:
                continue
            cursor.execute("SELECT name FROM drivers WHERE id = ?", (driver_id,))
            driver = cursor.fetchone()
            if not driver:
                continue
            
            # Flag DCR below 99%
            dcr = sc.get('dcr')
            if dcr and dcr < 99.0:
                candidates.append({
                    'week': week,
                    'driver_id': driver_id,
                    'driver_name': driver['name'],
                    'metric': 'DCR',
                    'reason': f"DCR below 99% ({dcr}%)",
                    'evidence': f"Driver performance data for {week}",
                    'confidence': 'High'
                })
            
            # Flag safety events
            safety = sc.get('safety_events')
            if safety and safety > 0:
                candidates.append({
                    'week': week,
                    'driver_id': driver_id,
                    'driver_name': driver['name'],
                    'metric': 'Safety',
                    'reason': f"{safety} safety events",
                    'evidence': f"Driver performance data for {week}",
                    'confidence': 'Medium'
                })
            
            # Flag high DSB defects
            dsb = sc.get('dsb_defects')
            if dsb and dsb > 2:
                candidates.append({
                    'week': week,
                    'driver_id': driver_id,
                    'driver_name': driver['name'],
                    'metric': 'DSB',
                    'reason': f"{dsb} DSB defects",
                    'evidence': f"Driver performance data for {week}",
                    'confidence': 'Medium'
                })
            
            # Flag high CDF negatives
            cdf = sc.get('cdf_negatives')
            if cdf and cdf > 10:
                candidates.append({
                    'week': week,
                    'driver_id': driver_id,
                    'driver_name': driver['name'],
                    'metric': 'CDF',
                    'reason': f"{cdf} CDF negatives",
                    'evidence': f"Driver performance data for {week}",
                    'confidence': 'Medium'
                })
        
        conn.close()
        self.send_json(candidates)

    # ========== Route Monitoring Handlers ==========

    def handle_get_live_route_monitor(self, query):
        """Return only same-day Amazon Delivery Execution data; never substitute weekly scorecards."""
        tenant = _tenant_from_headers(self.headers)
        snapshot_path = ROOT / 'data' / 'tenants' / tenant / 'amazon' / 'live-routes' / 'latest.json'
        if not snapshot_path.exists():
            self.send_json({
                'period': None, 'capturedAt': None, 'source': 'Amazon Delivery Execution',
                'live': False, 'stale': True, 'needsData': True, 'needsReauth': False,
                'message': 'No same-day route execution snapshot has been captured. Run npm run amazon:live-routes.',
                'routeCount': 0,
                'summary': {'assigned': 0, 'inProgress': 0, 'completed': 0, 'behind': 0, 'stalled': 0, 'lateDepartures': 0, 'multiRoute': 0},
                'routes': [],
            })
            return
        try:
            payload = json.loads(snapshot_path.read_text())
            captured = datetime.fromisoformat(payload['capturedAt'].replace('Z', '+00:00'))
            eastern_today = datetime.now(ZoneInfo('America/New_York')).date().isoformat()
            payload['stale'] = payload.get('period') != eastern_today or datetime.now(timezone.utc) - captured > timedelta(minutes=15)
            payload['live'] = not payload['stale']
            self.send_json(payload)
        except (OSError, ValueError, KeyError) as error:
            self.send_json({'error': f'Live route snapshot is invalid: {error}'}, status=500)
    
    def handle_get_route_monitor(self, query):
        """Return weekly scorecard route performance history."""
        conn = get_connection()
        cursor = conn.cursor()
        
        requested = (query.get('period') or query.get('date') or [None])[0]
        if not requested:
            cursor.execute("SELECT MAX(date) AS latest FROM amazon_routes")
            requested = cursor.fetchone()['latest']
        cursor.execute("""
            SELECT 
                route_code, driver_id, date, stops, packages, status,
                overall_score, pod, cdf, dsb, is_weekly_aggregate
            FROM amazon_routes
            WHERE date = ?
            ORDER BY route_code
        """, (requested,))
        routes = []
        for row in cursor.fetchall():
            route = dict(row)
            # Get driver name
            if route['driver_id']:
                cursor.execute("SELECT name FROM drivers WHERE id = ?", (route['driver_id'],))
                driver = cursor.fetchone()
                if driver:
                    route['driver_name'] = driver['name']
                else:
                    route['driver_name'] = 'Unknown'
            else:
                route['driver_name'] = 'Unknown'
            routes.append(route)
        
        conn.close()
        self.send_json({'period': requested, 'routes': routes, 'routeCount': len(routes),
                        'source': 'Amazon scorecard route aggregates'})

    # ========== Payroll Reconciliation Handlers ==========
    
    def handle_get_payroll(self, query):
        """Return payroll reconciliation summary."""
        conn = get_connection()
        cursor = conn.cursor()
        
        # Get ADP timecards
        cursor.execute("SELECT * FROM adp_timecards")
        timecards = [dict(row) for row in cursor.fetchall()]
        
        # Get Amazon routes
        cursor.execute("SELECT * FROM amazon_routes LIMIT 10")
        routes = [dict(row) for row in cursor.fetchall()]
        
        summaries = sorted((ROOT / 'data' / 'adp').glob('*/summary.json'))
        source_summary = {}
        if summaries:
            try:
                source_summary = json.loads(summaries[-1].read_text())
            except (ValueError, OSError):
                source_summary = {}
        result = {
            'timecards': timecards,
            'routes': routes,
            'summary': source_summary,
            'source': 'ADP Workforce Now API',
            'sourceStatus': source_summary.get('payrollOutputStatus') or 'Timecard data available',
        }
        
        conn.close()
        self.send_json(result)
    
    def handle_get_payroll_discrepancies(self, query):
        """Return payroll discrepancies."""
        conn = get_connection()
        cursor = conn.cursor()
        
        # Cross-reference ADP timecards with Amazon routes
        # Use the correct column name: duration_hours, not hat.hours or at.hours
        cursor.execute("""
            SELECT 
                at.driver_id, d.name as driver_name, at.date, at.duration_hours as adp_hours,
                ar.stops, ar.packages, ar.date as route_date
            FROM adp_timecards at
            JOIN drivers d ON at.driver_id = d.id
            LEFT JOIN amazon_routes ar ON at.driver_id = ar.driver_id AND at.date = ar.date
            WHERE at.duration_hours = 0 AND ar.driver_id IS NOT NULL
        """)
        discrepancies = [dict(row) for row in cursor.fetchall()]
        
        conn.close()
        self.send_json(discrepancies)

    # ========== PAVE Handlers ==========

    def handle_get_fleet_compliance(self):
        """Return the reconciled fleet compliance management view."""
        self.send_json(build_fleet_compliance_payload(_tenant_from_headers(self.headers)))
    
    def handle_get_pave_vehicles(self, query):
        """Return PAVE vehicle data with all Amazon wear and tear data points."""
        conn = get_connection()
        cursor = conn.cursor()
        
        # Get all vans with PAVE-related data
        cursor.execute("""
            SELECT 
                v.id, v.van_number, v.vin, v.make, v.model, v.year, v.status, v.ownership, v.notes, v.updated_at,
                p.inspection_date, p.inspector_name, p.overall_score, p.compliance_status, p.pave_status,
                p.exterior_score, p.body_damage, p.scratches_dents, p.rust, p.paint_condition, 
                p.decals_logos, p.lights, p.mirrors, p.windows,
                p.tire_condition, p.tread_depth, p.tire_pressure, p.spare_tire,
                p.interior_score, p.seat_condition, p.floor_mats, p.dashboard, 
                p.steering_wheel, p.pedals, p.cargo_area_cleanliness, p.odor,
                p.mechanical_score, p.engine, p.transmission, p.brakes, p.suspension, 
                p.exhaust, p.fluids, p.battery, p.heating_ac,
                p.fire_extinguisher, p.first_aid_kit, p.reflective_triangles, p.jump_starter,
                p.interior_cleanliness_score, p.exterior_cleanliness_score,
                p.registration_valid, p.insurance_valid, p.dot_inspection_date, 
                p.maintenance_records_up_to_date, p.next_inspection_due, p.notes as pave_notes
            FROM vans v
            LEFT JOIN pave_inspections p ON v.id = p.van_id
            ORDER BY v.van_number
        """)
        rows = cursor.fetchall()
        
        # Transform to the expected format for the dashboard
        vehicles = []
        for row in rows:
            van = dict(row)
            
            # Use PAVE data if available, otherwise determine from van status
            if van.get('overall_score'):
                compliance_status = van.get('compliance_status', 'unknown')
                pave_status = van.get('pave_status', 'grey')
                pave_score = van.get('overall_score', 0)
            elif van['status'] == 'OPERATIONAL':
                compliance_status = 'compliant'
                pave_status = 'green'
                pave_score = 100
            elif van['status'] == 'GROUND':
                compliance_status = 'non-compliant'
                pave_status = 'red'
                pave_score = 0
            else:
                compliance_status = 'unknown'
                pave_status = 'grey'
                pave_score = 50
            
            vehicles.append({
                'vin': van['vin'],
                'licensePlate': van.get('van_number', 'N/A'),
                'vanNumber': van.get('van_number', 'N/A'),
                'make': van.get('make', 'Unknown'),
                'model': van.get('model', 'Unknown'),
                'year': van.get('year'),
                'ownership': van.get('ownership', 'Unknown'),
                'status': van.get('status', 'Unknown'),
                'complianceStatus': compliance_status,
                'paveStatus': pave_status,
                'paveScore': pave_score,
                'lastPaveInspectionDate': van.get('inspection_date', van.get('updated_at', 'N/A')),
                'nextPaveInspectionDue': van.get('next_inspection_due', 'N/A'),
                'inspectorName': van.get('inspector_name', 'N/A'),
                
                # Exterior data points
                'exteriorScore': van.get('exterior_score'),
                'bodyDamage': van.get('body_damage', 'N/A'),
                'scratchesDents': van.get('scratches_dents', 'N/A'),
                'rust': van.get('rust', 'N/A'),
                'paintCondition': van.get('paint_condition', 'N/A'),
                'decalsLogos': van.get('decals_logos', 'N/A'),
                'lights': van.get('lights', 'N/A'),
                'mirrors': van.get('mirrors', 'N/A'),
                'windows': van.get('windows', 'N/A'),
                
                # Tire data points
                'tireCondition': van.get('tire_condition', 'N/A'),
                'treadDepth': van.get('tread_depth'),
                'tirePressure': van.get('tire_pressure', 'N/A'),
                'spareTire': van.get('spare_tire', 'N/A'),
                
                # Interior data points
                'interiorScore': van.get('interior_score'),
                'seatCondition': van.get('seat_condition', 'N/A'),
                'floorMats': van.get('floor_mats', 'N/A'),
                'dashboard': van.get('dashboard', 'N/A'),
                'steeringWheel': van.get('steering_wheel', 'N/A'),
                'pedals': van.get('pedals', 'N/A'),
                'cargoAreaCleanliness': van.get('cargo_area_cleanliness', 'N/A'),
                'odor': van.get('odor', 'N/A'),
                
                # Mechanical data points
                'mechanicalScore': van.get('mechanical_score'),
                'engine': van.get('engine', 'N/A'),
                'transmission': van.get('transmission', 'N/A'),
                'brakes': van.get('brakes', 'N/A'),
                'suspension': van.get('suspension', 'N/A'),
                'exhaust': van.get('exhaust', 'N/A'),
                'fluids': van.get('fluids', 'N/A'),
                'battery': van.get('battery', 'N/A'),
                'heatingAc': van.get('heating_ac', 'N/A'),
                
                # Safety equipment
                'fireExtinguisher': van.get('fire_extinguisher', 'N/A'),
                'firstAidKit': van.get('first_aid_kit', 'N/A'),
                'reflectiveTriangles': van.get('reflective_triangles', 'N/A'),
                'jumpStarter': van.get('jump_starter', 'N/A'),
                
                # Cleanliness scores
                'interiorCleanlinessScore': van.get('interior_cleanliness_score'),
                'exteriorCleanlinessScore': van.get('exterior_cleanliness_score'),
                
                # Documentation
                'registrationValid': van.get('registration_valid', 'N/A'),
                'insuranceValid': van.get('insurance_valid', 'N/A'),
                'dotInspectionDate': van.get('dot_inspection_date', 'N/A'),
                'maintenanceRecordsUpToDate': van.get('maintenance_records_up_to_date', 'N/A'),
                
                # Notes
                'paveNotes': van.get('pave_notes', van.get('notes', ''))
            })
        
        result = {
            'total': len(vehicles),
            'items': vehicles
        }

        tenant = _tenant_from_headers(self.headers)
        pave_upload, pave_preview = _latest_pave_upload_preview(tenant)
        if pave_preview:
            by_vin = {item['vin']: item for item in pave_preview['latestByVin']}
            for vehicle in result['items']:
                assessment = by_vin.get(vehicle.get('vin'))
                if not assessment:
                    continue
                vehicle.update({
                    'licensePlate': assessment.get('licensePlate') or vehicle.get('licensePlate'),
                    'paveStatus': 'red' if assessment.get('groundingRisk') or assessment.get('grade') == 2 else 'green',
                    'paveScore': assessment.get('conditionScore'),
                    'complianceStatus': 'non-compliant' if assessment.get('grade') == 2 else 'compliant',
                    'lastPaveInspectionDate': assessment.get('createdAt'),
                    'paveGrade': assessment.get('grade'),
                    'paveGradeLabel': assessment.get('gradeLabel'),
                    'hasNewDamage': assessment.get('hasNewDamage'),
                    'groundingRisk': assessment.get('groundingRisk'),
                    'sessionKey': assessment.get('sessionKey'),
                })
            result['source'] = {
                'provider': 'PAVE Fleet Dashboard CSV',
                'filename': pave_upload['originalFilename'],
                'confirmedAt': pave_upload['confirmedAt'],
                'summary': pave_preview['summary'],
            }
        
        conn.close()
        self.send_json(result)
    
    # ========== File Serving Handlers ==========
    
    def handle_dashboard_html(self):
        """Serve the main dashboard HTML."""
        if not DASHBOARD_PATH.exists():
            self.send_error(404, "Dashboard not found")
            return
        
        with open(DASHBOARD_PATH, 'rb') as f:
            content = f.read()
        
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', len(content))
        self.end_headers()
        self.wfile.write(content)
    
    def handle_form_html(self):
        """Serve the daily entry form HTML."""
        form_path = ROOT / "data/dashboards/daily-entry-form.html"
        if not form_path.exists():
            self.send_error(404, "Form not found")
            return
        
        with open(form_path, 'rb') as f:
            content = f.read()
        
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', len(content))
        self.end_headers()
        self.wfile.write(content)
    
    def handle_dashboard_file(self):
        """Serve static files from the dashboard directory."""
        parsed = urllib.parse.urlparse(self.path)
        file_path = DASHBOARD_DIR / parsed.path.lstrip('/data/dashboards/')
        
        if file_path.exists() and file_path.is_file():
            with open(file_path, 'rb') as f:
                content = f.read()
            
            content_type = 'text/html'
            if self.path.endswith('.css'):
                content_type = 'text/css'
            elif self.path.endswith('.js'):
                content_type = 'application/javascript'
            elif self.path.endswith('.json'):
                content_type = 'application/json'
            
            self.send_response(200)
            self.send_header('Content-Type', content_type + '; charset=utf-8')
            self.send_header('Content-Length', len(content))
            self.end_headers()
            self.wfile.write(content)
        else:
            self.send_error(404, f"File not found: {self.path}")
    
    def handle_redirect_to_dashboard(self):
        """Redirect to the main dashboard."""
        self.send_response(302)
        self.send_header('Location', '/amazon-dsp-kpi-dashboard.html')
        self.end_headers()
    
    # ========== Auth Handlers (Mock for Development) ==========
    
    def handle_get_current_user(self):
        """Return mock user for development."""
        mock_user = {
            "id": "dev-user",
            "email": "dev@example.com",
            "firstName": "Developer",
            "lastName": "User",
            "role": "admin",
            "status": "active"
        }
        self.send_json(mock_user)
    
    def handle_post_login(self):
        """Mock login endpoint for development."""
        content_length = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(content_length)
        data = json.loads(body)
        
        mock_response = {
            "user": {
                "id": "dev-user",
                "email": data.get('email', 'dev@example.com'),
                "firstName": "Developer",
                "lastName": "User",
                "role": "admin"
            },
            "token": "mock-token-for-dev",
            "refreshToken": "mock-refresh-token-for-dev",
            "expiresIn": 3600
        }
        self.send_json(mock_response)
    
    def handle_post_logout(self):
        """Mock logout endpoint for development."""
        self.send_json({"message": "Logged out successfully"})
    
    def handle_post_refresh(self):
        """Mock token refresh endpoint for development."""
        content_length = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(content_length)
        data = json.loads(body)
        
        mock_response = {
            "token": "mock-refreshed-token-for-dev",
            "refreshToken": data.get('refreshToken', 'mock-refresh-token-for-dev'),
            "expiresIn": 3600
        }
        self.send_json(mock_response)

    def handle_get_weekly_evaluations(self):
        """Expose the incumbent dashboard's weekly evaluation payload to React."""
        if _tenant_from_headers(self.headers) != DEFAULT_TENANT:
            self.send_json({'weeks': [], 'evaluations': {}, 'needsData': True,
                            'message': 'No tenant-scoped weekly evaluations have been generated'})
            return
        html = DASHBOARD_PATH.read_text(encoding='utf-8')
        match = re.search(
            r'window\.__WEEKLY_EVALUATIONS__=(\{.*?\});</script>',
            html,
            flags=re.DOTALL,
        )
        if not match:
            self.send_json({'weeks': [], 'evaluations': {}})
            return

        evaluations = json.loads(match.group(1))
        self.send_json({
            'weeks': list(evaluations.keys()),
            'evaluations': evaluations,
        })

    def handle_get_time_attendance_exceptions(self):
        """Return real ADP exceptions, reconciled against daily Amazon assignments."""
        tenant = _tenant_from_headers(self.headers)
        if tenant != DEFAULT_TENANT:
            self.send_json({
                'tenant': tenant, 'source': None, 'sourcePeriod': None,
                'capturedAt': None, 'needsData': True,
                'message': 'No tenant-scoped ADP timecards have been ingested',
                'coverage': {'employees': 0, 'dailyAssignments': 0,
                             'routeReconciliationAvailable': False},
                'exceptions': [],
            })
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM driver_route_assignments")
        assignment_count = cursor.fetchone()[0]
        cursor.execute("""
            SELECT employee, date, issue_type AS issueType, details
            FROM time_attendance_issues
            WHERE issue_type = 'Long shift' OR ? > 0
            ORDER BY date DESC, employee
        """, (assignment_count,))
        exceptions = [dict(row) for row in cursor.fetchall()]
        cursor.execute("SELECT MIN(date), MAX(date), COUNT(DISTINCT driver_id) FROM adp_timecards")
        first_date, last_date, employee_count = cursor.fetchone()
        conn.close()

        summaries = sorted((ROOT / 'data' / 'adp').glob('????-??-??_to_????-??-??/summary.json'))
        captured_at = None
        if summaries:
            try:
                captured_at = json.loads(summaries[-1].read_text()).get('capturedAt')
            except (OSError, ValueError):
                captured_at = None
        period = f'{first_date} to {last_date}' if first_date and last_date else None
        self.send_json({
            'source': 'ADP Workforce Now API',
            'sourcePeriod': period,
            'capturedAt': captured_at,
            'needsData': not period,
            'message': (
                f'Live ADP timecards for {employee_count} employees; '
                + ('reconciled with Amazon daily assignments' if assignment_count else
                   'Amazon daily assignments are not loaded, so route-based exceptions are withheld')
            ) if period else 'No ADP timecards have been ingested',
            'coverage': {
                'employees': employee_count or 0,
                'dailyAssignments': assignment_count,
                'routeReconciliationAvailable': bool(assignment_count),
            },
            'exceptions': exceptions,
        })

    def send_paginated(self, rows):
        """Return the pagination envelope expected by the React tables."""
        self.send_json({
            'data': rows,
            'meta': {
                'currentPage': 1,
                'totalPages': 1,
                'totalItems': len(rows),
                'itemsPerPage': max(len(rows), 1),
                'hasNextPage': False,
                'hasPreviousPage': False,
            },
        })

    def handle_get_fleet_cost_summary(self, query):
        """Return the reconciled three-month fleet totals in the React schema."""
        payload = build_fleet_cost_reconciliation(_tenant_from_headers(self.headers))
        summary = payload.get('summary', {})
        self.send_json({
            'period': payload.get('period'),
            'asOf': payload.get('asOf'),
            'totalCost': summary.get('threeMonthIncludedCost', 0),
            'totalFixedCost': summary.get('thirdPartyRentalCost', 0),
            'totalVariableCost': round(summary.get('lmrCost', 0) + summary.get('elementCost', 0), 2),
            'totalCapitalCost': 0,
            'totalOperatingCost': summary.get('threeMonthIncludedCost', 0),
            'costByCategory': {
                'thirdPartyRental': summary.get('thirdPartyRentalCost', 0),
                'lmr': summary.get('lmrCost', 0),
                'element': summary.get('elementCost', 0),
            },
            'costByVan': [],
            'costByDriver': [],
            'costPerMile': 0,
            'costPerDay': round(summary.get('threeMonthIncludedCost', 0) / 92, 2),
            'costPerRoute': 0,
            'costPerDelivery': 0,
            'fuelEfficiency': 0,
            'maintenanceCostPerMile': 0,
        })

    def handle_get_fuel_cost_analysis(self, query):
        """Return an empty but valid fuel-analysis envelope when fuel detail is unavailable."""
        self.send_json({
            'period': 'June–August 2026',
            'totalFuelCost': 0,
            'totalGallons': 0,
            'averagePricePerGallon': 0,
            'totalMiles': 0,
            'fuelEfficiency': 0,
            'costPerMile': 0,
            'byVan': [],
            'byDriver': [],
            'trends': [],
        })

    def handle_get_maintenance_cost_analysis(self, query):
        """Return an empty but valid maintenance-analysis envelope."""
        self.send_json({
            'period': 'June–August 2026',
            'totalMaintenanceCost': 0,
            'byVan': [],
            'byCategory': {},
            'averageCostPerMile': 0,
            'averageCostPerVan': 0,
            'trends': [],
        })

    def handle_get_fleet_cost_trends(self, query):
        """Return the monthly reconciliation used by the incumbent dashboard."""
        payload = build_fleet_cost_reconciliation(_tenant_from_headers(self.headers))
        self.send_json([{
            'period': row['month'], 'totalCost': row['includedCost'],
            'amazonCoverage': row['amazonCoverage'], 'difference': row['difference'],
            'status': row['status'], 'costByCategory': {}, 'costPerMile': 0, 'costPerDelivery': 0,
        } for row in payload.get('months', [])])
    

    # ========== Connections and tenant uploads ==========

    def _read_body(self):
        length = int(self.headers.get('Content-Length') or 0)
        if length <= 0:
            return b''
        if length > 25 * 1024 * 1024:
            raise ValueError('upload exceeds the 25 MB limit')
        return self.rfile.read(length)

    def handle_post_assistant_chat(self):
        tenant = _tenant_from_headers(self.headers)
        try:
            body = json.loads(self._read_body() or b'{}')
        except (TypeError, ValueError):
            self.send_json_status(400, {'error': 'invalid JSON body'})
            return
        page = body.get('page') if isinstance(body.get('page'), dict) else {}
        try:
            result = assistant_service.ask(
                tenant=tenant,
                actor=self.headers.get('x-actor') or 'dev-user',
                message=body.get('message'),
                history=body.get('history'),
                page=page,
                snapshot=build_assistant_snapshot(tenant, page.get('path') or '/dashboard'),
            )
            self.send_json_status(200, result)
        except ValueError as error:
            self.send_json_status(400, {'error': str(error)})
        except RuntimeError as error:
            status = 503 if 'not configured' in str(error) or 'unavailable' in str(error) else 502
            self.send_json_status(status, {'error': str(error)})

    def handle_put_assistant_config(self):
        """Store OpenAI credentials in the managed local vault; never echo the key."""
        tenant = _tenant_from_headers(self.headers)
        body = json.loads(self._read_body() or b'{}')
        api_key = str(body.get('apiKey') or '').strip()
        model = str(body.get('model') or 'gpt-5-mini').strip()
        organization = str(body.get('organization') or '').strip()
        project = str(body.get('project') or '').strip()
        if len(api_key) < 20 or len(api_key) > 512 or any(char.isspace() for char in api_key):
            raise ValueError('invalid OpenAI API key')
        if not re.fullmatch(r'[A-Za-z0-9._:-]{2,100}', model):
            raise ValueError('invalid model')
        for name, value in (('organization', organization), ('project', project)):
            if value and not re.fullmatch(r'[A-Za-z0-9_-]{2,160}', value):
                raise ValueError(f'invalid {name}')
        bundle = {'apiKey': api_key, 'model': model}
        if organization:
            bundle['organization'] = organization
        if project:
            bundle['project'] = project
        put_secret(tenant, 'openai', 'credentials', json.dumps(bundle, separators=(',', ':')))
        self.send_json_status(200, {
            'configured': True, 'provider': 'OpenAI', 'model': model,
            'organizationConfigured': bool(organization), 'projectConfigured': bool(project),
            'storage': 'Local managed secret store (AWS Secrets Manager in production)',
            'writeOnly': True, 'lastFour': api_key[-4:],
        })

    def handle_post_assistant_config_test(self):
        tenant = _tenant_from_headers(self.headers)
        credential = assistant_service.credentials(tenant)
        if not credential:
            self.send_json_status(409, {'error': 'AI assistant is not configured'})
            return
        headers = {'Authorization': f"Bearer {credential['apiKey']}"}
        if credential.get('organization'):
            headers['OpenAI-Organization'] = credential['organization']
        if credential.get('project'):
            headers['OpenAI-Project'] = credential['project']
        request = urllib.request.Request('https://api.openai.com/v1/models', headers=headers)
        try:
            with urllib.request.urlopen(request, timeout=15):
                pass
        except urllib.error.HTTPError:
            self.send_json_status(502, {'error': 'OpenAI credential validation failed'})
            return
        except (urllib.error.URLError, TimeoutError):
            self.send_json_status(502, {'error': 'OpenAI is temporarily unavailable'})
            return
        self.send_json_status(200, {'status': 'healthy', 'provider': 'OpenAI',
                                    'model': credential.get('model') or 'gpt-5-mini',
                                    'checkedAt': datetime.now(timezone.utc).isoformat(timespec='seconds')})

    def send_json_status(self, status, data):
        content = json.dumps(data, default=str).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', len(content))
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(content)

    def handle_post_upload(self, source):
        """Store an uploaded export and return a parse preview. Nothing is
        committed to reporting until the tenant confirms it."""
        tenant = _tenant_from_headers(self.headers)
        body = self._read_body()
        fields, files = parse_multipart(body, self.headers.get('Content-Type'))
        if 'file' not in files:
            self.send_json_status(400, {'error': 'no file field in upload'})
            return
        filename, payload = files['file']
        actor = fields.get('uploadedBy') or self.headers.get('x-actor') or 'local-dev@jeclogs.com'

        if source not in (*connections_registry.FINANCIAL_SOURCES, 'pave'):
            self.send_json_status(400, {'error': f'unsupported upload source: {source}'})
            return

        record = tenant_store.put_file(tenant, source, filename, payload, actor)
        existing_status = record.get('status')
        preview = (parse_pave_export(payload, filename) if source == 'pave'
                   else parse_financial_export(payload, filename, tenant=tenant,
                                               provider=fields.get('provider') or None))

        if preview.get('ok'):
            summary = dict(preview['summary'])
            summary['periods'] = preview.get('periods')
            summary['provider'] = preview.get('provider')
            summary['providerLabel'] = preview.get('providerLabel')
            record = tenant_store.update_file(
                tenant, record['id'], status='confirmed' if existing_status == 'confirmed' else 'parsed', provider=preview.get('provider'),
                periodKey=preview.get('periodKey'), parseSummary=summary)
        else:
            record = tenant_store.update_file(
                tenant, record['id'], status='rejected', rejectedReason=preview.get('error'))

        self.send_json_status(200 if preview.get('ok') else 422, {
            'tenant': tenant, 'upload': record, 'preview': preview,
        })

    def handle_post_upload_reparse(self, file_id):
        """Re-run the parser on stored bytes, e.g. after vendor rules change."""
        tenant = _tenant_from_headers(self.headers)
        record = tenant_store.get_file(tenant, file_id)
        if not record:
            self.send_json_status(404, {'error': 'unknown upload'})
            return
        preview = (parse_pave_export(tenant_store.read_bytes(tenant, file_id), record['originalFilename'])
                   if record['source'] == 'pave' else
                   parse_financial_export(tenant_store.read_bytes(tenant, file_id),
                                          record['originalFilename'], tenant=tenant))
        if preview.get('ok'):
            summary = dict(preview['summary'])
            summary['periods'] = preview.get('periods')
            summary['provider'] = preview.get('provider')
            summary['providerLabel'] = preview.get('providerLabel')
            record = tenant_store.update_file(
                tenant, file_id, status='parsed', provider=preview.get('provider'),
                periodKey=preview.get('periodKey'), parseSummary=summary)
        self.send_json_status(200, {'tenant': tenant, 'upload': record, 'preview': preview})

    def handle_post_upload_confirm(self, file_id):
        """Confirm a parsed upload so reporting screens start using it."""
        tenant = _tenant_from_headers(self.headers)
        record = tenant_store.get_file(tenant, file_id)
        if not record:
            self.send_json_status(404, {'error': 'unknown upload'})
            return
        if record['status'] not in ('parsed', 'confirmed'):
            self.send_json_status(409, {'error': f"upload is {record['status']} and cannot be confirmed"})
            return
        superseded = tenant_store.supersede_others(tenant, record['source'], file_id, record.get('periodKey'))
        record = tenant_store.update_file(
            tenant, file_id, status='confirmed',
            confirmedAt=datetime.now(timezone.utc).isoformat(timespec='seconds'))
        connections_registry.set_connection_state(
            tenant, record['source'], status='healthy',
            lastSuccessAt=record['confirmedAt'], lastError=None)
        self.send_json_status(200, {'tenant': tenant, 'upload': record, 'supersededCount': superseded})

    def handle_post_vendor_rules(self):
        """Replace the tenant's vendor classification rules."""
        tenant = _tenant_from_headers(self.headers)
        try:
            body = json.loads(self._read_body() or b'{}')
        except ValueError:
            self.send_json_status(400, {'error': 'invalid JSON body'})
            return
        rules = body.get('rules')
        if not isinstance(rules, list):
            self.send_json_status(400, {'error': 'rules must be a list'})
            return
        saved = save_vendor_rules(tenant, rules)
        _FLEET_COST_CACHE.clear()
        self.send_json_status(200, {'tenant': tenant, 'rules': saved})

    def handle_post_connection_test(self, connection_id):
        """Record a connection check. Browser-session sources report the real
        local session state instead of claiming success."""
        tenant = _tenant_from_headers(self.headers)
        entry = connections_registry.BY_ID.get(connection_id)
        if not entry:
            self.send_json_status(404, {'error': 'unknown connection'})
            return
        now = datetime.now(timezone.utc).isoformat(timespec='seconds')
        if connection_id == 'digits_api':
            public = next(item for item in build_connections_payload(tenant)['connections']
                          if item['id'] == connection_id)
            environment = public.get('environment') or 'development'
            try:
                from scripts.digits_api_client import list_entities
                entities = list_entities(tenant, environment)
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='healthy', lastCheckedAt=now,
                    lastSuccessAt=now, lastError=None, environment=environment,
                    secretReference=connections_registry.secret_reference(
                        tenant, connection_id, 'credentials'),
                    configuredFields=['clientId', 'clientSecret'])
                self.send_json_status(200, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'status': 'healthy', 'provider': 'Digits API',
                    'message': 'Read-only ledger access verified',
                    'entityCount': len(entities),
                })
                return
            except Exception:
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='degraded', lastCheckedAt=now,
                    lastError='Provider authentication failed', environment=environment)
                self.send_json_status(502, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'error': 'provider connection test failed',
                })
                return
        if connection_id == 'email_imap':
            try:
                host = get_secret(tenant, 'email-imap', 'production-host')
                port = int(get_secret(tenant, 'email-imap', 'production-port'))
                username = get_secret(tenant, 'email-imap', 'production-username')
                password = get_secret(tenant, 'email-imap', 'production-app-password')
                with imaplib.IMAP4_SSL(host, port, ssl_context=ssl.create_default_context(), timeout=30) as client:
                    client.login(username, password)
                    status, _ = client.select('INBOX', readonly=True)
                    if status != 'OK':
                        raise RuntimeError('read-only INBOX access unavailable')
                    client.logout()
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='healthy', lastCheckedAt=now,
                    lastSuccessAt=now, lastError=None,
                    secretReference=connections_registry.secret_reference(
                        tenant, connection_id, 'credentials'),
                    configuredFields=['host', 'port', 'username', 'appPassword'])
                self.send_json_status(200, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'status': 'healthy', 'provider': 'IMAP',
                    'message': 'Read-only INBOX access verified',
                })
                return
            except Exception:
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='degraded', lastCheckedAt=now,
                    lastError='IMAP authentication or read-only INBOX access failed.')
                self.send_json_status(502, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'error': 'provider connection test failed',
                })
                return
        if connection_id == 'adp':
            try:
                worker_count = _test_adp_connection(tenant)
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='healthy', lastCheckedAt=now,
                    lastSuccessAt=now, lastError=None, environment='production',
                    secretReference=connections_registry.secret_reference(
                        tenant, connection_id, 'credentials'),
                    configuredFields=['clientId', 'clientSecret', 'certificatePem', 'privateKeyPem'])
                self.send_json_status(200, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'status': 'healthy', 'provider': 'ADP Workforce Now',
                    'message': 'OAuth and read-only worker access verified',
                    'workerCountInProbe': worker_count,
                })
                return
            except Exception as error:
                safe_error = str(error) if str(error).startswith('ADP ') else 'ADP connection validation failed'
                state = connections_registry.set_connection_state(
                    tenant, connection_id, status='degraded', lastCheckedAt=now,
                    lastError=safe_error, environment='production')
                self.send_json_status(502, {
                    'tenant': tenant, 'connection': connection_id, 'state': state,
                    'error': safe_error,
                })
                return
        if entry['authKind'] == 'manual_upload':
            latest = tenant_store.latest_confirmed(tenant, connection_id)
            ok = latest is not None
            state = connections_registry.set_connection_state(
                tenant, connection_id,
                status='healthy' if ok else 'not_connected',
                lastCheckedAt=now,
                lastError=None if ok else 'No confirmed upload yet.')
        else:
            state = connections_registry.set_connection_state(
                tenant, connection_id, status='pending', lastCheckedAt=now,
                lastError='Connector not provisioned in local development.')
        self.send_json_status(200, {'tenant': tenant, 'connection': connection_id, 'state': state})

    def handle_post_connections_refresh(self):
        tenant = _tenant_from_headers(self.headers)
        existing = _CONNECTION_REFRESH_JOBS.get(tenant, {})
        if existing.get('status') == 'running':
            self.send_json_status(202, existing)
            return
        job = {'status': 'running',
               'startedAt': datetime.now(timezone.utc).isoformat(timespec='seconds'),
               'finishedAt': None, 'sources': []}
        _CONNECTION_REFRESH_JOBS[tenant] = job
        threading.Thread(target=refresh_all_connection_data, args=(tenant,), daemon=True).start()
        self.send_json_status(202, job)

    def handle_put_connection_credentials(self, connection_id):
        """Write local connector secrets without returning their values."""
        tenant = _tenant_from_headers(self.headers)
        definitions = connections_registry.BY_ID
        if connection_id not in ('digits_api', 'adp', 'email_imap', 'pave'):
            raise ValueError('this connection does not accept credentials')
        body = json.loads(self._read_body() or b'{}')
        environment = body.get('environment')
        credentials = body.get('credentials') or {}
        if environment not in ('development', 'production'):
            raise ValueError('invalid connector environment')
        expected = [field['name'] for field in definitions[connection_id].get('credentialFields', [])]
        for name in expected:
            value = credentials.get(name)
            if not isinstance(value, str) or len(value.strip()) < 2 or len(value) > 65536:
                raise ValueError(f'invalid credential field: {name}')
        integration = 'digits' if connection_id == 'digits_api' else connection_id.replace('_', '-')
        for name in expected:
            stored_name = {'clientId': 'client-id', 'clientSecret': 'client-secret',
                           'certificatePem': 'certificate-pem', 'privateKeyPem': 'private-key-pem',
                           'appPassword': 'app-password'}.get(name, name)
            put_secret(tenant, integration, f'{environment}-{stored_name}', credentials[name])
        reference = connections_registry.secret_reference(tenant, connection_id, 'credentials')
        connections_registry.set_connection_state(
            tenant, connection_id, status='pending', secretReference=reference,
            environment=environment, configuredFields=expected,
            lastError=None)
        public = next(item for item in build_connections_payload(tenant)['connections']
                      if item['id'] == connection_id)
        self.send_json_status(200, {'connection': public})

    def handle_post_connection_reconnect(self, connection_id):
        tenant = _tenant_from_headers(self.headers)
        urls = {
            'amazon': 'https://logistics.amazon.com/dspconsolev2',
            'pave': connections_registry.PAVE_LOGIN_URL,
        }
        if connection_id not in urls:
            raise ValueError('this connection does not use a browser reconnect flow')
        if not urls[connection_id]:
            raise ValueError('reconnect adapter is unavailable')
        state = connections_registry.set_connection_state(
            tenant, connection_id, status='needs_reauth',
            lastCheckedAt=datetime.now(timezone.utc).isoformat(timespec='seconds'),
            lastError=('Complete Amazon sign-in and MFA in the reconnect browser.'
                       if connection_id == 'amazon' else
                       'Complete PAVE sign-in in the reconnect browser.'))
        launch_mode = 'external_url'
        if connection_id in ('amazon', 'pave'):
            log_path = ROOT / '.openclaw' / f'{connection_id}-login.log'
            log_path.parent.mkdir(parents=True, exist_ok=True)
            log = open(log_path, 'ab', buffering=0)
            child_environment = os.environ.copy()
            if connection_id == 'pave':
                child_environment['PAVE_USERNAME'] = get_secret(tenant, 'pave', 'production-username')
                child_environment['PAVE_PASSWORD'] = get_secret(tenant, 'pave', 'production-password')
            process = subprocess.Popen(
                ['node', str(ROOT / 'scripts' / ('amazon_logistics_login.mjs' if connection_id == 'amazon' else 'pave_login.mjs'))],
                cwd=str(ROOT), stdout=log, stderr=subprocess.STDOUT,
                start_new_session=True, env=child_environment)
            def watch_login():
                return_code = process.wait()
                log.close()
                checked_at = datetime.now(timezone.utc).isoformat(timespec='seconds')
                if return_code == 0:
                    connections_registry.set_connection_state(
                        tenant, connection_id, status='healthy', lastSuccessAt=checked_at,
                        lastCheckedAt=checked_at, lastError=None)
                else:
                    connections_registry.set_connection_state(
                        tenant, connection_id, status='needs_reauth', lastCheckedAt=checked_at,
                        lastError=f'{connections_registry.BY_ID[connection_id]["displayName"]} sign-in did not complete. Reopen the managed browser and finish authentication.')
            threading.Thread(target=watch_login, daemon=True, name=f'{connection_id}-login-{tenant}').start()
            launch_mode = 'local_managed_browser'
        self.send_json_status(200, {
            'connection': connection_id, 'status': state['status'],
            'authorizationUrl': urls[connection_id],
            'launchMode': launch_mode,
            'message': ('Complete Amazon sign-in and MFA once. The shared session authorizes all Amazon-backed features.'
                        if connection_id == 'amazon' else
                        'Complete PAVE sign-in. This session is separate from Amazon and feeds fleet compliance data.'),
        })

    def handle_put_feature(self, feature_id):
        tenant = _tenant_from_headers(self.headers)
        body = json.loads(self._read_body() or b'{}')
        if not isinstance(body.get('enabled'), bool):
            raise ValueError('enabled must be a boolean')
        self.send_json_status(200, {'feature': _set_feature_override(
            tenant, feature_id, body['enabled'])})

    def handle_post_member_invitation(self):
        tenant = _tenant_from_headers(self.headers)
        body = json.loads(self._read_body() or b'{}')
        email_address = str(body.get('email') or '').strip().lower()
        given_name = str(body.get('givenName') or '').strip()
        family_name = str(body.get('familyName') or '').strip()
        role = str(body.get('role') or 'viewer')
        if not _valid_member_email(email_address) or not given_name or not family_name or role not in {'owner', 'admin', 'reviewer', 'analyst', 'viewer'}:
            self.send_json_status(400, {'error': 'first name, last name, valid email, and tenant role are required'})
            return
        members = _local_members(tenant)
        if any(item.get('email', '').lower() == email_address for item in members):
            self.send_json_status(409, {'error': 'that email is already a tenant member'})
            return
        member = {
            'identitySubject': f'invited:{uuid.uuid4()}', 'email': email_address,
            'givenName': given_name, 'familyName': family_name, 'role': role,
            'status': 'invited', 'createdAt': datetime.now(timezone.utc).isoformat(timespec='seconds')
        }
        members.append(member)
        _save_local_members(tenant, members)
        self.send_json_status(201, {'member': member, 'invitationSent': False})

    def handle_post_member_invitation_resend(self, identity_subject):
        tenant = _tenant_from_headers(self.headers)
        members = _local_members(tenant)
        member = next((item for item in members if item.get('identitySubject') == identity_subject), None)
        if not member:
            self.send_json_status(404, {'error': 'member not found'})
            return
        if member.get('status') != 'invited':
            self.send_json_status(409, {'error': 'only pending invitations can be resent'})
            return
        member['lastInvitationResentAt'] = datetime.now(timezone.utc).isoformat(timespec='seconds')
        _save_local_members(tenant, members)
        self.send_json_status(200, {'member': member, 'invitationSent': False, 'deliveryMode': 'development'})

    def handle_put_member(self, identity_subject):
        tenant = _tenant_from_headers(self.headers)
        body = json.loads(self._read_body() or b'{}')
        role = str(body.get('role') or '')
        status = str(body.get('status') or '')
        if role not in {'owner', 'admin', 'reviewer', 'analyst', 'viewer'} or status not in {'invited', 'active', 'disabled'}:
            self.send_json_status(400, {'error': 'valid role and status required'})
            return
        members = _local_members(tenant)
        member = next((item for item in members if item.get('identitySubject') == identity_subject), None)
        if not member:
            self.send_json_status(404, {'error': 'member not found'})
            return
        member.update({'role': role, 'status': status})
        _save_local_members(tenant, members)
        self.send_json_status(200, {'member': member})

    def handle_post_super_admin_tenant(self):
        body = json.loads(self._read_body() or b'{}')
        slug = str(body.get('slug') or '').strip()
        display_name = str(body.get('displayName') or '').strip()
        owner_email = str(body.get('ownerEmail') or '').strip().lower()
        if not re.fullmatch(r'[a-z][a-z0-9-]{2,62}', slug):
            self.send_json_status(400, {'error': 'invalid tenant slug'})
            return
        if not display_name or not _valid_member_email(owner_email):
            self.send_json_status(400, {'error': 'display name and valid owner email are required'})
            return
        tenants = _local_tenants()
        if any(item['slug'] == slug for item in tenants):
            self.send_json_status(409, {'error': 'tenant slug already exists'})
            return
        now = datetime.now(timezone.utc).isoformat(timespec='seconds')
        tenant = {'id': f'local-{uuid.uuid4()}', 'slug': slug, 'displayName': display_name,
                  'status': 'active', 'createdAt': now,
                  'moduleIds': body.get('moduleIds') if isinstance(body.get('moduleIds'), list) else []}
        tenants.append(tenant)
        _save_local_tenants(tenants)
        owner = {'identitySubject': f'invited:{uuid.uuid4()}', 'email': owner_email,
                 'givenName': str(body.get('ownerGivenName') or '').strip(),
                 'familyName': str(body.get('ownerFamilyName') or '').strip(),
                 'role': 'owner', 'status': 'invited', 'createdAt': now}
        _save_local_members(slug, [owner])
        self.send_json_status(201, {'tenant': tenant, 'owner': {**owner, 'invitationSent': False}})

    def handle_put_super_admin_status(self, tenant_slug):
        body = json.loads(self._read_body() or b'{}')
        status = str(body.get('status') or '')
        if status not in {'active', 'suspended', 'closed'}:
            self.send_json_status(400, {'error': 'invalid tenant status'})
            return
        tenants = _local_tenants()
        tenant = next((item for item in tenants if item['slug'] == tenant_slug), None)
        if not tenant:
            self.send_json_status(404, {'error': 'tenant not found'})
            return
        tenant['status'] = status
        tenant['updatedAt'] = datetime.now(timezone.utc).isoformat(timespec='seconds')
        _save_local_tenants(tenants)
        self.send_json_status(200, {'tenant': tenant})

    def handle_post_super_admin_impersonation(self):
        body = json.loads(self._read_body() or b'{}')
        tenant_slug = str(body.get('tenantSlug') or '')
        target_subject = str(body.get('targetSubject') or '')
        reason = str(body.get('reason') or '').strip()
        target = next((item for item in _local_members(tenant_slug)
                       if item.get('identitySubject') == target_subject and item.get('status') == 'active'), None)
        if not target:
            self.send_json_status(404, {'error': 'active target user not found in tenant'})
            return
        if not 10 <= len(reason) <= 500:
            self.send_json_status(400, {'error': 'reason must be 10-500 characters'})
            return
        token = f'local-support-{uuid.uuid4()}'
        _LOCAL_SUPPORT_SESSIONS[token] = {**target, 'tenantSlug': tenant_slug, 'reason': reason,
                                          'expiresAt': int((datetime.now(timezone.utc).timestamp() + 900) * 1000)}
        self.send_json_status(201, {'token': token, 'expiresInSeconds': 900,
                                    'target': {'email': target['email'], 'role': target['role'],
                                               'tenantSlug': tenant_slug}})

    def send_binary(self, content, content_type, filename=None):
        self.send_response(200)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', len(content))
        if filename:
            self.send_header('Content-Disposition', f'attachment; filename="{filename}"')
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(content)

    def send_json(self, data):
        """Send JSON response."""
        content = json.dumps(data, default=str).encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', len(content))
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(content)


def run_server(port=8000):
    """Run the JECS API server."""
    # Parse scorecards before accepting traffic so the first dashboard request
    # uses the source-aware in-memory cache instead of blocking page render.
    build_performance_dashboard_payload()
    server_address = ('', port)
    httpd = HTTPServer(server_address, JecsAPIHandler)
    print(f"JECS API server running on http://localhost:{port}")
    print(f"Open the dashboard at: http://localhost:{port}/amazon-dsp-kpi-dashboard.html")
    print("Press Ctrl+C to stop the server")
    threading.Thread(
        target=run_connection_scheduler, daemon=True,
        name="startup-connection-health").start()
    httpd.serve_forever()


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='JECS Unified API Server')
    parser.add_argument('--port', type=int, default=8000, help='Port to run the server on (default: 8000)')
    args = parser.parse_args()
    run_server(args.port)
