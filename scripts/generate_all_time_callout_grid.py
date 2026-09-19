#!/usr/bin/env python3
"""Generate the daily-brief repeat-absence grid with all available callout totals."""

from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


WORKBOOK = Path(
    "/Users/claw/.openclaw/media/inbound/"
    "Attendance_Callout_Summary_Apr-Aug_2026---1955d9c4-ea32-4d8c-af92-755888430d59.xlsx"
)
OUTPUT = Path("data/attendance_reports/daily-report-callout-totals-2026-09-02.pdf")

TARGETS = [
    ("Christopher Platas", "Christopher Platas", ["2026-08-24", "2026-08-25"]),
    ("Kionn Beverly", "Kionn Beverley", ["2026-08-25", "2026-08-26"]),
    ("Jada Williams", "Jada Williams", ["2026-08-26", "2026-08-30"]),
    ("Lesley Sicar", "Lesley Sicar", ["2026-08-28"]),
    ("Maria Diaz", "Maria Diaz", ["2026-08-20", "2026-08-30"]),
    ("William Avila", "William Avila", ["2026-08-23"]),
]


def workbook_callouts():
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Attendance Detail"]
    result = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        date, employee, _, _, _, counted, *_ = row
        if counted == "Yes":
            result.setdefault(employee, []).append(date.date().isoformat())
    return result


def pretty_dates(dates):
    from datetime import date

    values = []
    for value in sorted(dates):
        day = date.fromisoformat(value)
        values.append(f"{day.strftime('%b')} {day.day}")
    return ", ".join(values)


def main():
    historical = workbook_callouts()
    rows = []
    for display_name, workbook_name, recent_dates in TARGETS:
        all_dates = sorted(set(historical.get(workbook_name, []) + recent_dates))
        rows.append((display_name, all_dates))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(OUTPUT),
        pagesize=landscape(letter),
        rightMargin=0.42 * inch,
        leftMargin=0.42 * inch,
        topMargin=0.36 * inch,
        bottomMargin=0.36 * inch,
        title="Repeat Absence Callout Totals",
        author="JECS Operations",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "Title2", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=18, leading=21, textColor=colors.HexColor("#17365D"),
        alignment=TA_CENTER, spaceAfter=4,
    )
    subtitle = ParagraphStyle(
        "Subtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=9, leading=11, textColor=colors.HexColor("#4A5568"),
        alignment=TA_CENTER,
    )
    cell = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8.2, leading=10)
    cell_bold = ParagraphStyle("CellBold", parent=cell, fontName="Helvetica-Bold")
    center = ParagraphStyle("Center", parent=cell_bold, alignment=TA_CENTER)

    story = [
        Paragraph("REPEAT ABSENCE CALLOUT TOTALS", title),
        Paragraph("Employees listed in the Daily Operations Brief for September 1, 2026", subtitle),
        Paragraph("All available attendance documentation reviewed: April 28–September 1, 2026", subtitle),
        Spacer(1, 0.16 * inch),
    ]

    data = [[
        Paragraph("Status", center), Paragraph("Employee", center),
        Paragraph("Verified callout dates — all available data", center),
        Paragraph("Total<br/>callouts", center),
    ]]
    for display_name, dates in rows:
        data.append([
            Paragraph("WATCH", center),
            Paragraph(display_name, cell_bold),
            Paragraph(pretty_dates(dates), cell),
            Paragraph(str(len(dates)), center),
        ])

    table = Table(data, colWidths=[0.78*inch, 1.75*inch, 6.42*inch, 0.95*inch], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#FFF2CC")),
        ("TEXTCOLOR", (0, 1), (0, -1), colors.HexColor("#8A5A00")),
        ("BACKGROUND", (-1, 1), (-1, -1), colors.HexColor("#E2F0D9")),
        ("TEXTCOLOR", (-1, 1), (-1, -1), colors.HexColor("#1F5E2C")),
        ("ROWBACKGROUNDS", (1, 1), (-2, -1), [colors.white, colors.HexColor("#F7F9FC")]),
        ("GRID", (0, 0), (-1, -1), 0.55, colors.HexColor("#B7C3D0")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story += [table, Spacer(1, 0.14 * inch)]
    story.append(Paragraph(
        "Totals count only entries explicitly documented as a callout. No-call/no-show and other attendance events are excluded. "
        "William Avila also has one documented no-call/no-show on August 30.",
        ParagraphStyle("Foot", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#4A5568")),
    ))
    doc.build(story)
    print(OUTPUT)
    for name, dates in rows:
        print(f"{name}: {len(dates)} ({pretty_dates(dates)})")


if __name__ == "__main__":
    main()
